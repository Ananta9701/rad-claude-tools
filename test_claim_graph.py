#!/usr/bin/env python3
"""test_claim_graph.py — claim_graph.py 단독 테스트. deck_toolkit 이 없는 프로젝트(저자·리뷰어)에서 그대로 돈다.

    python test_claim_graph.py      # 실패 0 이어야 함
test_toolkit.py(발표 프로젝트)는 이 파일을 import 해 같은 테스트를 함께 돌린다.
"""
import copy, io, os, re, subprocess, sys, traceback
sys.dont_write_bytecode = True   # /mnt/project 는 대화창 안에서 쓰기 가능 — __pycache__ 를 남기지 않는다 (v2.3.1)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import claim_graph as CGm


def _manifest_version(fname):
    """같은 폴더의 TOOLS_MANIFEST.md 에 적힌 판. 없으면 None (Z1 이후: EXPECT_VERSION 만 맞추고 manifest 를 안 올린 사고 방지)."""
    import os, re
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'TOOLS_MANIFEST.md')
    if not os.path.exists(p):
        return None
    m = re.search(r'\| `%s` \| v([0-9.]+)' % re.escape(fname), open(p, encoding='utf8').read())
    return m.group(1) if m else None

EXPECT_VERSION = '16.21'

def t_version_matches_manifest():
    assert getattr(CGm, '__version__', None) == EXPECT_VERSION, (getattr(CGm, '__version__', None), EXPECT_VERSION)
    mv = _manifest_version('claim_graph.py')
    assert mv is None or mv == EXPECT_VERSION, ('TOOLS_MANIFEST.md 의 판', mv, '코드', EXPECT_VERSION)

GRAPH = [
    {'id': 'a', 'statement': 'A', 'evidence': 'e', 'sites': ['slide:7'], 'keys': ['title'],
     'status': 'accepted', 'depends_on': []},
    {'id': 'b', 'statement': 'B', 'evidence': 'e', 'sites': ['notes:7'], 'keys': ['note'],
     'depends_on': [{'id': 'a', 'type': 'premise', 'weight': 1.0}]},
    {'id': 'c', 'statement': 'C', 'evidence': 'e', 'sites': ['notes:8'], 'keys': [],
     'depends_on': [{'id': 'b', 'type': 'context', 'weight': 0.2}]},
]

MD = '/tmp/cg_doc.md'
open(MD, 'w', encoding='utf8').write('''# Results
FA decreased in the lesion group (0.36 vs 0.44, P = 0.014). Therefore marker A loss precedes marker B change.
However, the sample was small (n = 24).
# Discussion
Index X was higher in lesions (Kim et al. 2021).
''')

def _resolve_dict(d):
    return lambda site: d[site]

def t_mapgraph_clean():
    probs, order = CGm.mapgraph(copy.deepcopy(GRAPH), io.StringIO())
    assert not [p for p in probs if not p.startswith('[참고]')], probs
    assert order.index('a') < order.index('b') < order.index('c'), order

def t_mapgraph_unknown_id():
    g = copy.deepcopy(GRAPH); g[1]['depends_on'] = [{'id': 'zzz'}]
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('정의되지 않음' in p for p in probs), probs

def t_mapgraph_cycle_needs_anchor():
    g = copy.deepcopy(GRAPH); g[0]['depends_on'] = [{'id': 'c'}]
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('순환' in p and 'anchor' in p for p in probs), probs
    g[0]['anchor'] = True
    probs, order = CGm.mapgraph(g, io.StringIO())
    assert not any('순환' in p for p in probs), probs
    assert order[0] == 'a', order

def t_mapgraph_superseded_with_sites():
    g = copy.deepcopy(GRAPH); g[2]['status'] = 'superseded'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('superseded' in p for p in probs), probs

def t_mapgraph_supersedes_needs_forbidden():
    g = copy.deepcopy(GRAPH); g[0]['supersedes'] = {'statement': 'the lesion is far denser than metastasis'}
    probs, _ = CGm.mapgraph(g, io.StringIO())
    hit = [p for p in probs if 'forbidden 이 비어' in p]
    assert hit and not hit[0].startswith('[참고]') and 'far denser' in hit[0], probs
    g[0]['forbidden'] = ['far denser']
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not any('forbidden 이 비어' in p for p in probs), probs

def t_mapgraph_role_rules():
    g = copy.deepcopy(GRAPH)
    g[0]['role'] = 'evidence'; g[1]['role'] = 'main'; g[2]['role'] = 'claim'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('a: evidence 인데 걸린 caveat' in p for p in probs), probs
    assert any('c: role=claim 인데 premise 간선이 없음' in p for p in probs), probs
    assert not any('b: role=main 인데' in p for p in probs), probs
    g[0]['depends_on'] = [{'id': 'c', 'type': 'caveat'}]; g[2]['role'] = 'caveat'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not any('걸린 caveat' in p for p in probs), probs

def t_mapgraph_main_role_count():
    g = copy.deepcopy(GRAPH); g[0]['role'] = 'evidence'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('role=main' in p for p in probs), probs
    g[1]['role'] = 'main'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not any('role=main 인 주장이' in p for p in probs), probs

def t_impact_propagates_with_weight():
    rows = CGm.impact(copy.deepcopy(GRAPH), ['a'], io.StringIO())
    d = {cid: s for cid, s, _ in rows}
    assert d['b'] == 1.0 and abs(d['c'] - 0.2) < 1e-9, d

def t_impact_stops_on_cycle():
    g = copy.deepcopy(GRAPH); g[0]['depends_on'] = [{'id': 'c'}]; g[0]['anchor'] = True
    rows = CGm.impact(g, ['a'], io.StringIO())
    assert 'a' not in [r[0] for r in rows]

def t_freeze_stale_with_dict_resolver():
    d = {'slide:7': 'title x', 'notes:7': 'note y', 'notes:8': 'z'}
    g = CGm.mapfreeze(_resolve_dict(d), copy.deepcopy(GRAPH), at='2026-01-01')
    assert CGm.mapstale(_resolve_dict(d), g, io.StringIO())['changed'] == []
    d['notes:7'] = 'note y changed'
    r = CGm.mapstale(_resolve_dict(d), g, io.StringIO())
    assert r['changed'] == ['b'] and [x[0] for x in r['suspect']] == ['c'], r
    g[0]['statement'] = 'A changed'
    assert 'a' in CGm.mapstale(_resolve_dict(d), g, io.StringIO())['changed']
    g.append({'id': 'new', 'statement': '', 'sites': []})
    assert CGm.mapstale(_resolve_dict(d), g, io.StringIO())['unverified'] == ['new']

def t_save_load_claims_roundtrip():
    p = '/tmp/claims_rt.json'
    CGm.save_claims(p, copy.deepcopy(GRAPH), deck_name='X', note='n')
    dn, note, cl = CGm.load_claims_meta(p)
    assert dn == 'X' and [c['id'] for c in cl] == ['a', 'b', 'c']
    assert CGm.load_claims(p)[0]['id'] == 'a'

def t_docsource_resolve():
    src = CGm.DocSource(MD)
    assert src.resolve('doc:p:2').startswith('FA decreased')
    assert 'n = 24' in src.resolve('doc:find:sample was small')
    sec = src.resolve('doc:sec:Results')
    assert 'FA decreased' in sec and 'Index X' not in sec
    try:
        src.resolve('doc:find:the'); raise AssertionError('모호한 find 가 통과됨')
    except KeyError:
        pass

def t_extract_candidates_and_edges():
    src = CGm.DocSource(MD)
    c = CGm.extract(src.units(), stream=io.StringIO())
    fa = next(x for x in c if x['statement'].startswith('FA decreased'))
    assert fa['score'] >= 3 and fa['status'] == 'proposed' and fa['origin'] == 'extract'
    th = next(x for x in c if x['statement'].startswith('Therefore'))
    assert th['depends_on'][0]['id'] == fa['id'] and th['depends_on'][0]['type'] == 'premise'
    hw = next(x for x in c if x['statement'].startswith('However'))
    assert hw.get('role') == 'caveat'

def t_extract_skips_reference_lines():
    units = [('doc:p:1', 'Shin JI, et al. Korean J Radiol. 2024;25(1):62-73.'),
             ('doc:p:2', 'References) Gi T, et al.')]
    assert CGm.extract(units, stream=io.StringIO()) == []

def t_scaffold_orders_by_site():
    buf = io.StringIO()
    sites = CGm.scaffold(copy.deepcopy(GRAPH), buf)
    assert set(sites) == {'slide:7', 'notes:7', 'notes:8'} and '[a]' in buf.getvalue()

def t_docx_mapcheck_freeze_stale():
    src = CGm.DocSource(MD)
    cl = [{'id': 'fa', 'statement': 'FA down', 'evidence': 'e', 'sites': ['doc:find:FA decreased'],
           'keys': ['decreased'], 'forbidden': ['increased'], 'depends_on': []},
          {'id': 'cond', 'statement': 'cond up', 'evidence': 'e', 'sites': ['doc:sec:Discussion'],
           'keys': ['higher'], 'depends_on': [{'id': 'fa', 'type': 'premise'}]}]
    probs, _ = CGm.mapcheck(src.resolve, copy.deepcopy(cl), io.StringIO())
    assert not probs, probs
    g = CGm.mapfreeze(src.resolve, copy.deepcopy(cl), at='2026-01-01')
    assert CGm.mapstale(src.resolve, g, io.StringIO())['changed'] == []
    src.paras[1] = src.paras[1].replace('0.36', '0.37')
    r = CGm.mapstale(src.resolve, g, io.StringIO())
    assert r['changed'] == ['fa'] and [x[0] for x in r['suspect']] == ['cond'], r

def _mk_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph('Abstract').runs[0].bold = True
    d.add_paragraph('Objective: FA decreased (0.36 vs 0.44, P = 0.014).')
    d.add_paragraph('RESULTS')
    d.add_paragraph('Index X was higher in lesions.')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'Measure'; t.cell(0, 1).text = 'Value'
    t.cell(1, 0).text = 'IDX'; t.cell(1, 1).text = '0.612'
    d.add_paragraph('Limitations').runs[0].bold = True
    d.add_paragraph('The sample was small.')
    d.save(path)

def t_docx_heuristic_headings_and_table():
    p = '/tmp/cg_test.docx'; _mk_docx(p)
    src = CGm.DocSource(p)
    heads = [src.paras[i] for i in sorted(src.headings)]
    assert heads == ['Abstract', 'RESULTS', 'Limitations'], heads
    assert len(src.tables) == 1 and '0.612' in src.resolve('doc:tbl:1')
    assert 'higher in lesions' in src.resolve('doc:sec:RESULTS')
    assert 'Value' not in src.resolve('doc:sec:Limitations')

def t_mapdiff_pairs_by_sites_and_keys():
    a = copy.deepcopy(GRAPH)
    b = [{'id': 'x1', 'sites': ['slide:7'], 'keys': ['title'], 'depends_on': []},
         {'id': 'x2', 'sites': ['notes:7'], 'keys': [], 'depends_on': [{'id': 'x1', 'type': 'support', 'weight': 0.7}]},
         {'id': 'x9', 'sites': ['doc:find:zzz'], 'keys': ['reviewer only'], 'depends_on': []},
         {'id': 'lim', 'sites': ['doc:find:limit'], 'keys': [], 'depends_on': []}]
    b[0]['depends_on'] = [{'id': 'lim', 'type': 'caveat', 'weight': 0.5}]
    r = CGm.mapdiff(a, b, 'A', 'B', io.StringIO())
    assert ('a', 'x1') in r['both'] and ('b', 'x2') in r['both'], r['both']
    assert r['only_a'] == ['c'] and set(r['only_b']) == {'x9', 'lim'}, r
    assert any('premise' in d and 'support' in d for d in r['edge_diff']), r['edge_diff']
    assert any('a/x1' in d for d in r['caveat_diff']), r['caveat_diff']

def t_save_claims_preserves_meta_and_doc_wins():
    # v15.4 (#14): 상위 키 보존, doc 이 stale deck 을 이김
    import json
    p = '/tmp/cg_meta.json'
    json.dump({'deck': 'old_v44', 'doc': 'X_v45.docx', 'version': 'v5', 'protocol': 'v15', 'claims': copy.deepcopy(GRAPH)},
              open(p, 'w', encoding='utf8'))
    meta, cl = CGm.load_claims_full(p)
    assert meta == {'deck': 'old_v44', 'doc': 'X_v45.docx', 'version': 'v5', 'protocol': 'v15'}
    dn, _, _ = CGm.load_claims_meta(p)
    assert dn == 'X_v45.docx', dn   # v15.3 까지는 'old_v44' 가 나왔다
    buf, old = io.StringIO(), sys.stdout; sys.stdout = buf
    try:
        CGm.save_claims(p, cl, meta=meta, doc='/x/y/X_v46__3_.docx')
    finally:
        sys.stdout = old
    d = json.load(open(p, encoding='utf8'))
    assert d['doc'] == 'X_v46' and 'deck' not in d   # v15.4.1: 확장자 없음(규약 §3) and d['version'] == 'v5' and d['protocol'] == 'v15', d.keys()
    assert '[경고]' in buf.getvalue() and 'stale' in buf.getvalue()
    # 덱 경로는 그대로 deck 키
    CGm.save_claims(p, cl, deck_name='D.pptx', note='n')
    d = json.load(open(p, encoding='utf8')); assert d['deck'] == 'D.pptx' and d['note'] == 'n'

def t_norm_name_strips_suffix():
    assert CGm.doc_name('/x/DOC_v47_clean__2_.docx') == 'DOC_v47_clean' and CGm.doc_name('DOC_v47_clean') == 'DOC_v47_clean'
    assert CGm.norm_name('/a/b/claim_graph__8_.py') == 'claim_graph.py'
    assert CGm.norm_name('DOC_v45_clean__1_.docx') == 'DOC_v45_clean.docx'
    assert CGm.norm_name('plain.docx') == 'plain.docx' and CGm.norm_name(None) is None

def t_mapdiff_role_mismatch_excluded_and_manual_pairs():
    # v15.4 (#9): 같은 자리라도 caveat↔evidence 는 짝짓지 않는다; --pairs 로 수동 확정
    a = [{'id': 'pv-5mm', 'role': 'caveat', 'sites': ['doc:find:5 mm'], 'keys': [], 'depends_on': []},
         {'id': 'idx-md', 'role': 'evidence', 'sites': ['doc:find:IDX'], 'keys': ['idx md'], 'depends_on': []}]
    b = [{'id': 'ev-null', 'role': 'evidence', 'sites': ['doc:find:5 mm'], 'keys': [], 'depends_on': []},
         {'id': 'ev-idx-md', 'role': 'evidence', 'sites': ['doc:find:idx'], 'keys': [], 'depends_on': []}]
    r = CGm.mapdiff(a, b, 'A', 'B', io.StringIO())
    assert ('pv-5mm', 'ev-null') not in r['both'] and ('idx-md', 'ev-idx-md') in r['both'], r['both']
    r = CGm.mapdiff(a, b, 'A', 'B', io.StringIO(), pairs={'pv-5mm': 'ev-null'})
    assert ('pv-5mm', 'ev-null') in r['both'] and r['only_a'] == [] and r['only_b'] == [], r
    # role 이 한쪽에만 있으면 종전대로 짝짓는다
    del a[0]['role']
    r = CGm.mapdiff(a, b, 'A', 'B', io.StringIO())
    assert ('pv-5mm', 'ev-null') in r['both']

def t_remap_refs():
    # v15.4 (#4): 매핑 {29:26,30:27,31:28,13:None,12:12}
    cl = [{'id': 'wb', 'statement': 'Prior work [29-31] and [13] agree.', 'evidence': 'refs 29-31',
           'sites': [], 'keys': ['29-31', 'refs 29-31', 'whole brain'], 'depends_on': []},
          {'id': 'n', 'statement': 'n = 200 [12]', 'evidence': '', 'sites': [], 'keys': ['42', '99'], 'depends_on': [],
           'verified': {'at': '2026-01-01'}}]
    refmap = {29: 26, 30: 27, 31: 28, 13: None, 12: 12}
    buf = io.StringIO()
    ch = CGm.remap_refs(copy.deepcopy(cl), refmap, stream=buf)
    c2 = copy.deepcopy(cl); CGm.remap_refs(c2, refmap, stream=io.StringIO())
    assert c2[0]['statement'] == 'Prior work [26-28] and agree.', c2[0]['statement']
    assert c2[0]['keys'] == ['29-31', 'refs 26-28', 'whole brain'], c2[0]['keys']   # 순수 숫자 key 는 절대 보존 (v15.4.1)
    assert c2[1]['statement'] == 'n = 200 [12]' and c2[1]['keys'] == ['42', '99']
    assert '삭제된 문헌' in buf.getvalue() and "순수 숫자 key" in buf.getvalue()
    c4 = copy.deepcopy(cl); buf = io.StringIO(); CGm.remap_refs(c4, {**refmap, 12: 11}, stream=buf)
    assert c4[1]['statement'] == 'n = 200 [11]' and 'mapfreeze' in buf.getvalue()

def t_mapcheck_refs_key_matches_bracket_citation():
    # v15.4.2: `refs 26-28` key 는 자리의 [26-28] / [26–28] / [24,26-28] 인용으로 적중, [26,27] 만 있으면 실패
    texts = {'doc:find:a': 'Prior work [26-28] showed.', 'doc:find:b': 'Prior work [24,26–28] showed.',
             'doc:find:c': 'Prior work [26,27] showed.', 'doc:find:d': 'no citation here'}
    def mk(site, key):
        return [{'id': 'x', 'statement': 's', 'evidence': 'e', 'sites': [site], 'keys': [key], 'depends_on': []}]
    ok = lambda site, key: not [p for p in CGm.mapcheck(_resolve_dict(texts), mk(site, key), io.StringIO())[0] if '반영되지 않음' in p]
    assert ok('doc:find:a', 'refs 26-28') and ok('doc:find:b', 'refs 26-28') and ok('doc:find:a', 'refs: 27')
    assert not ok('doc:find:c', 'refs 26-28') and not ok('doc:find:d', 'refs 26-28')
    assert ok('doc:find:a', 'prior work') and ok('doc:find:a', '26-28') and not ok('doc:find:a', '26-29')   # 리터럴 key 는 종전대로 부분문자열

def t_remap_refs_then_mapcheck_passes():
    # 리뷰어 요청 테스트: remap-refs 적용 후 mapcheck 통과
    texts = {'doc:find:p': 'Previous studies [26-28] reported.'}
    cl = [{'id': 'wb', 'statement': 'reported [29-31]', 'evidence': 'refs 29-31', 'sites': ['doc:find:p'],
           'keys': ['refs 29-31'], 'depends_on': []}]
    CGm.remap_refs(cl, {29: 26, 30: 27, 31: 28}, stream=io.StringIO())
    assert cl[0]['keys'] == ['refs 26-28']
    probs, _ = CGm.mapcheck(_resolve_dict(texts), cl, io.StringIO())
    assert not [p for p in probs if '반영되지 않음' in p], probs

def t_nums_sep_checks_only_prefix():
    # v15.5 리뷰어 요청: evidence "원고 | 재현" 에서 구분자 앞쪽만 --nums
    texts = {'doc:find:p': 'IDX 0.612 vs 0.523 in Table 2.'}
    cl = [{'id': 'x', 'statement': 's', 'evidence': '원고 0.612 vs 0.523 | 재현 0.612 vs 0.523, r 0.843', 'sites': ['doc:find:p'],
           'keys': ['idx'], 'depends_on': []}]
    probs, _ = CGm.mapcheck(_resolve_dict(texts), cl, io.StringIO(), nums=True)
    assert any('0.843' in p for p in probs), probs
    probs, _ = CGm.mapcheck(_resolve_dict(texts), cl, io.StringIO(), nums=True, nums_sep='|')
    assert not [p for p in probs if '수치' in p or '0.843' in p], probs

def t_mapstale_detects_keys_change_only_when_hash_present():
    # v15.5 리뷰어 §2-2: keys_hash 별도 필드 — 구판 freeze 는 무영향
    texts = {'doc:find:p': 'Prior work [26-28] reported.'}
    cl = [{'id': 'wb', 'statement': 's', 'evidence': 'e', 'sites': ['doc:find:p'], 'keys': ['prior work'], 'depends_on': []}]
    CGm.mapfreeze(_resolve_dict(texts), cl)
    assert 'keys' in cl[0]['verified']
    r = CGm.mapstale(_resolve_dict(texts), cl, io.StringIO()); assert r['changed'] == []
    cl[0]['keys'] = ['refs 26-28']
    buf = io.StringIO(); r = CGm.mapstale(_resolve_dict(texts), cl, buf)
    assert r['changed'] == ['wb'] and 'keys 가 바뀜' in buf.getvalue()
    del cl[0]['verified']['keys']            # 구판 freeze 흉내
    r = CGm.mapstale(_resolve_dict(texts), cl, io.StringIO()); assert r['changed'] == []

def t_mapdiff_one_to_many_pairs():
    a = [{'id': 's10', 'role': 'evidence', 'sites': ['doc:find:S10'], 'keys': ['s10'], 'depends_on': []}]
    b = [{'id': 'nested', 'role': 'evidence', 'sites': ['doc:find:nested'], 'keys': [], 'depends_on': []},
         {'id': 'delta-r2', 'role': 'evidence', 'sites': ['doc:find:delta'], 'keys': [], 'depends_on': []}]
    r = CGm.mapdiff(a, b, 'A', 'B', io.StringIO(), pairs={'s10': ['nested', 'delta-r2']})
    assert r['only_a'] == [] and r['only_b'] == [] and r['both'][0][0] == 's10', r
    assert CGm.parse_pairs('s10=nested,s10=delta-r2') == {'s10': ['nested', 'delta-r2']}

def t_mapgraph_keys_forbidden_overlap():
    g = copy.deepcopy(GRAPH)
    g[0]['keys'] = list(g[0].get('keys', [])) + ['수리 권고']; g[0]['forbidden'] = ['수리 권고']; g[0]['supersedes'] = {'statement': 'old'}
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('keys 와 forbidden 에 같은 표현' in p for p in probs), probs

def t_selfcheck():
    import hashlib, shutil, tempfile, subprocess
    d = tempfile.mkdtemp()
    src = os.path.dirname(os.path.abspath(CGm.__file__))
    for f in ('claim_graph.py', 'test_claim_graph.py'):
        shutil.copy(os.path.join(src, f), d)
    h = lambda f: hashlib.sha256(open(os.path.join(d, f), 'rb').read()).hexdigest()[:12]
    man = ('# TOOLS_MANIFEST\n\n**manifest 판: v9 · 릴리스 v9.9 · 2026-01-01**\n\n## 1.\n\n'
           '| 파일 | 판 | 해시 | 테스트 | 크기 |\n|---|---|---|---|---|\n'
           '| `claim_graph.py` | v%s | `%s` | — | 1 |\n| `test_claim_graph.py` | v%s 동반 | `%s` | — | 1 |\n'
           '| `deck_toolkit.py` | v1 | `000000000000` | — | 1 |\n\n'
           '## 2. 배포표\n\n| 파일 | 저자 | 발표 |\n|---|:-:|:-:|\n'
           '| claim_graph.py · test_claim_graph.py | ○ | ○ |\n| deck_toolkit.py | × | ○ |\n'
           '| TOOLS_MANIFEST.md (이 파일) · RELEASE.md | ○ | ○ |\n\n## 3. 절차\n'
           % (CGm.__version__, h('claim_graph.py'), CGm.__version__, h('test_claim_graph.py')))
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'w', encoding='utf8').write(man)
    mh = hashlib.sha256(open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'rb').read()).hexdigest()[:12]
    rel_ok = ('# RELEASE v9.9 — manifest v9 — x\n\n<!-- sets:begin -->\n### 저자 (3)\n| 파일 | 판 | 해시 | 변경 |\n|---|---|---|---|\n'
              '| `claim_graph.py` | v%s | `%s` | — |\n| `TOOLS_MANIFEST.md` | v9 | `%s` | ○ |\n<!-- sets:end -->\n' % (CGm.__version__, h('claim_graph.py'), mh))
    open(os.path.join(d, 'RELEASE.md'), 'w', encoding='utf8').write(rel_ok)
    for junk in ('TOOLS_MANIFEST_v3.md', 'x__2_.py', '260910_manifest만_발표_v9.md', 'DOC_v47_claims.json'):
        open(os.path.join(d, junk), 'w').write('x')
    buf = io.StringIO(); r = CGm.selfcheck(d, stream=buf)
    assert r['ok'] and r['project'] == '저자', (r, buf.getvalue())
    assert set(r['delete']) == {'TOOLS_MANIFEST_v3.md', 'x__2_.py', '260910_manifest만_발표_v9.md'} and r['other'] == ['DOC_v47_claims.json'], r
    assert '②′ RELEASE §3 대조' in buf.getvalue()
    # v15.7: --role — 전체 세트(GitHub)처럼 다른 역할 도구가 같이 있어도 역할대로 보고, 안 쓰는 도구는 삭제 후보가 아니다
    open(os.path.join(d, 'deck_toolkit.py'), 'w').write("__version__ = '1'\n")
    buf = io.StringIO(); r = CGm.selfcheck(d, stream=buf, role='저자')
    assert r['project'] == '저자' and 'deck_toolkit.py' not in r['delete'] and '역할(지정)' in buf.getvalue(), (r, buf.getvalue())
    assert CGm.selfcheck(d, stream=io.StringIO(), role='없는역할')['ok'] is False
    # --compare: 예비 폴더와 대조
    d2 = tempfile.mkdtemp()
    for f in ('claim_graph.py', 'TOOLS_MANIFEST.md'):
        shutil.copy(os.path.join(d, f), d2)
    buf = io.StringIO(); CGm.selfcheck(d, stream=buf, role='저자', compare=d2)
    assert '예비 폴더' in buf.getvalue() and '없는 파일 2' in buf.getvalue() and '다름' in buf.getvalue(), buf.getvalue()
    os.remove(os.path.join(d, 'deck_toolkit.py'))
    # v15.8: git 저장소 폴더면 받은 커밋을 적고, .git 은 세트 밖 파일로 잡지 않는다
    r0 = subprocess.run(['git', 'init', '-q', d], capture_output=True)
    if r0.returncode == 0:
        subprocess.run(['git', '-C', d, '-c', 'user.email=t@t', '-c', 'user.name=t', 'add', '-A'], capture_output=True)
        subprocess.run(['git', '-C', d, '-c', 'user.email=t@t', '-c', 'user.name=t', 'commit', '-qm', 'x'], capture_output=True)
        buf = io.StringIO(); r = CGm.selfcheck(d, stream=buf, role='저자')
        assert '받은 커밋' in buf.getvalue() and '.git' not in r['other'], buf.getvalue()
    # ②′: manifest 를 바꿔치기하면(자기 기준으로는 ○) RELEASE 대조가 잡는다
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'a', encoding='utf8').write('\n<!-- tampered -->\n')
    r = CGm.selfcheck(d, stream=io.StringIO()); assert any('②′ TOOLS_MANIFEST.md' in p for p in r['problems']), r['problems']
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'w', encoding='utf8').write(man)
    open(os.path.join(d, 'RELEASE.md'), 'w', encoding='utf8').write('# RELEASE v9.8 — manifest v8 — x\n')
    r = CGm.selfcheck(d, stream=io.StringIO()); assert not r['ok'] and any('1단계' in p for p in r['problems'])
    s = open(os.path.join(d, 'claim_graph.py'), encoding='utf8').read().replace("__version__ = '%s'" % CGm.__version__, "__version__ = '0.0'")
    open(os.path.join(d, 'claim_graph.py'), 'w', encoding='utf8').write(s)
    r = CGm.selfcheck(d, stream=io.StringIO()); assert any('2단계' in p for p in r['problems']) and any('3단계' in p for p in r['problems'])

def t_selfcheck_tests_leave_no_pycache():
    # 리뷰어 v2.3 수령 [결함]: /mnt/project 는 대화창 안에서 쓰기 가능 — selfcheck --tests 가 __pycache__ 를 남겼다
    import hashlib, shutil, tempfile
    d = tempfile.mkdtemp()
    shutil.copy(CGm.__file__, d)
    open(os.path.join(d, 'test_x.py'), 'w').write("EXPECT_VERSION = '1'\nimport sys, os\nsys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))\nimport claim_graph\nprint('통과 1 / 실패 0  (전체 1)')\n")
    h = lambda f: hashlib.sha256(open(os.path.join(d, f), 'rb').read()).hexdigest()[:12]
    man = ('**manifest 판: v1 · 릴리스 v1 · x**\n\n| `claim_graph.py` | v%s | `%s` | — | 1 |\n| `test_x.py` | v1 동반 | `%s` | — | 1 |\n\n'
           '## 2.\n\n| 파일 | 저자 |\n|---|:-:|\n| claim_graph.py · test_x.py | ○ |\n| TOOLS_MANIFEST.md · RELEASE.md | ○ |\n\n## 3.\n' % (CGm.__version__, h('claim_graph.py'), h('test_x.py')))
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'w', encoding='utf8').write(man)
    mh = h('TOOLS_MANIFEST.md')
    open(os.path.join(d, 'RELEASE.md'), 'w', encoding='utf8').write('# RELEASE v1 — manifest v1 — x\n\n### 저자 (3)\n| `TOOLS_MANIFEST.md` | v1 | `%s` | ○ |\n<!-- sets:end -->\n' % mh)
    r = CGm.selfcheck(d, run_tests=True, stream=io.StringIO())
    assert r['ok'], r['problems']
    assert not os.path.exists(os.path.join(d, '__pycache__')), os.listdir(d)

def t_selfcheck_code_only_release_hash():
    # v15.8.2 (코드 v2.43): 코드 전용 5개(HISTORY 등)는 비공개 저장소에서 릴리스 사이에도 바뀐다 — ②′ RELEASE 대조에서 뺀다.
    # 공개 파일(claim_graph.py)은 그대로 잡는다
    import hashlib, shutil, tempfile
    code_only = ('HISTORY.md', 'PRIVATE_TERMS.txt', 'CODE_PROJECT_README.md', 'release.py', 'GITHUB_README.md')
    d = tempfile.mkdtemp()
    shutil.copy(CGm.__file__, d)
    h = lambda f: hashlib.sha256(open(os.path.join(d, f), 'rb').read()).hexdigest()[:12]
    man = ('**manifest 판: v1 · 릴리스 v1 · x**\n\n| `claim_graph.py` | v%s | `%s` | — | 1 |\n\n'
           '## 2.\n\n| 파일 | 저자 |\n|---|:-:|\n| claim_graph.py | ○ |\n| TOOLS_MANIFEST.md · RELEASE.md | ○ |\n\n## 3.\n' % (CGm.__version__, h('claim_graph.py')))
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'w', encoding='utf8').write(man)
    for f in code_only:
        open(os.path.join(d, f), 'w', encoding='utf8').write('릴리스 뒤에 고친 내용\n')
    rows = ''.join('| `%s` | — | `000000000000` | ○ |\n' % f for f in code_only)
    def rel(cg_hash):
        open(os.path.join(d, 'RELEASE.md'), 'w', encoding='utf8').write(
            '# RELEASE v1 — manifest v1 — x\n\n### 저자 (3)\n| `claim_graph.py` | v1 | `%s` | — |\n| `TOOLS_MANIFEST.md` | v1 | `%s` | ○ |\n%s<!-- sets:end -->\n'
            % (cg_hash, h('TOOLS_MANIFEST.md'), rows))
    rel(h('claim_graph.py'))          # 성공 길: 코드 전용 해시가 모두 달라도 통과
    buf = io.StringIO(); r = CGm.selfcheck(d, stream=buf)
    assert r['ok'] and not any('②′' in p for p in r['problems']), (r['problems'], buf.getvalue())
    rel('000000000000')               # 실패 길: 공개 파일 해시가 다르면 여전히 잡는다
    r = CGm.selfcheck(d, stream=io.StringIO())
    assert not r['ok'] and any('②′ claim_graph.py' in p for p in r['problems']), r['problems']
    assert not any(('②′ %s' % f) in p for f in code_only for p in r['problems']), r['problems']
    assert tuple(CGm.CODE_ONLY) == code_only, CGm.CODE_ONLY

def t_v1583_selfcheck_unreadable_manifest_fails():
    # 코드 리뷰 09-28 [결함]: 판을 읽을 수 없는 manifest(판 없음 = 판 없음)로 '통과 — 작업 시작 가능' 이 나왔다
    import tempfile
    d = tempfile.mkdtemp()
    open(os.path.join(d, 'TOOLS_MANIFEST.md'), 'w', encoding='utf8').write('garbage\n')
    buf = io.StringIO(); r = CGm.selfcheck(d, stream=buf)
    assert not r['ok'] and any('판을 읽지 못했다' in p for p in r['problems']), (r, buf.getvalue())
    assert '통과 — 작업 시작 가능' not in buf.getvalue()


def t_v1584_unreadable_sites_not_unchanged():
    # 코드 리뷰 09-28 ⑦: 읽을 수 없는 자리(지운 슬라이드·바뀐 절 제목)가 freeze·stale 두 번 다 None 이면 '바뀐 것 없음'
    import copy
    texts = {'s:1': 'alpha text', 's:2': 'beta text'}
    def resolve(site):
        return texts[site]
    g = [{'id': 'x', 'statement': 'X', 'evidence': 'e', 'sites': ['s:1', 's:gone']}]
    try:                                                   # 실패 길: 읽을 수 없는 자리가 있으면 검증 기록을 하지 않는다
        CGm.mapfreeze(resolve, copy.deepcopy(g), at='2026-01-01'); assert False, '멈추지 않았다'
    except SystemExit as e:
        assert 's:gone' in str(e) and '읽지 못' in str(e), e
    old = copy.deepcopy(g)                                 # 구판 freeze 가 None 을 적어 둔 그래프 — 여전히 못 읽으면 알린다
    old[0]['verified'] = {'at': '2026-01-01', 'sites': {'s:1': CGm._fingerprint('alpha text'), 's:gone': None},
                          'evidence': CGm._fingerprint('e|X')}
    buf = io.StringIO(); r = CGm.mapstale(resolve, old, buf)
    assert r['changed'] == ['x'] and '읽을 수 없다' in buf.getvalue(), (r, buf.getvalue())
    ok = [{'id': 'y', 'statement': 'Y', 'evidence': 'e', 'sites': ['s:1', 's:2']}]   # 성공 길: 다 읽히면 전처럼
    CGm.mapfreeze(resolve, ok, at='2026-01-01')
    r = CGm.mapstale(resolve, ok, io.StringIO()); assert r['changed'] == [] and r['unverified'] == [], r


def t_claim_graph_cli():
    p = '/tmp/cg_cli.json'
    CGm.save_claims(p, copy.deepcopy(GRAPH))
    import json
    json.dump({'map': {'1': 1}}, open('/tmp/cg_refmap.json', 'w'))
    r = subprocess.run([sys.executable, CGm.__file__, 'remap-refs', '--claims', p, '--map', '/tmp/cg_refmap.json', '-o', '/tmp/cg_cli2.json'], capture_output=True, text=True)
    assert r.returncode == 0 and os.path.exists('/tmp/cg_cli2.json'), r.stderr[-300:]
    r = subprocess.run([sys.executable, CGm.__file__, 'mapdiff', p, p, '--pairs', 'a=a'], capture_output=True, text=True)
    assert r.returncode == 0 and '(a) 양쪽이 잡은 주장 3개' in r.stdout, r.stdout[-300:] + r.stderr[-300:]
    r = subprocess.run([sys.executable, CGm.__file__, 'impact', '--claims', p, 'a', '--sites'], capture_output=True, text=True)
    assert r.returncode == 0 and r.stdout.strip() and '===' not in r.stdout, r.stdout
    # v15.4.1: 두 번 적용 중단, --force 로 통과
    r = subprocess.run([sys.executable, CGm.__file__, 'remap-refs', '--claims', '/tmp/cg_cli2.json', '--map', '/tmp/cg_refmap.json', '-o', '/tmp/cg_cli3.json'], capture_output=True, text=True)
    assert r.returncode == 2 and '이미 적용' in r.stdout, r.stdout
    r = subprocess.run([sys.executable, CGm.__file__, 'remap-refs', '--claims', '/tmp/cg_cli2.json', '--map', '/tmp/cg_refmap.json', '-o', '/tmp/cg_cli3.json', '--force'], capture_output=True, text=True)
    assert r.returncode == 0 and len(json.load(open('/tmp/cg_cli3.json'))['refs_maps_applied']) == 2
    # v15.4.1: --pairs 텍스트 파일·잘못된 형식
    open('/tmp/cg_pairs.txt', 'w').write('# 주석\na=a\nb=b\n')
    r = subprocess.run([sys.executable, CGm.__file__, 'mapdiff', p, p, '--pairs', '/tmp/cg_pairs.txt'], capture_output=True, text=True)
    assert r.returncode == 0 and '(a) 양쪽이 잡은 주장 3개' in r.stdout, r.stderr[-300:]
    open('/tmp/cg_pairs_bad.txt', 'w').write('nonsense\n')
    r = subprocess.run([sys.executable, CGm.__file__, 'mapdiff', p, p, '--pairs', '/tmp/cg_pairs_bad.txt'], capture_output=True, text=True)
    assert r.returncode != 0 and '--pairs 파일 형식' in (r.stderr + r.stdout), r.stderr
    here = os.path.dirname(os.path.abspath(__file__))
    for args in (['mapgraph', '--claims', p], ['scaffold', '--claims', p], ['mapreport', '--claims', p],
                 ['extract', MD, '-o', '/tmp/cg_ex.json'], ['mapdiff', p, p], ['mapcheck', MD, '--claims', p, '--nums'] if False else ['mapdiff', p, p],
                 ['mapcheck', MD, '--claims', p] if False else ['impact', '--claims', p, 'a'],
                 ['mapfreeze', MD, '--claims', '/tmp/cg_docg.json', '-o', '/tmp/cg_fz.json'],
                 ['mapstale', MD, '--claims', '/tmp/cg_fz.json']):
        if args[0] == 'mapfreeze':   # v15.8.4: 원고(md)에서 읽히는 자리로 — 덱 자리(slide:7) 그래프는 이제 멈춘다(아래)
            json.dump({'doc': 'cg_doc', 'claims': [{'id': 'a', 'statement': 'A', 'evidence': 'e', 'sites': ['doc:sec:Results']}]},
                      open('/tmp/cg_docg.json', 'w', encoding='utf8'))
        r = subprocess.run([sys.executable, os.path.join(here, 'claim_graph.py')] + args,
                           capture_output=True, text=True)
        assert r.returncode == 0, (args, r.stderr[-400:])
    r = subprocess.run([sys.executable, os.path.join(here, 'claim_graph.py'), 'mapfreeze', MD, '--claims', p, '-o', '/tmp/cg_fz2.json'],
                       capture_output=True, text=True)
    assert r.returncode != 0 and 'slide:7' in (r.stdout + r.stderr), (r.returncode, r.stderr[-300:])   # 읽지 못하는 자리 → 멈춤


def t_confidence_and_weight_rules():
    g = copy.deepcopy(GRAPH)
    g[0]['confidence'] = 'low'; g[1]['confidence'] = 'bogus'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('약한 고리: b 가 기대는 a' in p for p in probs), probs
    assert any('confidence "bogus"' in p and not p.startswith('[참고]') for p in probs), probs
    g[1]['confidence'] = 'mid'; g[1]['depends_on'][0]['weight'] = 0.9
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert any('weight 0.90 는 type 기본값' in p and p.startswith('[참고]') for p in probs), probs
    assert not any('약한 고리: c' in p for p in probs), probs      # context 간선은 약한 고리 아님

def t_numtokens_and_mapcheck_nums():
    assert CGm._numtokens('Shin KJR 2024;25(1):62-73 (n=1079)') == ['1079']
    assert CGm._numtokens('Ulano AJR 2016;207(2):362-368') == []
    assert CGm._numtokens('0.72 vs 0.61, P = 0.004, 784±162 HU') == ['0.72', '0.61', '0.004', '784', '162']
    src = CGm.DocSource(MD)
    cl = [{'id': 'fa', 'statement': 'x', 'evidence': 'FA 0.36 vs 0.44, P = 0.014',
           'sites': ['doc:find:FA decreased'], 'keys': ['decreased'], 'depends_on': []}]
    probs, _ = CGm.mapcheck(src.resolve, copy.deepcopy(cl), io.StringIO(), nums=True)
    assert not probs, probs
    cl[0]['evidence'] = 'FA 0.39 vs 0.45'
    probs, _ = CGm.mapcheck(src.resolve, cl, io.StringIO(), nums=True)
    assert any('수치 0.39' in p and p.startswith('[참고]') for p in probs), probs
    probs, _ = CGm.mapcheck(src.resolve, cl, io.StringIO())
    assert not probs, probs                                        # nums 안 켜면 안 본다


# ---------------------------------------------------------------- v16 (5판: 근거 칸 · 반박 · 짝 저장 · 이력 · 관계도)
def _probs(g):
    return CGm.mapgraph(g, io.StringIO())[0]


def t_v16_sources_validation():
    g = copy.deepcopy(GRAPH)
    g[0]['sources'] = [{'kind': '문헌', 'what': '10.1000/abc', 'at': '[p.5 · PDF 5]', 'element': '방향', 'verdict': '부합'},
                       {'kind': '교과서', 'what': '시험책', 'at': 'p.56'}, {'kind': '덱', 'at': 'slide@260'}]
    assert not [p for p in _probs(g) if 'sources' in p], _probs(g)                 # 성공 길: 맞는 근거 칸은 조용
    g[0]['sources'] = [{'kind': '잡지', 'what': 'x', 'at': 'y'}, {'kind': '문헌', 'at': 'p.1'}, {'kind': '문헌', 'what': 'd', 'at': 'p.1', 'verdict': '맞음'},
                       {'kind': '덱', 'at': 'slide:7'}, {'kind': '문헌', 'what': 'd'}, 'x']
    ps = _probs(g)
    assert any('kind "잡지"' in p for p in ps) and any('what' in p and '비어' in p for p in ps) and any('verdict "맞음"' in p for p in ps), ps
    assert any(p.startswith('[참고]') and 'slide@sldId' in p for p in ps) and any(p.startswith('[참고]') and 'at(' in p for p in ps), ps
    assert any('sources[6]' in p for p in ps), ps
    g[0]['sources'] = 'x'
    assert any('목록이어야' in p for p in _probs(g))


def t_v16_rebuttal_edge():
    g = copy.deepcopy(GRAPH)
    g.append({'id': 'r', 'statement': 'R', 'role': 'rebuttal', 'depends_on': []})
    g[1]['depends_on'].append({'id': 'r', 'type': 'rebuttal'})
    assert not [p for p in _probs(g) if not p.startswith('[참고]')], _probs(g)   # 성공 길: 새 type 을 받는다
    rows = dict((cid, s) for cid, s, _ in CGm.impact(g, ['r'], io.StringIO()))
    assert abs(rows['b'] - 0.5) < 1e-9 and 'c' in rows, rows                          # caveat 과 같은 무게로 하류까지
    g[1]['depends_on'][-1]['type'] = 'caveat'
    assert any('type rebuttal' in p and p.startswith('[참고]') for p in _probs(g))      # 한계로 적은 반박 노드는 알린다
    g[1]['depends_on'][-1]['type'] = 'rebut'
    assert any('type "rebut"' in p for p in _probs(g))                                  # 실패 길: 모르는 type
    g[1]['depends_on'][-1]['type'] = 'rebuttal'
    g[0]['sources'] = [{'kind': '문헌', 'what': 'd', 'at': 'p.1', 'verdict': '반대 방향'}]
    assert any('반박 노드' in p for p in _probs(g))                                      # a 에는 rebuttal 간선이 없다
    g[1]['sources'] = [{'kind': '문헌', 'what': 'd', 'at': 'p.1', 'verdict': '반대 방향'}]
    assert not any(p.startswith('[참고] b:') and '반박 노드' in p for p in _probs(g))


def t_v16_supersedes_history():
    g = copy.deepcopy(GRAPH)
    g[0]['supersedes'] = [{'statement': 'whole sample alpha beta', 'retracted': '2026-08-01'},
                          {'statement': 'subgroup gamma delta', 'retracted': '2026-09-01'}]
    ps = _probs(g)
    assert any('supersedes 가 있는데 forbidden' in p and 'gamma' in p for p in ps), ps  # 가장 최근 것에서 후보
    assert CGm._supersedes({'supersedes': {'statement': 's'}}) == [{'statement': 's'}] and CGm._supersedes({}) == []


def _store(root, doi='10.1000/abc', sha='aaaa', md=None):
    d = os.path.join(root, doi.replace('/', '_')); os.makedirs(d, exist_ok=True)
    open(os.path.join(d, 'meta.md'), 'w', encoding='utf8').write('<!-- lit: doi=%s pages=2 blank=0 sha=%s -->\n# t\n' % (doi, sha))
    open(os.path.join(d, 'paper.md'), 'w', encoding='utf8').write(md)
    return d


OLD_PDF_MD = '# t\n\n> 원문 PDF 의 글자층(literature.py v0.6). 쪽 표지 [p.N].\n\n[p.1]\n\nIntro words here.\n\n[p.2]\n\nThe sensitivity was 92 percent in lesions.\n'
NEW_PDF_MD = '# t\n\n> 원문 PDF 의 글자층(literature.py v0.8.1). 쪽 표지 [p.인쇄 · PDF N].\n\n[p.e11 · PDF 1]\n\nIntro words here.\n\n[p.e12 · PDF 2]\n\nThe sensitivity was 92 percent in lesions.\n'


def t_v16_mapstale_sources():
    import tempfile
    root = tempfile.mkdtemp(prefix='cgsrc_')
    d = _store(root, md=OLD_PDF_MD)
    cl = [{'id': 's', 'statement': 'S', 'evidence': 'e', 'sites': [], 'keys': [], 'depends_on': [],
           'sources': [{'kind': '문헌', 'what': '10.1000/abc', 'at': '[p.2]', 'verdict': '부합'},
                       {'kind': '문헌', 'what': '10.1000/abc', 'at': '[p.1]', 'element': 'x', 'verdict': '근거 없음'}]},
          {'id': 't', 'statement': 'T', 'sites': [], 'keys': [], 'depends_on': [{'id': 's', 'type': 'premise'}]}]
    CGm.mapfreeze(lambda s: '', cl, sources=root)
    assert len(cl[0]['verified']['sources']) == 2, cl[0]['verified']
    # 1. 쪽 표지 형식·머리말만 바뀐 새 ingest(v0.7+) — 알림 없음, '형식만' 한 줄
    open(os.path.join(d, 'paper.md'), 'w', encoding='utf8').write(NEW_PDF_MD)
    out = io.StringIO(); r = CGm.mapstale(lambda s: '', cl, out, sources=root)
    assert not r['changed'] and len(r['format_only']) == 2 and '검증 이후 바뀐 것 없음' in out.getvalue(), out.getvalue()
    # 2. 원 파일 같음(sha)인데 1쪽 글이 늘었다(변환 개선) — '변환', 하류로 안 번짐, '근거 없음' 판정만 다시 볼 것
    open(os.path.join(d, 'paper.md'), 'w', encoding='utf8').write(NEW_PDF_MD.replace('Intro words here.', 'Intro words here. - listed item now visible'))
    out = io.StringIO(); r = CGm.mapstale(lambda s: '', cl, out, sources=root)
    assert not r['changed'] and r['converted'] and len(r['recheck']) == 1 and '근거 없음' in r['recheck'][0], out.getvalue()
    # 2-1. 원 파일 같음인데 새 변환에 그 자리 표지가 없다 — [변경] 이 아니라 '변환' + at 을 고치라는 줄
    open(os.path.join(d, 'paper.md'), 'w', encoding='utf8').write('# t\n\n[§ Whole]\n\nIntro words here. The sensitivity was 92 percent in lesions.\n')
    out = io.StringIO(); r = CGm.mapstale(lambda s: '', cl, out, sources=root)
    assert not r['changed'] and any('at 을 새 표지로' in x for x in r['recheck']), out.getvalue()
    # 3. 원 파일이 바뀌고(다른 PDF) 2쪽 글도 다름 — [변경], 하류 t 까지
    _store(root, sha='bbbb', md=NEW_PDF_MD.replace('92 percent', '90 percent'))
    out = io.StringIO(); r = CGm.mapstale(lambda s: '', cl, out, sources=root)
    assert r['changed'] == ['s'] and any(x[0] == 't' for x in r['suspect']) and '원 파일도 다름' in out.getvalue(), out.getvalue()
    # 4. 폴더를 주지 않으면 근거 원문은 보지 않는다(전처럼)
    r = CGm.mapstale(lambda s: '', cl, io.StringIO())
    assert not r['changed'], r
    # 5. 보관소에서 그 DOI 가 없어졌다 — [변경](조용히 '같음' 이 되지 않는다)
    import shutil; shutil.rmtree(d)
    out = io.StringIO(); r = CGm.mapstale(lambda s: '', cl, out, sources=root)
    assert r['changed'] == ['s'] and '찾을 수 없음' in out.getvalue(), out.getvalue()
    # 6. 폴더 없이 다시 freeze 하면 전 근거 기록을 그대로 둔다
    CGm.mapfreeze(lambda s: '', cl)
    assert len(cl[0]['verified']['sources']) == 2


def t_v16_sources_textbook_and_section():
    import tempfile
    root = tempfile.mkdtemp(prefix='cgtb_')
    bd = os.path.join(root, '01_시험책'); os.makedirs(bd)
    open(os.path.join(bd, '03_장.md'), 'w', encoding='utf8').write('[p.55 · PDF 69]\n\nalpha\n\n[p.56 · PDF 70]\n\nbeta finding.\n')
    cl = [{'id': 'q', 'statement': 'Q', 'sites': [], 'keys': [], 'depends_on': [],
           'sources': [{'kind': '교과서', 'what': '시험책', 'at': 'p.56'}, {'kind': '교과서', 'what': '없는책', 'at': 'p.1'}]}]
    out = io.StringIO(); CGm.mapfreeze(lambda s: '', cl, sources=root, stream=out)
    assert list(cl[0]['verified']['sources']) == ['교과서|시험책|p.56|'] and '없는책' in out.getvalue(), out.getvalue()   # 못 찾은 근거는 알리고 기록 안 함
    open(os.path.join(bd, '03_장.md'), 'w', encoding='utf8').write('[p.55 · PDF 69]\n\nalpha changed\n\n[p.56 · PDF 70]\n\nbeta finding.\n')
    assert not CGm.mapstale(lambda s: '', cl, io.StringIO(), sources=root)['changed']      # 다른 쪽이 바뀐 것은 무관
    open(os.path.join(bd, '03_장.md'), 'w', encoding='utf8').write('[p.55 · PDF 69]\n\nalpha\n\n[p.56 · PDF 70]\n\nbeta finding revised.\n')
    assert CGm.mapstale(lambda s: '', cl, io.StringIO(), sources=root)['changed'] == ['q']  # 원 sha 가 없는 교과서는 글이 바뀌면 [변경]
    # 절 표지(oa XML md): 상자 뒤 다시 붙은 같은 절 표지도 한 절로
    md = '# t\n\n> Europe PMC 전문 XML(PMC1)에서(변환 v0.8.1)\n\n[§ Methods]\n\nfirst part.\n\n[§ 상자 · K]\n\nbox.\n\n[§ Methods]\n\nsecond part.\n'
    assert 'first part' in CGm._at_text(md, '[§ Methods]') and 'second part' in CGm._at_text(md, '[§ Methods]') and 'box' not in CGm._at_text(md, '[§ Methods]')
    try:
        CGm._at_text(md, '[§ Results]'); assert False
    except KeyError:
        pass
    assert CGm._page_of('[p.e12 · PDF 2]') == ('e12', 2) and CGm._page_of('[p.— · PDF 3]') == (None, 3) and CGm._page_of('[p.7]') == (None, 7) and CGm._page_of('p.7') == ('7', None)


def t_v16_mapdraw():
    g = copy.deepcopy(GRAPH)
    g[0]['role'] = 'evidence'; g[1]['role'] = 'main'; g[1]['statement'] = 'He said "yes" <b>'
    g.append({'id': 'r', 'statement': 'R', 'role': 'rebuttal', 'depends_on': []})
    g.append({'id': 'z', 'statement': 'Z', 'depends_on': []})
    g[1]['depends_on'].append({'id': 'r', 'type': 'rebuttal'})
    md = CGm.mapdraw(g, text=True, all_edges=True)                                          # v16.0 모양은 --all-edges 로 남는다
    assert md.count('```mermaid') == 1 and 'flowchart BT' in md, md
    assert 'n1 ==> n2' in md and 'n2 -.-> n3' in md and 'n4 -- 반박 --x n2' in md, md       # 근거 → 주장 방향, 종류별 선
    assert '#quot;yes#quot;' in md and '<b>' not in md and '#lt;b#gt;' in md, md             # 따옴표·꺾쇠는 Mermaid 가 깨지지 않게
    assert 'n1["a<br/>evidence·mid' in md and '#lt;br' not in md, md                          # 줄바꿈 <br/> 은 그대로(v16 첫 빌드 결함)
    assert 'n5' in md                                                                       # --all-edges 는 외톨이도 그린다
    md = CGm.mapdraw(g, changed=['a'])
    assert 'n5' not in md and 'n4' not in md and 'class n1 changed' in md and 'class n2 must' in md, md   # impact: 경로만, 색
    md = CGm.mapdraw(g, changed=['nope'])                                                   # 실패 길: 없는 id — 멈추지 않고 빈 그림
    assert 'flowchart BT' in md and 'n1' not in md, md


def t_v16_mapdraw_and_saved_pairs_cli():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgcli_')
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    A = os.path.join(d, 'a.json'); B = os.path.join(d, 'b.json')
    json.dump({'doc': 'x.docx', 'claims': [{'id': 'a1', 'statement': 's', 'sites': ['doc:p:1'], 'keys': ['k1'], 'depends_on': []}]}, open(A, 'w'))
    json.dump({'claims': [{'id': 'b9', 'statement': 's', 'sites': ['doc:p:9'], 'keys': ['k9'], 'depends_on': []}]}, open(B, 'w'))
    run = lambda *a: subprocess.run([sys.executable, cg] + list(a), capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    r = run('mapdiff', A, B, '--labels', '저자', '리뷰어')
    assert '양쪽이 잡은 주장 0개' in r.stdout, r.stdout                                      # 자리·keys 가 달라 자동으로는 못 짝짓는다
    r = run('mapdiff', A, B, '--labels', '저자', '리뷰어', '--save-pairs')
    assert r.returncode == 2 and '--pairs 와 함께' in r.stdout, r.stdout                     # 실패 길: 짝 없이 저장하지 않는다
    r = run('mapdiff', A, B, '--labels', '저자', '리뷰어', '--pairs', 'a1=b9', '--save-pairs')
    assert json.load(open(A))['pairs_with'] == {'리뷰어': {'a1': 'b9'}} and json.load(open(A))['doc'] == 'x.docx', open(A).read()
    r = run('mapdiff', A, B, '--labels', '저자', '리뷰어')
    assert '양쪽이 잡은 주장 1개' in r.stdout and 'pairs_with' in r.stdout, r.stdout           # 다음부터 --pairs 없이
    r = run('mapdiff', B, A, '--labels', '리뷰어', '저자')
    assert '양쪽이 잡은 주장 1개' in r.stdout, r.stdout                                      # 반대쪽에서 불러도 뒤집어 쓴다
    out = os.path.join(d, '관계도.md')
    r = run('mapdraw', '--claims', A, '-o', out)
    t = open(out, encoding='utf8').read()
    assert r.returncode == 0 and 'flowchart LR' in t and '`a1`' in t, (r.stderr, t)          # 간선 없는 한 주장 — 그림 밖 목록으로
    r = run('mapdraw', '--claims', A, '-o', out, '--all-edges')
    assert r.returncode == 0 and 'flowchart BT' in open(out, encoding='utf8').read(), r.stderr
    r = run('mapdraw', '--claims', A, '-o', out, '--impact', 'a1', '--text')
    assert r.returncode == 0 and 'class n1 changed' in open(out, encoding='utf8').read()



def t_v161_mapdraw_compact_default():
    # 사용자 09-29: 저자 51주장·간선 83(caveat 46) 전체 그림이 6614×1033 px — 왼쪽→오른쪽, caveat 접기, 역할별 색, 외톨이는 목록
    g = [{'id': 'm', 'role': 'main', 'depends_on': [{'id': 'e', 'type': 'premise'}, {'id': 'cv', 'type': 'caveat'}]},
         {'id': 'e', 'role': 'evidence', 'depends_on': [{'id': 'cv', 'type': 'caveat'}, {'id': 'cv2', 'type': 'caveat'}, {'id': 'bg', 'type': 'context'}]},
         {'id': 'bg', 'role': 'background', 'depends_on': []},
         {'id': 'cv', 'role': 'caveat', 'depends_on': []},
         {'id': 'cv2', 'role': 'caveat', 'depends_on': []},
         {'id': 'cvk', 'role': 'caveat', 'depends_on': []},                                   # caveat 이지만 context 간선으로 끼어 있다 — 접지 않는다
         {'id': 'k', 'role': 'claim', 'depends_on': [{'id': 'cvk', 'type': 'context'}, {'id': 'm', 'type': 'support'}]},
         {'id': 'lone', 'role': 'method', 'statement': 'lonely', 'depends_on': []},
         {'id': 'onlycav', 'role': 'claim', 'depends_on': [{'id': 'cv', 'type': 'caveat'}]}]
    md = CGm.mapdraw(g)
    block = md.split('```mermaid')[1].split('```')[0]
    assert 'flowchart LR' in block and 'flowchart BT' not in block, md                       # 근거 왼쪽 → main 오른쪽
    assert '"e<br/>evidence·mid<br/>한계 2"' in block and '"m<br/>main·mid<br/>한계 1"' in block, block   # 걸린 caveat 수
    assert '"cv<br/>' not in block and '"cv2<br/>' not in block and '한계 .->' not in block, block       # caveat 상자·선은 접힘
    assert '"cvk<br/>caveat·mid"' in block, block                                             # 다른 간선에 낀 caveat 은 남긴다
    assert 'classDef r_main' in block and 'classDef r_evidence' in block and 'classDef r_claim' in block and 'classDef r_background' in block, block
    assert '"lone' not in block and '"onlycav' not in block, block                           # 외톨이는 그림에 없고
    after = md.split('```')[-1]
    assert '`lone` (방법)' in after and '`onlycav` (해석, 한계 1)' in after and '접은 caveat 2개' in after, after   # 그림 밖 목록에
    assert '역할·confidence' in md                                                            # 범례에 상자 글 뜻(발표 09-29)
    # --impact 는 v16.0 그대로(저자 확인 끝남): 아래→위, caveat 도 경로에 있으면 그린다
    mi = CGm.mapdraw(g, changed=['cv'])
    assert 'flowchart BT' in mi and 'class n4 changed' in mi and 'classDef r_' not in mi, mi
    # 실패 길: 간선이 하나도 없는 그래프(발표 09-20 판 모양) — 빈 그림·목록 대신 한 줄(v16.2)
    none = [{'id': 'x%02d' % k, 'depends_on': []} for k in range(21)]
    md = CGm.mapdraw(none)
    assert '```mermaid' not in md and '간선이 하나도 없는 그래프(주장 21개)' in md and 'x05' not in md, md
    assert 'flowchart BT' in CGm.mapdraw(none, all_edges=True) and '"x05' in CGm.mapdraw(none, all_edges=True)   # 상자는 --all-edges 로
    one = CGm.mapdraw([{'id': 'solo', 'depends_on': []}])                                     # 주장 하나는 그래프 모양 그대로(목록 한 줄)
    assert '`solo`' in one, one


def t_v161_mapgraph_isolated_note():
    g = copy.deepcopy(GRAPH) + [{'id': 'z1', 'depends_on': []}, {'id': 'z2', 'depends_on': []}]
    ps = [p for p in CGm.mapgraph(g, io.StringIO())[0] if '외톨이' in p]
    assert ps == ['[참고] 간선이 하나도 없는 주장(외톨이) 2개: z1, z2'], ps                    # 한 줄로, 이름과 함께
    assert not [p for p in CGm.mapgraph(copy.deepcopy(GRAPH), io.StringIO())[0] if '외톨이' in p]   # 다 이어졌으면 없음
    assert not [p for p in CGm.mapgraph([{'id': 'solo'}], io.StringIO())[0] if '외톨이' in p]        # 주장 하나뿐이면 알리지 않음
    ps = [p for p in CGm.mapgraph([{'id': 'x%02d' % k} for k in range(21)], io.StringIO())[0] if '간선이 하나도' in p]
    assert ps == ['[참고] 간선이 하나도 없는 그래프(주장 21개) — 관계(depends_on)를 아직 적지 않았다'], ps   # v16.2: 21개 이름을 늘어놓지 않는다
    # v16.3 (사용자 09-29, 부관리자 재현): role=claim 21개·간선 0 — "premise 간선이 없음" 이 21줄 나오던 것도 요약 한 줄로
    g21 = [{'id': 'x%02d' % k, 'role': 'claim'} for k in range(20)] + [{'id': 'ev', 'role': 'evidence'}]
    ps = [p for p in CGm.mapgraph(g21, io.StringIO())[0] if p.startswith('[참고]')]
    assert '[참고] 간선이 하나도 없는 그래프(주장 21개) — 관계(depends_on)를 아직 적지 않았다 (주장별 간선 [참고] premise 없음 20 · caveat 없음 1 를 이 줄로 합침)' in ps, ps
    assert len(ps) == 2 and any('role=main 인 주장이 0개' in p for p in ps), ps               # 그래프 전체 한 줄(main 개수)은 그대로
    # 실패 길: 간선이 하나라도 있으면 주장별 [참고] 는 전처럼 줄마다
    g21[0]['depends_on'] = [{'id': 'ev', 'type': 'premise'}]
    ps = CGm.mapgraph(g21, io.StringIO())[0]
    assert sum('premise 간선이 없음' in p for p in ps) == 19 and any('ev: evidence 인데 걸린 caveat 이 없음' in p for p in ps), ps



def t_v164_case_example_in_doc_is_valid():
    # v16.4·v16.6: CLAIM_GRAPH.md §3-2-1 증례 예시가 규약대로 돈다 — 문서가 도구와 어긋나지 않게
    import json, re as _re
    doc = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'CLAIM_GRAPH.md')
    if not os.path.exists(doc):
        return
    t = open(doc, encoding='utf8').read()
    g = json.loads(_re.search(r'## 3-2-1.*?```json\n(.*?)```', t, _re.S).group(1))
    cl = g['claims']
    assert CGm.mapgraph(copy.deepcopy(cl), io.StringIO(), kind=g['kind'])[0] == []               # kind 증례 — 알림 없음
    assert [r[0] for r in CGm.impact(cl, ['f-calc'], io.StringIO())] == ['ddx-gran', 'dx', 'ddx-malig']
    md = CGm.mapdraw(cl)
    assert '-- 반박 --x' in md and 'class ' in md and 'excluded' in md and 'subgraph g1["증례1"]' in md, md



GAPG = [
    {'id': 'm', 'role': 'main', 'sources': [{'kind': '문헌', 'what': '10.1000/a', 'at': 'p.1'}, {'kind': '문헌', 'what': '10.1000/b', 'at': 'p.2'}],
     'depends_on': [{'id': 'e', 'type': 'premise'}, {'id': 'lowc', 'type': 'support'}, {'id': 'k', 'type': 'premise'}]},
    {'id': 'k', 'role': 'claim', 'keys': ['alpha term', 'beta term'], 'depends_on': [{'id': 'e', 'type': 'premise'}, {'id': 'b', 'type': 'support'}]},
    {'id': 'k2', 'role': 'claim', 'sources': [{'kind': '문헌', 'what': '10.1000/c'}, {'kind': '교과서', 'what': '책', 'at': 'p.5'}],
     'depends_on': [{'id': 'e', 'type': 'support'}]},
    {'id': 'b', 'role': 'background', 'sources': [{'kind': '문헌', 'what': '10.1000/d', 'at': 'p.3'}], 'depends_on': []},
    {'id': 'e', 'role': 'evidence', 'depends_on': []},
    {'id': 'lowc', 'role': 'claim', 'confidence': 'low', 'sources': [{'kind': '문헌', 'what': '10.1000/e'}, {'kind': '문헌', 'what': '10.1000/f'}],
     'depends_on': [{'id': 'e', 'type': 'premise'}, {'id': 'b', 'type': 'premise'}]},
    {'id': 'x', 'role': 'method', 'depends_on': []},
    {'id': 'old', 'role': 'claim', 'status': 'superseded', 'depends_on': []},
]


def t_v165_find_gaps():
    # 사용자 09-29 (가): 문헌 없음(claim·main·background 만) · 근거 하나(sources 1 또는 받침 간선 1) · 약한 고리 · 외톨이
    rows, noedge = CGm.find_gaps(copy.deepcopy(GAPG))
    got = {cid: ks for cid, _, ks, _ in rows}
    assert not noedge and set(got) == {'k', 'k2', 'b', 'lowc', 'x'}, got                    # m 은 sources 2·받침 3 — 공백 없음
    assert got['k'] == ['문헌 없음'], got                                                     # 받침 간선 2 라 '근거 하나' 는 아님
    assert got['k2'] == ['근거 하나 — 받침 간선 1(e)'], got                                    # 교과서도 원문 근거로 센다(sources 2)
    assert got['b'] == ['근거 하나 — sources 1'], got
    assert got['lowc'][0].startswith('약한 고리 — confidence low, 기대는 주장 1(m)'), got
    assert got['x'] == ['외톨이'], got
    assert 'e' not in got and 'old' not in got                                               # evidence 는 문헌 공백에서 빼고, 철회는 뺀다
    assert all(t_ == '' for _, _, _, t_ in rows), rows                                        # v16.7: 검색어는 keys 에서 가져오지 않는다(비움)
    rows, noedge = CGm.find_gaps([{'id': 'p', 'role': 'claim'}, {'id': 'q', 'role': 'claim'}])   # 간선 0 그래프: 외톨이를 줄마다 내지 않는다
    assert noedge and all('외톨이' not in ks for _, _, ks, _ in rows) and len(rows) == 2, rows


def t_v165_gaps_table_and_to_instr():
    md = CGm.gaps_table(copy.deepcopy(GAPG), '시험원고')
    rows = [l for l in md.splitlines() if l.startswith('| G')]
    assert len(rows) == 10 and rows[0].startswith('| G01-받침 | `k`') and rows[1].startswith('| G01-반박 | `k`'), rows[:2]   # 공백마다 받침·반박
    assert 'AI 제안' in md and '반박 줄도' in md and 'evidence' in md, md[:600]
    # 채운 표: G01 받침은 AI 제안 DOI, 반박은 비움 · G02 받침은 DOI 없이 글만 · G03 받침은 이미 판정 · G04 받침·반박이 같은 DOI
    fill = {'G01-받침': ('https://doi.org/10.1234/ABC.5', 'AI 제안', '', 'alpha term; beta term'), 'G02-받침': ('Kim 2020 어딘가', 'AI 제안', '', 'x'),
            'G03-받침': ('10.9/zz', '사람', '부합', 'x'), 'G04-받침': ('10.1234/abc.5', '사람', '', 'y'), 'G04-반박': ('10.5555/r1', 'AI 제안', '', 'z')}
    out = []
    for l in md.splitlines():
        c = [x.strip() for x in l.strip().strip('|').split('|')]
        if c and c[0] in fill:
            cand, src, ver, terms = fill[c[0]]; c[6], c[7], c[9], c[5] = cand, src, ver, terms
            l = '| ' + ' | '.join(c) + ' |'
        out.append(l)
    instr, notes = CGm.gaps_to_instr('\n'.join(out), '시험원고')
    assert '1. 후보 G01-받침 (AI 제안). doi:10.1234/abc.5' in instr and '2. 후보 G04-반박 (AI 제안). doi:10.5555/r1' in instr, instr
    assert '10.9/zz' not in instr and instr.count('doi:10.1234/abc.5') == 1, instr           # 판정 있는 줄은 빼고, 같은 DOI 는 한 번
    assert '| G04-받침 | 1 |' in instr and '| G04-반박 | 2 |' in instr, instr
    assert any('G01: 받침 줄만' in n for n in notes) and any('G02-받침: 후보 칸에 DOI 가 없다' in n for n in notes), notes   # 실패 길: 알림
    assert not any('G04' in n for n in notes), notes
    try:                                                                                        # literature 가 있는 세트면 그 파서로 읽힌다
        import literature as LT
    except ImportError:
        return
    refs, cl = LT.parse_refs(instr), LT.parse_claims(instr)
    assert [r['doi'] for r in refs] == ['10.1234/abc.5', '10.5555/r1'] and [c['refs'] for c in cl] == [[1], [1], [2]], (refs, cl)
    assert cl[0]['terms'] == ['alpha term', 'beta term'], cl


def t_v165_mapgraph_sources_store():
    import tempfile
    root = tempfile.mkdtemp(prefix='cggs_')
    d = os.path.join(root, '10.1000_a'); os.makedirs(d)
    open(os.path.join(d, 'meta.md'), 'w', encoding='utf8').write('<!-- lit: doi=10.1000/a pages=1 blank=0 sha=aa -->\n')
    g = [{'id': 'k', 'role': 'claim', 'depends_on': [], 'sources': [
        {'kind': '문헌', 'what': '10.1000/a', 'at': 'p.1', 'verdict': '부합'},
        {'kind': '문헌', 'what': '10.1000/a', 'at': 'p.2'},
        {'kind': '문헌', 'what': '10.1000/zzz', 'at': 'p.1', 'verdict': '부합'}]}]
    ps = CGm.mapgraph(copy.deepcopy(g), io.StringIO(), sources=root)[0]
    hard = [p for p in ps if not p.startswith('[참고]')]
    assert hard == ['k: sources[3]: 문헌 10.1000/zzz 가 보관소에 없다 — 원문을 받지 않은 근거는 sources 에 넣지 않는다(작업표에 둔다)'], ps   # [필수]
    assert any(p.startswith('[참고] k: sources[2]: 문헌 10.1000/a 판정') for p in ps) and not any('sources[1]' in p for p in ps), ps
    assert not [p for p in CGm.mapgraph(copy.deepcopy(g), io.StringIO())[0] if '보관소' in p or '판정' in p]   # 보관소를 안 주면 전처럼


def t_v1611_gaps_part2_title_counts_overlap():
    # 다음 할 일 ③(09-29 실물 v9): 머리 줄은 "인용 있음 6" 인데 2부 제목은 "2개" — 나머지 4개는 다른 공백과 겹쳐 1부로 갔다는 말이 없었다
    g = [{'id': 'c1', 'role': 'claim', 'statement': 'X is higher [12,14].', 'depends_on': []},
         {'id': 'c2', 'role': 'claim', 'statement': 'Y is linked [7].', 'depends_on': []},
         {'id': 'c3', 'role': 'claim', 'statement': 'Z plausibly follows', 'depends_on': []},
         {'id': 'c4', 'role': 'background', 'statement': 'W [3]', 'confidence': 'low', 'depends_on': []},
         {'id': 'c5', 'role': 'claim', 'statement': 'V [5]', 'confidence': 'low', 'depends_on': []},
         {'id': 'd', 'role': 'main', 'sources': [{'kind': '문헌', 'what': '10.1/a'}, {'kind': '문헌', 'what': '10.1/b'}],
          'depends_on': [{'id': x, 'type': 'premise'} for x in ('c1', 'c2', 'c4', 'c5')] + [{'id': 'c3', 'type': 'context'}]}]
    md = CGm.gaps_table(copy.deepcopy(g), 'T')
    assert '문헌 없음(인용 있음 — sources 미기입) 4' in md.splitlines()[2], md.splitlines()[2]
    p1, p2 = md.split('## 2. 인용 있음')
    assert p2.startswith(' — sources 미기입 2개 (+ 다른 공백과 겹쳐 1부에 간 2개)\n'), p2[:120]            # 2 + 2 = 머리 줄 4
    note = [l for l in p2.splitlines() if l.startswith('> 인용 있음 4개 중 2개')]
    assert len(note) == 1 and 'G02 `c4`' in note[0] and 'G03 `c5`' in note[0] and '받침 줄' in note[0], p2
    assert '| G02-받침 | `c4`' in p1 and '| G03-받침 | `c5`' in p1, p1
    # 겹친 것이 없으면 제목·알림은 전과 같다
    g2 = [c for c in copy.deepcopy(g) if c['id'] not in ('c4', 'c5')]
    g2[-1]['depends_on'] = [e for e in g2[-1]['depends_on'] if e['id'] not in ('c4', 'c5')]
    md2 = CGm.gaps_table(g2, 'T')
    assert '## 2. 인용 있음 — sources 미기입 2개\n' in md2 and '겹쳐' not in md2, md2
    # 작업표 → 검증지시 는 제목·알림 줄과 상관없이 표 줄만 읽는다
    row = next(l for l in md.splitlines() if l.startswith('| G02-받침'))
    c = [x.strip() for x in row.strip().strip('|').split('|')]
    c[5], c[6], c[7] = 'W', '10.1000/zz', '사람'                                                 # 검색어 · 후보 DOI · 출처
    ins, notes = CGm.gaps_to_instr(md.replace(row, '| ' + ' | '.join(c) + ' |'), 'T')
    assert 'doi:10.1000/zz' in ins and '| G02-받침 | 1 | c4 받침(사람) | W |' in ins, (ins, notes)


# ---------------------------------------------------------------- 구연 덧붙임 (v16.12, 사용자 09-30)
ORAL_AUTHOR = {'doc': 'Fake_manuscript_v3', 'claims': [
    {'id': 'ev-a', 'role': 'evidence', 'statement': 'Alpha sentence from the manuscript', 'evidence': 'n=40',
     'sites': ['doc:find:Alpha sentence'], 'keys': ['alpha sentence'], 'depends_on': [], 'verified': {'at': '2026-01-01', 'sites': {}}},
    {'id': 'ev-b', 'role': 'evidence', 'statement': 'Beta sentence from the manuscript', 'sites': ['doc:find:Beta sentence'],
     'keys': ['beta sentence'], 'depends_on': []},
    {'id': 'cv-1', 'role': 'caveat', 'statement': 'Gamma limitation sentence', 'sites': ['doc:find:Gamma'], 'keys': ['gamma'], 'depends_on': []},
    {'id': 'cl-1', 'role': 'claim', 'statement': 'Delta claim sentence', 'sites': ['doc:find:Delta'], 'keys': ['delta claim'],
     'depends_on': [{'id': 'ev-a', 'type': 'premise'}, {'id': 'cv-1', 'type': 'caveat'}]},
    {'id': 'mn', 'role': 'main', 'statement': 'Epsilon main sentence', 'sites': ['doc:find:Epsilon'], 'keys': ['epsilon'],
     'forbidden': ['old epsilon wording'], 'depends_on': [{'id': 'cl-1', 'type': 'premise'}, {'id': 'ev-b', 'type': 'support'}]},
    {'id': 'old', 'role': 'claim', 'status': 'superseded', 'statement': 'Zeta retracted', 'sites': [], 'depends_on': []},
    {'id': 'down', 'role': 'claim', 'statement': 'Eta downstream sentence', 'sites': ['doc:find:Eta'], 'keys': ['eta'],
     'depends_on': [{'id': 'mn', 'type': 'premise'}]}]}


def _oral_files(d, overlay_edit=None):
    import json
    A = os.path.join(d, 'author.json'); json.dump(ORAL_AUTHOR, open(A, 'w'), ensure_ascii=False)
    O = os.path.join(d, 'oral.json')
    ov = CGm.oral_init(A, deck='Fake_deck')
    if overlay_edit:
        overlay_edit(ov)
    json.dump(ov, open(O, 'w'), ensure_ascii=False)
    return A, O


def t_v1612_oral_snapshot_is_hashes_only():
    # 사용자 09-30: 덧붙임 스냅숏은 원고 문장을 옮기지 않는다(파일 크기·원고 퍼짐) — 짝짓기에 쓰는 원고 자리·keys 도 해시로
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgor_')
    A, O = _oral_files(d)
    ov = json.load(open(O))
    assert ov['kind'] == '구연' and ov['deck'] == 'Fake_deck' and ov['use'] == {} and ov['claims'] == [], ov
    src = ov['source']
    assert src['file'] == 'author.json' and src['doc'] == 'Fake_manuscript_v3' and src['n'] == 7 and len(src['sha']) == 16, src
    assert set(src['snap']) == {c['id'] for c in ORAL_AUTHOR['claims']} and src['snap']['cl-1']['role'] == 'claim'
    blob = json.dumps(ov, ensure_ascii=False).lower()
    for c in ORAL_AUTHOR['claims']:                                         # 문장·원고 자리·keys 어느 것도 글자로 남지 않는다
        for t in [c['statement']] + c.get('sites', []) + c.get('keys', []) + c.get('forbidden', []):
            assert t.lower() not in blob, t
    # 해시로도 mapdiff 짝 점수가 원문과 같다(정확히 같은 것만 세므로)
    snap = src['snap']
    assert snap['ev-a']['sites'] == [CGm._oral_h(CGm._norm_site('doc:find:Alpha sentence'))] and snap['ev-a']['keys'] == [CGm._oral_h('alpha sentence')]
    assert snap['ev-a']['text'] == CGm._fingerprint('n=40|Alpha sentence from the manuscript')
    try:                                                                     # 실패 길: 이미 있는 덧붙임은 덮어쓰지 않는다
        CGm.oral_init(A, deck='x', out=O); assert False
    except SystemExit as e:
        assert '이미 있다' in str(e), e


def t_v1612_oral_merge_success():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgor_')
    def edit(ov):
        ov['use'] = {'mn': {'sites': ['slide@260'], 'keys': ['엡실론']}, 'ev-b': {'sites': ['notes@260']}}
        ov['claims'] = [{'id': 'p-intro', 'role': 'claim', 'statement': '도입', 'sites': ['slide@256'], 'keys': ['도입'],
                         'depends_on': [{'id': 'ev-b', 'type': 'support'}]}]
    A, O = _oral_files(d, edit)
    before = open(A, 'rb').read()
    meta, cl, probs = CGm.load_oral(O, A)
    by = {c['id']: c for c in cl}
    assert meta.get('deck') == 'Fake_deck' and 'doc' not in meta, meta
    assert set(by) == {'mn', 'ev-b', 'cl-1', 'ev-a', 'cv-1', 'p-intro'}, sorted(by)       # 쓴 것 + 상류 전부, 하류 down·철회 old 는 빠짐
    assert by['mn']['sites'] == ['slide@260'] and by['mn']['keys'] == ['엡실론'] and by['mn']['forbidden'] == ['old epsilon wording']
    assert by['ev-b']['sites'] == ['notes@260'] and by['ev-b']['keys'] == []
    for off in ('cl-1', 'ev-a', 'cv-1'):                                     # 무대 밖: 자리·keys·원고 verified 없음
        assert by[off]['offstage'] is True and by[off]['sites'] == [] and by[off]['keys'] == [] and 'verified' not in by[off], by[off]
    assert 'verified' not in by['mn'] and 'offstage' not in by['mn']
    assert [p for p in probs if not p.startswith('[참고]')] == [], probs
    kn = [p for p in probs if 'keys 없는 주장' in p]
    assert kn == ['[참고] 화면 keys 없는 주장 1개(ev-b) — 그 주장은 keys 검사를 하지 않는다(화면 표현을 use.keys 에)'], probs
    ps, _ = CGm.mapgraph(cl, stream=io.StringIO())
    assert not [p for p in ps if not p.startswith('[참고]')], ps
    md = CGm.mapdraw(cl)
    assert 'classDef offstage' in md and 'p-intro' in md.replace('p_intro', 'p-intro'), md[-600:]
    assert open(A, 'rb').read() == before                                    # 저자 파일은 읽기만


def t_v1612_oral_merge_failures():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgor_')
    cases = [
        (lambda ov: ov['use'].update({'gone': {'sites': ['slide@1']}}), '저자 파일에 없다'),
        (lambda ov: ov['use'].update({'old': {'sites': ['slide@1']}}), '철회한 주장'),
        (lambda ov: ov['use'].update({'mn': {'sites': ['doc:find:Epsilon']}}), '화면 자리'),
        (lambda ov: ov['claims'].append({'id': 'intro', 'statement': 'x', 'depends_on': []}), 'p- 로 시작'),
        (lambda ov: ov['claims'].extend([{'id': 'p-a', 'statement': 'x', 'depends_on': []},
                                          {'id': 'p-b', 'statement': 'y', 'depends_on': [{'id': 'p-a'}]}]), '저자 id 에만'),
        (lambda ov: ov['source'].update({'sha': '0' * 16}), '판이 다르다'),
    ]
    for edit, word in cases:
        A, O = _oral_files(d, edit)
        _, _, probs = CGm.load_oral(O, A)
        hard = [p for p in probs if not p.startswith('[참고]')]
        assert hard and any(word in p for p in hard), (word, probs)
    # CLI: oral check 종료 코드 1 / 성공 0, oral init 덮어쓰기 거부
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    run = lambda *a: subprocess.run([sys.executable, cg] + list(a), capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    A, O = _oral_files(d, cases[0][0])
    r = run('oral', 'check', '--author', A, '--oral', O)
    assert r.returncode == 1 and '저자 파일에 없다' in r.stdout, (r.stdout, r.stderr)
    A, O = _oral_files(d, lambda ov: ov['use'].update({'mn': {'sites': ['slide@260'], 'keys': ['e']}}))
    r = run('oral', 'check', '--author', A, '--oral', O)
    assert r.returncode == 0 and '무대 밖 상류 4개' in r.stdout, (r.stdout, r.stderr)
    r = run('mapdraw', '--oral', O, '--author', A, '-o', os.path.join(d, 'g.md'))
    assert r.returncode == 0 and 'offstage' in open(os.path.join(d, 'g.md'), encoding='utf8').read(), (r.stdout, r.stderr)
    r = run('impact', '--oral', O, '--author', A, 'ev-a')
    assert r.returncode == 0 and 'mn' in r.stdout, (r.stdout, r.stderr)
    r = run('mapgraph', '--oral', O)                                         # 둘 중 하나만 주면 고칠 말
    assert r.returncode == 2 and '--author' in r.stdout, (r.stdout, r.stderr)
    r = run('oral', 'init', '--author', A, '-o', O)
    assert r.returncode != 0 and '이미 있다' in (r.stdout + r.stderr), (r.stdout, r.stderr)


def _author_v2(**kw):
    """가짜 저자 새 판: mn 글 바뀜 · ev-b → ev-b2 이름 바뀜(원고 자리 같음) · down 지움 · new-1 새 주장."""
    import copy as _c
    a = _c.deepcopy(ORAL_AUTHOR)
    by = {c['id']: c for c in a['claims']}
    by['mn']['statement'] = 'Epsilon main sentence, revised'
    by['ev-b']['id'] = 'ev-b2'
    for c in a['claims']:
        for e in c.get('depends_on', []):
            if isinstance(e, dict) and e['id'] == 'ev-b':
                e['id'] = 'ev-b2'
    a['claims'] = [c for c in a['claims'] if c['id'] != 'down']
    a['claims'].append({'id': 'new-1', 'role': 'claim', 'statement': 'Theta new', 'sites': ['doc:find:Theta'], 'keys': ['theta'], 'depends_on': []})
    for k, v in kw.items():
        v(a)
    return a


def _oral_v2_files(d, overlay_edit=None, **kw):
    import json
    A, O = _oral_files(d, overlay_edit)
    A2 = os.path.join(d, 'author_v2.json'); json.dump(_author_v2(**kw), open(A2, 'w'), ensure_ascii=False)
    return A, O, A2


def _use_mn_evb(ov):
    ov['use'] = {'mn': {'sites': ['slide@260'], 'keys': ['e']}, 'ev-b': {'sites': ['notes@260'], 'keys': ['b'], 'verified': {'at': 'x'}}}
    ov['claims'] = [{'id': 'p-intro', 'role': 'claim', 'statement': '도입', 'sites': ['slide@256'], 'keys': ['도입'],
                     'depends_on': [{'id': 'ev-b', 'type': 'support'}]}]


def t_v1613_oral_sync_stops_on_vanished_id_with_candidate():
    # v2.71 (사용자 09-30 결정 4): 쓴 id 가 새 판에 없으면 --pairs·--drop 전까지 멈춘다 — 스냅숏 해시로 후보를 찾는다
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgsy_')
    A, O, A2 = _oral_v2_files(d, _use_mn_evb)
    ov = json.load(open(O)); am, ac = CGm.load_claims_full(A2)
    new, lines, hard = CGm.oral_sync(ov, am, ac, CGm._file_sha(A2), 'author_v2.json')
    assert new is None and len(hard) == 1 and 'ev-b' in hard[0] and 'ev-b2' in hard[0] and '--pairs ev-b=' in hard[0] and 'notes@260' in hard[0], hard
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    run = lambda *a: subprocess.run([sys.executable, cg] + list(a), capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    O2 = os.path.join(d, 'oral_v2.json')
    r = run('oral', 'sync', '--author', A2, '--oral', O, '-o', O2)
    assert r.returncode == 1 and '[필수]' in r.stdout and not os.path.exists(O2), (r.stdout, r.stderr)
    # 실패 길 둘: 없는 새 id 로 짝 · 철회된 주장이 화면에
    new, _, hard = CGm.oral_sync(ov, am, ac, 'x' * 16, 'v2', pairs={'ev-b': 'nope'})
    assert new is None and any('nope' in h for h in hard), hard
    am3, ac3 = am, [dict(c, status='superseded') if c['id'] == 'mn' else c for c in ac]
    new, _, hard = CGm.oral_sync(ov, am3, ac3, 'x' * 16, 'v2', pairs={'ev-b': 'ev-b2'})
    assert new is None and any('철회' in h and 'mn' in h for h in hard), hard


def t_v1613_oral_sync_success():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgsy_')
    A, O, A2 = _oral_v2_files(d, _use_mn_evb)
    before = (open(A, 'rb').read(), open(A2, 'rb').read())
    ov = json.load(open(O)); am, ac = CGm.load_claims_full(A2)
    new, lines, hard = CGm.oral_sync(ov, am, ac, CGm._file_sha(A2), 'author_v2.json', pairs={'ev-b': 'ev-b2'})
    assert not hard and new, (lines, hard)
    txt = '\n'.join(lines)
    assert '[옮김] ev-b → ev-b2' in txt and '[변경] 쓴 주장 mn' in txt and '[참고] 새 저자 주장 1개(new-1)' in txt and '[참고] 저자가 뺀 주장 1개(down)' in txt, txt
    assert set(new['use']) == {'mn', 'ev-b2'} and new['use']['ev-b2'] == {'sites': ['notes@260'], 'keys': ['b']}, new['use']   # 옮긴 주장은 확인 기록 없이
    assert new['claims'][0]['depends_on'] == [{'id': 'ev-b2', 'type': 'support'}], new['claims']            # 발표 주장의 기댐도 따라간다
    assert new['source']['sha'] == CGm._file_sha(A2) and new['source']['file'] == 'author_v2.json' and 'ev-b2' in new['source']['snap']
    assert new['synced'][-1]['from'] == ov['source']['sha'] and new['synced'][-1]['moved'] == {'ev-b': ['ev-b2']}, new['synced']
    json.dump(new, open(O, 'w'), ensure_ascii=False)
    _, cl, probs = CGm.load_oral(O, A2)
    assert not [p for p in probs if not p.startswith('[참고]')], probs                               # 새 판과 맞는다
    # --drop · 1:N 쪼갬 · 같은 판이면 할 일 없음
    ov = json.load(open(_oral_files(d, _use_mn_evb)[1]))
    new, lines, hard = CGm.oral_sync(ov, am, ac, 'y' * 16, 'v2', drop=['ev-b'])
    assert not hard and 'ev-b' not in new['use'] and new['claims'][0]['depends_on'] == [] and any('[뺌] ev-b' in l for l in lines), lines
    ac4 = ac + [{'id': 'ev-b3', 'role': 'evidence', 'statement': 'split', 'sites': [], 'depends_on': []}]
    new, lines, hard = CGm.oral_sync(ov, am, ac4, 'z' * 16, 'v2', pairs={'ev-b': ['ev-b2', 'ev-b3']})
    assert not hard and set(new['use']) >= {'ev-b2', 'ev-b3'} and len(new['claims'][0]['depends_on']) == 2, (new['use'], new['claims'])
    new, lines, hard = CGm.oral_sync(ov, *CGm.load_claims_full(A), CGm._file_sha(A), 'author.json')
    assert new is None and not hard and any('같은 판' in l for l in lines), lines
    assert (open(A, 'rb').read(), open(A2, 'rb').read()) == before                                  # 저자 파일은 읽기만


# ---------------------------------------------------------------- 초점 그림 (v16.15, 사용자 09-30)
FOCUS_G = [
    {'id': 'e3', 'role': 'evidence', 'statement': 'Third-level evidence far upstream', 'depends_on': []},
    {'id': 'e2', 'role': 'evidence', 'statement': 'Second-level evidence', 'depends_on': [{'id': 'e3', 'type': 'support'}, {'id': 'cv2', 'type': 'caveat'}]},
    {'id': 'e1', 'role': 'evidence', 'statement': 'First-level evidence with a fairly long statement that is cut', 'depends_on': [{'id': 'e2', 'type': 'premise'}, {'id': 'cv1', 'type': 'caveat'}]},
    {'id': 'cv0', 'role': 'caveat', 'statement': 'Limit on the focus itself', 'depends_on': []},
    {'id': 'cv1', 'role': 'caveat', 'statement': 'Limit on first-level evidence', 'depends_on': []},
    {'id': 'cv2', 'role': 'caveat', 'statement': 'Limit on second-level evidence', 'depends_on': []},
    {'id': 'rb', 'role': 'rebuttal', 'statement': 'Counter finding', 'depends_on': []},
    {'id': 'f', 'role': 'claim', 'statement': 'The focused claim sentence that is long enough to be wrapped over more than one line in the box',
     'depends_on': [{'id': 'e1', 'type': 'premise'}, {'id': 'cv0', 'type': 'caveat'}, {'id': 'rb', 'type': 'rebuttal'}]},
    {'id': 'd1', 'role': 'main', 'statement': 'Downstream conclusion', 'depends_on': [{'id': 'f', 'type': 'premise'}]},
    {'id': 'far', 'role': 'claim', 'statement': 'Unrelated', 'depends_on': []}]


def t_v1615_focus_graph_up_down_and_labels():
    import copy as _c
    fg = CGm.focus_graph(_c.deepcopy(FOCUS_G), ['f'])
    k = fg['kind']
    assert set(k) == {'f', 'e1', 'e2', 'cv0', 'cv1', 'rb', 'd1'}, k                # 받침 2단계 · 한계 직접+1단계 · 반박 · 영향
    assert k['f'] == 'focus' and k['e1'] == k['e2'] == 'base' and k['cv0'] == k['cv1'] == 'limit' and k['rb'] == 'rebut' and k['d1'] == 'impact'
    assert 'e3' not in k and 'cv2' not in k and 'far' not in k                      # 3단계·2단계의 한계·무관은 빠짐
    md = CGm.focus_mermaid(fg)
    assert 'subgraph' in md and '받침' in md and '영향' in md and 'classDef focus' in md, md
    full = FOCUS_G[7]['statement']
    lab_f = re.search(r'\bn\d+\["([^"]*)"\]', [l for l in md.splitlines() if 'The focused' in l][0]).group(1)
    assert lab_f.replace('<br/>', ' ') == full and '<br/>' in lab_f, lab_f         # 선택한 주장은 전문(줄바꿈)
    lab_e1 = [l for l in md.splitlines() if 'First-level' in l][0]
    assert 'First-level evidence with a fairly long…' in lab_e1 and 'is cut' not in lab_e1, lab_e1   # 나머지는 앞 40자
    ids = CGm.focus_mermaid(fg, ids=True)
    assert '"f"' in ids and 'The focused' not in ids, ids
    leg = CGm.focus_legend(fg)
    assert [x for x, _ in leg] == ['선택한 주장', '핵심 근거', '한계', '반박', '영향받는 결론'], leg
    fg2 = CGm.focus_graph(_c.deepcopy([c for c in FOCUS_G if c['id'] != 'rb']), ['f'])
    assert '반박' not in [x for x, _ in CGm.focus_legend(fg2)]                        # 반박 선이 없으면 범례에서도 뺀다
    try:
        CGm.focus_graph(_c.deepcopy(FOCUS_G), ['nope']); assert False
    except ValueError as e:
        assert 'nope' in str(e)


def t_v1615_focus_oral_screens_and_offstage():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgfo_')
    A = os.path.join(d, 'a.json'); json.dump({'doc': 'x', 'claims': FOCUS_G}, open(A, 'w'), ensure_ascii=False)
    ov = CGm.oral_init(A, deck='d'); ov['use'] = {'f': {'sites': ['slide@260'], 'keys': ['k']}, 'd1': {'sites': ['slide@261'], 'keys': ['k']}}
    O = os.path.join(d, 'o.json'); json.dump(ov, open(O, 'w'), ensure_ascii=False)
    _, cl, _ = CGm.load_oral(O, A)
    fg = CGm.focus_graph(cl, ['f'])
    md = CGm.focus_mermaid(fg, screen=lambda s: {'slide@260': '화면 12', 'slide@261': '화면 13'}.get(s, s))
    assert '화면 12' in md and '화면 13' in md and '_off fill' in md and '화면에 없음' in md, md     # v16.16: 무대 밖은 종류 색을 흐리게
    assert '화면에 없음' in [x for x, _ in CGm.focus_legend(fg)]
    md2 = CGm.focus_mermaid(fg)                                                         # 덱을 안 주면 slide@ID
    assert 'slide@260' in md2


def t_v1615_focus_mermaid_js_order_and_no_browser():
    import tempfile
    d = tempfile.mkdtemp(prefix='cgmj_')
    given = os.path.join(d, 'given.js'); envf = os.path.join(d, 'env.js')
    nm = os.path.join(d, 'node_modules', 'mermaid', 'dist'); os.makedirs(nm); npmf = os.path.join(nm, 'mermaid.min.js')
    for p in (given, envf, npmf):
        open(p, 'w').write('/* fake */')
    old = os.environ.pop('CLAIM_GRAPH_MERMAID_JS', None)
    try:
        assert CGm.find_mermaid_js(given, cwd=d) == (given, '--mermaid-js')
        os.environ['CLAIM_GRAPH_MERMAID_JS'] = envf
        assert CGm.find_mermaid_js(None, cwd=d) == (envf, '환경변수')
        del os.environ['CLAIM_GRAPH_MERMAID_JS']
        assert CGm.find_mermaid_js(None, cwd=d) == (npmf, 'npm node_modules')
        src, how = CGm.find_mermaid_js(None, cwd=tempfile.mkdtemp())
        assert how == 'CDN' and src.startswith('https://') and '@11' in src, (src, how)
        try:
            CGm.find_mermaid_js(os.path.join(d, 'missing.js'), cwd=d); assert False
        except ValueError as e:
            assert 'missing.js' in str(e)
    finally:
        os.environ.pop('CLAIM_GRAPH_MERMAID_JS', None)
        if old is not None:
            os.environ['CLAIM_GRAPH_MERMAID_JS'] = old
    html = CGm.focus_html(CGm.focus_graph([dict(c) for c in FOCUS_G], ['f']), npmf)
    assert 'file://' + npmf in html and 'flowchart' in html and '선택한 주장' in html and '핵심 근거' in html, html[:400]
    # 브라우저가 없으면 md·html 을 남기고 알린다(조용히 넘어가지 않음) — 종료 코드 1
    import json
    A = os.path.join(d, 'a.json'); json.dump({'claims': FOCUS_G}, open(A, 'w'), ensure_ascii=False)
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', CLAIM_GRAPH_BROWSER='none', CLAIM_GRAPH_MERMAID_JS=npmf)
    r = subprocess.run([sys.executable, cg, 'focus', 'f', '--claims', A, '-o', os.path.join(d, 'f.md'), '--png', os.path.join(d, 'f.png')],
                       capture_output=True, text=True, env=env)
    assert r.returncode == 1 and '브라우저' in r.stdout and os.path.exists(os.path.join(d, 'f.md')) and os.path.exists(os.path.join(d, 'f.html')) \
        and not os.path.exists(os.path.join(d, 'f.png')), (r.stdout, r.stderr)
    r = subprocess.run([sys.executable, cg, 'focus', 'f', '--claims', A, '-o', os.path.join(d, 'g.md')], capture_output=True, text=True, env=env)
    assert r.returncode == 0 and '```mermaid' in open(os.path.join(d, 'g.md'), encoding='utf8').read(), (r.stdout, r.stderr)
    assert CGm.find_browser(env={'CLAIM_GRAPH_BROWSER': 'none'}) is None
    assert CGm.find_browser(env={'CLAIM_GRAPH_BROWSER': npmf}) == npmf                 # 준 경로를 그대로


def _fake_browser(d, title, svg):
    """가짜 브라우저(v2.74 시험): --dump-dom 이면 그 제목·svg 로 DOM, --screenshot= 이면 작은 PNG. 실제 Chrome 처럼 끝나지 않고 기다린다."""
    p = os.path.join(d, 'fakebrowser_%s' % re.sub(r'\W', '', title)[:12])
    open(p, 'w').write('''#!%s
import sys, time
a = sys.argv[1:]
if '--dump-dom' in a:
    sys.stdout.write('<html><head><title>%s</title></head><body><pre class="mermaid">%s</pre></body></html>\\n'); sys.stdout.flush()
for x in a:
    if x.startswith('--screenshot='):
        from PIL import Image, ImageDraw
        im = Image.new('RGB', (400, 300), 'white'); ImageDraw.Draw(im).rectangle((50, 50, 200, 150), fill='black'); im.save(x.split('=', 1)[1])
time.sleep(30)
''' % (sys.executable, title, '<svg></svg>' if svg else 'flowchart LR'))
    os.chmod(p, 0o755)
    return p


def t_v1616_focus_mermaid_failure_no_png():
    # 부관리자 09-30 [결함]: jsdelivr 403 으로 mermaid 가 안 돌았는데 mermaid 글자가 찍힌 PNG 를 저장하고 종료 코드 0 — 찍기 전에 SVG 를 확인한다
    import json, tempfile, time
    assert CGm._mermaid_status('<title>done</title><pre class="mermaid"><svg></svg></pre>') == (True, '')
    ok, why = CGm._mermaid_status('<title>mermaid-load-failed</title><pre class="mermaid">flowchart</pre>')
    assert not ok and '받지 못함' in why, why
    ok, why = CGm._mermaid_status('<title>mermaid-error: Parse error</title>')
    assert not ok and 'Parse error' in why, why
    ok, why = CGm._mermaid_status('<title>f</title><pre class="mermaid">flowchart</pre>')
    assert not ok and '그리지 못함' in why, why
    ok, why = CGm._mermaid_status('<title>done</title><pre class="mermaid">flowchart</pre>')          # 제목만 done 이고 svg 없음
    assert not ok, why
    d = tempfile.mkdtemp(prefix='cgmf_')
    A = os.path.join(d, 'a.json'); json.dump({'claims': FOCUS_G}, open(A, 'w'), ensure_ascii=False)
    js = os.path.join(d, 'mermaid.min.js'); open(js, 'w').write('/* fake */')
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    html = CGm.focus_html(CGm.focus_graph([dict(c) for c in FOCUS_G], ['f']), js)
    assert 'onerror="document.title=' + "'mermaid-load-failed'" + '"' in html and 'mermaid-error' in html, html[-600:]
    for title, svg, code, word in (('mermaid-load-failed', False, 1, '받지 못함'), ('f', False, 1, '그리지 못함'), ('done', True, 0, 'PNG:')):
        png = os.path.join(d, 'f_%s_%s.png' % (code, re.sub(r'\W', '', title)[:4]))
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', CLAIM_GRAPH_BROWSER=_fake_browser(d, title, svg), CLAIM_GRAPH_MERMAID_JS=js)
        t0 = time.time()
        r = subprocess.run([sys.executable, cg, 'focus', 'f', '--claims', A, '-o', os.path.join(d, 'f.md'), '--png', png], capture_output=True, text=True, env=env)
        assert r.returncode == code and word in r.stdout and time.time() - t0 < 25, (title, r.returncode, r.stdout, r.stderr)   # 끝나지 않는 브라우저는 끈다
        if code:
            assert not os.path.exists(png) and '[!] mermaid 가 그리지 못했다' in r.stdout and 'npm install mermaid@11' in r.stdout, r.stdout
        else:
            assert os.path.exists(png)


def t_v1616_focus_background_and_offstage_colors():
    # 사용자 09-30: 배경은 핵심 근거와 다른 색·범례 "배경" / 무대 밖 받침은 종류 색을 흐리게(범례와 맞게)
    import copy as _c, json, tempfile
    g = _c.deepcopy(FOCUS_G)
    g.append({'id': 'bg', 'role': 'background', 'statement': 'Background knowledge', 'depends_on': []})
    g.append({'id': 'ctx', 'role': 'claim', 'statement': 'Context only', 'depends_on': []})
    by = {c['id']: c for c in g}
    by['f']['depends_on'] += [{'id': 'bg', 'type': 'premise'}, {'id': 'ctx', 'type': 'context'}]
    fg = CGm.focus_graph(g, ['f'])
    assert fg['kind']['bg'] == 'context' and fg['kind']['ctx'] == 'context' and fg['kind']['e1'] == 'base', fg['kind']
    md = CGm.focus_mermaid(fg)
    assert 'classDef context' in md and CGm._FOCUS_STYLE['context'] != CGm._FOCUS_STYLE['base'], md
    assert [w for w, _ in CGm.focus_legend(fg)] == ['선택한 주장', '핵심 근거', '배경', '한계', '반박', '영향받는 결론']
    assert '배경' not in [w for w, _ in CGm.focus_legend(CGm.focus_graph(_c.deepcopy(FOCUS_G), ['f']))]       # 없으면 범례에서도 뺀다
    d = tempfile.mkdtemp(prefix='cgfc_')
    A = os.path.join(d, 'a.json'); json.dump({'claims': FOCUS_G}, open(A, 'w'), ensure_ascii=False)
    ov = CGm.oral_init(A, deck='d'); ov['use'] = {'f': {'sites': ['slide@1'], 'keys': ['k']}}
    O = os.path.join(d, 'o.json'); json.dump(ov, open(O, 'w'), ensure_ascii=False)
    fg = CGm.focus_graph(CGm.load_oral(O, A)[1], ['f'])
    md = CGm.focus_mermaid(fg)
    for k in ('base', 'limit', 'rebut'):                                              # 무대 밖이어도 종류 색(흐리게) — 회색 하나가 아니다
        assert 'classDef %s_off' % k in md, md[-900:]
    fills = set(re.findall(r'classDef \w+_off fill:(#\w+)', md))
    assert len(fills) >= 3 and '#fafafa' not in fills, fills
    assert CGm.focus_legend(fg)[-1][0] == '화면에 없음'


def t_v1610_mapfreeze_sources_missing_doi_stops():
    # 사용자 09-29: mapfreeze --sources 가 보관소에 없는 DOI 를 만나면 mapgraph 처럼 [필수] — 아무것도 기록하지 않고 멈춘다(전에는 [참고] 뒤 기록)
    import json, tempfile
    root = tempfile.mkdtemp(prefix='cgfz_')
    _store(root, md=NEW_PDF_MD)                                                          # 10.1000/abc 만 보관소에 있다
    ok = [{'id': 's', 'statement': 'S', 'evidence': 'e', 'sites': [], 'keys': [], 'depends_on': [],
           'sources': [{'kind': '문헌', 'what': '10.1000/abc', 'at': '[p.e12 · PDF 2]', 'verdict': '부합'},
                       {'kind': '문헌', 'what': 'https://doi.org/10.1000/ABC', 'at': '[p.e11 · PDF 1]', 'verdict': '부합'}]}]
    bad = copy.deepcopy(ok); bad[0]['sources'].append({'kind': '문헌', 'what': '10.1000/zzz', 'at': 'p.1', 'verdict': '부합'})
    bad.append({'id': 't', 'statement': 'T', 'sites': [], 'keys': [], 'depends_on': [], 'sources': [{'kind': '문헌', 'what': '10.1000/yyy'}]})
    # 성공 길: 모두 보관소에 있으면(대문자·https://doi.org/ 도 같은 DOI) 기록한다
    out = io.StringIO(); CGm.mapfreeze(lambda s_: '', ok, sources=root, stream=out)
    assert len(ok[0]['verified']['sources']) == 2 and '[참고]' not in out.getvalue(), (ok[0]['verified'], out.getvalue())
    # 실패 길: 하나라도 없으면 SystemExit — 어느 주장에도 verified 가 생기지 않는다
    before = copy.deepcopy(bad)
    try:
        CGm.mapfreeze(lambda s_: '', bad, sources=root, stream=io.StringIO()); assert False, '멈추지 않았다'
    except SystemExit as e:
        msg = str(e)
    assert msg.startswith('[멈춤] [필수] 근거 문헌 2곳') and 's: sources[3]: 문헌 10.1000/zzz' in msg and 't: sources[1]: 문헌 10.1000/yyy' in msg, msg
    assert '10.1000/abc' not in msg and bad == before, bad
    # 보관소 폴더를 잘못 주어도(없는 폴더) 멈춘다 — 조용히 근거 없이 기록하지 않는다
    try:
        CGm.mapfreeze(lambda s_: '', copy.deepcopy(ok), sources=os.path.join(root, 'nope'), stream=io.StringIO()); assert False
    except SystemExit as e:
        assert '근거 문헌 2곳' in str(e), e
    # --sources 를 안 주면 전처럼(문헌 근거는 보지 않고 기록)
    b2 = copy.deepcopy(bad); CGm.mapfreeze(lambda s_: '', b2)
    assert all('verified' in c and 'sources' not in c['verified'] for c in b2), b2
    # 교과서 근거를 못 찾는 것은 전처럼 [참고](결정은 문헌 DOI 만)
    tb = [{'id': 'q', 'statement': 'Q', 'sites': [], 'keys': [], 'depends_on': [], 'sources': [{'kind': '교과서', 'what': '없는책', 'at': 'p.1'}]}]
    out = io.StringIO(); CGm.mapfreeze(lambda s_: '', tb, sources=root, stream=out)
    assert 'verified' in tb[0] and '[참고] 근거 원문 1곳' in out.getvalue(), out.getvalue()
    # CLI: 종료 코드 1, 출력 파일을 만들지 않는다
    d = tempfile.mkdtemp(prefix='cgfzc_')
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    doc = os.path.join(d, 'Doc_v1.md'); open(doc, 'w', encoding='utf8').write('# Doc\n\ntext\n')
    A = os.path.join(d, 'a.json'); json.dump({'doc': 'Doc_v1.md', 'claims': before}, open(A, 'w'), ensure_ascii=False)
    O = os.path.join(d, 'o.json')
    run = lambda *a: subprocess.run([sys.executable, cg] + list(a), capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    r = run('mapfreeze', doc, '--claims', A, '--sources', root, '-o', O)
    assert r.returncode == 1 and '[필수]' in r.stderr and '10.1000/zzz' in r.stderr and not os.path.exists(O), (r.returncode, r.stdout, r.stderr)
    json.dump({'doc': 'Doc_v1.md', 'claims': ok}, open(A, 'w'), ensure_ascii=False)
    r = run('mapfreeze', doc, '--claims', A, '--sources', root, '-o', O)
    assert r.returncode == 0 and os.path.exists(O) and '기록 완료' in r.stdout, (r.stdout, r.stderr)


def t_v165_gaps_cli():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cggc_')
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    A = os.path.join(d, 'a.json'); json.dump({'doc': 'Doc_v1.docx', 'claims': GAPG}, open(A, 'w'), ensure_ascii=False)
    run = lambda *a: subprocess.run([sys.executable, cg] + list(a), capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    r = run('gaps', '--claims', A, '-o', os.path.join(d, 't.md'))
    assert r.returncode == 0 and '공백 5' in r.stdout and '# Doc_v1 — 근거 공백 작업표' in open(os.path.join(d, 't.md'), encoding='utf8').read(), (r.stdout, r.stderr)
    r = run('gaps', '--to-instr', os.path.join(d, 't.md'), '-o', os.path.join(d, 'i.md'), '--name', 'Doc_v1')
    assert r.returncode == 0 and '> 원고: Doc_v1' in open(os.path.join(d, 'i.md'), encoding='utf8').read(), r.stdout
    r = run('gaps', '-o', os.path.join(d, 'x.md'))
    assert r.returncode == 2 and '--claims' in r.stdout, r.stdout                           # 실패 길: 입력 없음
    r = run('mapgraph', '--claims', A, '--sources', d)
    assert r.returncode == 1 and '보관소에 없다' in r.stdout, r.stdout                          # 보관소에 없는 문헌 근거 = [필수] → 1



def t_v166_excluded_ddx_and_case_kind():
    # 발표 K23 (09-29): 배제된 감별이 확정 진단과 같은 상자로 그려짐 · premise 없음 [참고] 가 매번 · 논문용 [참고] 16줄
    g = [{'id': 'f1', 'group': 'c1', 'role': 'evidence', 'forbidden': ['old'], 'depends_on': []},
         {'id': 'ok', 'group': 'c1', 'role': 'claim', 'depends_on': [{'id': 'f1', 'type': 'premise'}]},
         {'id': 'ex', 'group': 'c1', 'role': 'claim', 'status': 'excluded', 'sites': ['slide@9'], 'depends_on': [{'id': 'f1', 'type': 'rebuttal'}]},
         {'id': 'ex2', 'group': 'c1', 'role': 'claim', 'status': 'excluded', 'depends_on': []},
         {'id': 'dx', 'group': 'c1', 'role': 'main', 'depends_on': [{'id': 'ok', 'type': 'premise'}]},
         {'id': 'f9', 'group': 'c2', 'role': 'evidence', 'depends_on': []},
         {'id': 'dx9', 'group': 'c2', 'role': 'main', 'depends_on': [{'id': 'f9', 'type': 'premise'}]}]
    ps = CGm.mapgraph(copy.deepcopy(g), io.StringIO())[0]
    assert not [p for p in ps if not p.startswith('[참고]')], ps                                # excluded 는 맞는 status(sites 가 있어도 됨)
    assert not any('ex:' in p and 'premise' in p for p in ps), ps                             # 배제 감별은 premise 없음 [참고] 를 내지 않고
    assert any(p.startswith('[참고] ex2: 배제(excluded)인데 rebuttal 간선이 없음') for p in ps), ps   # 배제 근거가 없으면 알린다
    assert any('forbidden 이 있는데' in p for p in ps) and any('caveat 이 없음' in p for p in ps) and any('role=main 인 주장이 2개' in p for p in ps), ps
    ps = CGm.mapgraph(copy.deepcopy(g), io.StringIO(), kind='증례')[0]                          # 증례 모드: 논문용 [참고] 셋을 끈다
    assert not any('forbidden 이 있는데' in p or 'caveat 이 없음' in p or 'role=main 인' in p for p in ps), ps
    assert any('ex2: 배제' in p for p in ps), ps                                                # 배제 근거 [참고] 는 증례에서도
    bad = copy.deepcopy(g); bad[2]['status'] = 'excludd'
    assert any('status "excludd"' in p for p in CGm.mapgraph(bad, io.StringIO())[0])           # 실패 길: 모르는 status
    md = CGm.mapdraw(copy.deepcopy(g))
    blk = md.split('```mermaid')[1].split('```')[0]
    assert 'subgraph g1["c1"]' in blk and 'subgraph g2["c2"]' in blk and blk.count('  end') == 2, blk   # 증례마다 묶음
    assert '  g1 ~~~ g2' in blk and blk.count('direction LR') == 2, blk                          # v16.8: 처음 나온 순서로 위→아래(보이지 않는 연결)
    rev = CGm.mapdraw([x for x in g if x['group'] == 'c2'] + [x for x in g if x['group'] == 'c1'])
    assert 'subgraph g1["c2"]' in rev and '  g1 ~~~ g2' in rev, rev                             # 처음 나온 순서 = 파일 순서
    one = CGm.mapdraw([x for x in g if x['group'] == 'c1'])
    assert '~~~' not in one and 'direction LR' not in one, one                                  # 실패 길: 묶음 하나면 연결·방향을 넣지 않는다
    assert '"ex<br/>claim·mid<br/>배제"' in blk and 'stroke-dasharray' in blk and 'fill:#ffffff' in blk, blk
    rc = [l for l in blk.splitlines() if l.strip().startswith('class ') and l.strip().endswith(' r_claim')][0]
    ex_nid = 'n3'
    assert ex_nid not in rc.split()[1].split(','), rc                                         # 배제 상자는 초록(claim) 색을 받지 않는다
    assert '배제된 감별' in md.split('```')[-1], md                                            # 범례 한 줄
    md = CGm.mapdraw([dict(x, group=None) for x in g])                                        # group 이 없으면 묶지 않는다
    assert 'subgraph' not in md
    rows, _ = CGm.find_gaps(copy.deepcopy(g))
    assert not any(cid in ('ex', 'ex2') for cid, _, _, _ in rows), rows                       # 배제 감별은 근거 공백에서 뺀다


def t_v166_cli_kind():
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgk_')
    cg = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'claim_graph.py')
    A = os.path.join(d, 'a.json')
    cl = [{'id': 'f', 'role': 'evidence', 'forbidden': ['x'], 'depends_on': []}, {'id': 'm', 'role': 'main', 'depends_on': [{'id': 'f', 'type': 'premise'}]}]
    json.dump({'kind': '증례', 'claims': cl}, open(A, 'w'), ensure_ascii=False)
    r = subprocess.run([sys.executable, cg, 'mapgraph', '--claims', A], capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    assert r.returncode == 0 and 'forbidden 이 있는데' not in r.stdout and 'caveat 이 없음' not in r.stdout, r.stdout
    json.dump({'claims': cl}, open(A, 'w'), ensure_ascii=False)
    r = subprocess.run([sys.executable, cg, 'mapgraph', '--claims', A], capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    assert 'forbidden 이 있는데' in r.stdout, r.stdout                                          # kind 가 없으면 전처럼



def t_v167_gaps_cited_split_and_blank_terms():
    # 저자 09-29(claims v10): "문헌 없음" 14개 중 6개는 원고 인용 [n] 이 이미 있는데 sources 만 빈 것 — 따로. 검색어는 keys 가 아니라 비움
    g = [{'id': 'c1', 'role': 'claim', 'statement': 'X is higher [12,14].', 'keys': ['argue', 'refs 12-14'], 'depends_on': []},
         {'id': 'c2', 'role': 'claim', 'statement': 'Y is linked', 'sites': ['doc:find:Y was linked [7\u20139]'], 'depends_on': []},
         {'id': 'c3', 'role': 'claim', 'statement': 'Z plausibly follows', 'keys': ['pooling'], 'depends_on': []},
         {'id': 'c4', 'role': 'background', 'statement': 'W [3]', 'confidence': 'low', 'depends_on': []},
         {'id': 'd', 'role': 'main', 'sources': [{'kind': '문헌', 'what': '10.1/a'}, {'kind': '문헌', 'what': '10.1/b'}],
          'depends_on': [{'id': 'c1', 'type': 'premise'}, {'id': 'c2', 'type': 'premise'}, {'id': 'c3', 'type': 'context'}, {'id': 'c4', 'type': 'support'}]}]
    got = {cid: ks for cid, _, ks, _ in CGm.find_gaps(copy.deepcopy(g))[0]}
    assert got['c1'] == [CGm.GAP_CITED] and got['c2'] == [CGm.GAP_CITED], got               # 문장 또는 자리에 [n]
    assert got['c3'] == ['문헌 없음'], got                                                     # 실패 길: 인용이 없으면 그대로 공백
    assert got['c4'][0] == CGm.GAP_CITED and got['c4'][1].startswith('약한 고리'), got
    md = CGm.gaps_table(copy.deepcopy(g), 'T')
    p1, p2 = md.split('## 2. 인용 있음')
    assert '## 1. 찾을 공백 2개' in p1 and '`c3`' in p1 and '`c4`' in p1 and '`c1`' not in p1, p1   # 인용 + 다른 공백(c4)은 1부에
    assert '— sources 미기입 2개' in p2 and '| G03-기입 | `c1`' in p2 and '| G04-기입 | `c2`' in p2 and 'G03-받침' not in md, p2
    assert 'argue' not in md and 'pooling' not in md and 'refs 12-14' not in md, md            # keys 가 검색어로 들어가지 않는다
    assert '문헌 없음(인용 있음 — sources 미기입) 3' in md and '문헌 없음 1' in md, md.splitlines()[2]   # 머리 요약도 두 부류를 나눠 센다
    rows = []
    for l in md.splitlines():
        c = [x.strip() for x in l.strip().strip('|').split('|')]
        if c and c[0] == 'G03-기입':
            c[6], c[7] = '10.7777/cited', '사람'
            l = '| ' + ' | '.join(c) + ' |'
        rows.append(l)
    instr, notes = CGm.gaps_to_instr('\n'.join(rows), 'T')
    assert 'doi:10.7777/cited' in instr and '| G03-기입 | 1 |' in instr, instr                     # 기입 줄도 검증지시로
    assert any('G03-기입: 검색어 칸이 비었다' in n for n in notes), notes                       # 검색어가 비면 알린다


def run():
    results = []
    for name, fn in sorted(globals().items()):
        if name.startswith('t_'):
            try:
                fn(); results.append(('PASS', name[2:], ''))
            except Exception as e:
                results.append(('FAIL', name[2:], '%s: %s' % (type(e).__name__, e)))
                traceback.print_exc()
    w = max(len(n) for _, n, _ in results)
    for st, n, msg in results:
        print('%-4s %-*s %s' % (st, w, n, msg))
    f = sum(1 for s, _, _ in results if s == 'FAIL')
    print('\n통과 %d / 건너뜀 0 / 실패 %d  (전체 %d)' % (len(results) - f, f, len(results)))   # v15.8: 건너뜀 칸을 모든 테스트에 같은 모양으로(발표 요청)
    return f


def t_v169_nums_suppl_tagged_one_line():
    """저자 3a: evidence 맨 앞 `Suppl S{n}` 표지(실물 15개 모두 이 모양 — 수치는 뒤 조각)의 못 찾은 수치는
    주장마다 따로 적지 않고 한 줄로. 표지 없는 evidence 는 전처럼 주장마다."""
    site = {'doc:p:1': 'Main text r = 0.40 only.'}
    cl = [{'id': 's1', 'statement': 'x', 'evidence': 'Suppl S3; r = 0.63, n = 42', 'sites': ['doc:p:1'], 'keys': []},
          {'id': 's2', 'statement': 'y', 'evidence': 'Suppl S4; 12.5%', 'sites': ['doc:p:1'], 'keys': []},
          {'id': 'm1', 'statement': 'z', 'evidence': 'r = 0.51', 'sites': ['doc:p:1'], 'keys': []},
          {'id': 'ok', 'statement': 'w', 'evidence': 'Suppl S5; r = 0.40', 'sites': ['doc:p:1'], 'keys': []}]
    buf = io.StringIO()
    probs, _ = CGm.mapcheck(_resolve_dict(site), cl, stream=buf, nums=True)
    per = [p for p in probs if 'evidence 의 수치' in p]
    assert per == ['[참고] m1: evidence 의 수치 0.51 가 sites 어디에도 없음'], probs          # 실패 길(옛 코드): s1·s2 도 한 줄씩
    one = [p for p in probs if '보충자료' in p]
    assert one == ['[참고] 보충자료 표지(Suppl S…)가 든 evidence 2개(s1, s2)의 수치 3개가 sites 에 없음 — sites 는 본문만 보므로 따로 적지 않는다'], probs
    assert not any('ok' in p for p in probs)                                                    # 성공 길: 찾은 수치는 조용


def t_v169_remap_refs_suppl():
    """저자 3b: 보충 표 재번호 — `Suppl S{n}` 도 매핑으로(`--suppl`). 본문 [n] 은 건드리지 않는다."""
    cl = [{'id': 'a', 'statement': 'A [3]', 'evidence': 'Suppl S3; r = 0.6', 'keys': []},
          {'id': 'b', 'statement': 'B', 'evidence': 'Suppl S9; x', 'keys': []},
          {'id': 'c', 'statement': 'C', 'evidence': 'Supplementary Table S4 and S2', 'keys': []}]
    buf = io.StringIO()
    ch = CGm.remap_refs(cl, {}, stream=buf, suppl={3: 4, 4: 5, 2: 2})
    assert cl[0]['evidence'] == 'Suppl S4; r = 0.6' and cl[0]['statement'] == 'A [3]', cl[0]
    assert cl[1]['evidence'] == 'Suppl S9; x' and 'S9 매핑에 없음' in buf.getvalue(), buf.getvalue()   # 실패 길: 없는 번호는 그대로 + 경고
    assert cl[2]['evidence'] == 'Supplementary Table S5 and S2', cl[2]
    assert len(ch) == 2
    import json, subprocess
    p = '/tmp/cg_suppl.json'; CGm.save_claims(p, [dict(c) for c in cl], meta={'doc': 'm.docx'})
    json.dump({'map': {'4': 7}}, open('/tmp/cg_smap.json', 'w'))
    run = lambda *a: subprocess.run([sys.executable, CGm.__file__, 'remap-refs', '--claims'] + list(a), capture_output=True, text=True)
    r = run(p, '--suppl', '/tmp/cg_smap.json', '-o', '/tmp/cg_suppl2.json')
    assert r.returncode == 0 and json.load(open('/tmp/cg_suppl2.json'))['claims'][0]['evidence'] == 'Suppl S7; r = 0.6', r.stdout + r.stderr
    r = run('/tmp/cg_suppl2.json', '--suppl', '/tmp/cg_smap.json', '-o', '/tmp/cg_suppl3.json')
    assert r.returncode == 2 and '이미 적용' in r.stdout, r.stdout                                  # 두 번 적용 거부
    r = run(p, '-o', '/tmp/cg_suppl4.json')
    assert r.returncode == 2 and '--map' in r.stdout and '--suppl' in r.stdout, r.stdout + r.stderr  # 둘 다 없으면 멈춤


def t_v169_add_and_link_cli():
    """저자 3c: add/link 가 간선 weight 를 type 기본값으로 넣는다. 없는 id·겹친 간선·순환은 멈춘다."""
    import json, subprocess
    p = '/tmp/cg_add.json'; CGm.save_claims(p, copy.deepcopy(GRAPH), meta={'doc': 'm.docx', 'note': 'n'})
    run = lambda *a: subprocess.run([sys.executable, CGm.__file__] + list(a), capture_output=True, text=True)
    r = run('add', '--claims', p, '--id', 'd', '--statement', 'D', '--role', 'claim', '--dep', 'b:support', '--dep', 'a',
            '--site', 'notes:9', '--key', 'dee', '-o', '/tmp/cg_add2.json')
    assert r.returncode == 0, r.stdout + r.stderr
    d = json.load(open('/tmp/cg_add2.json'))
    assert d['doc'] == 'm.docx' and d['note'] == 'n'                                                 # 맨 위 칸 보존
    new = [c for c in d['claims'] if c['id'] == 'd'][0]
    assert new['depends_on'] == [{'id': 'b', 'type': 'support', 'weight': 0.7}, {'id': 'a', 'type': 'premise', 'weight': 1.0}], new
    assert new['sites'] == ['notes:9'] and new['keys'] == ['dee'] and new['role'] == 'claim'
    r = run('link', '--claims', '/tmp/cg_add2.json', 'c', 'a', '--type', 'caveat', '-o', '/tmp/cg_add3.json')
    assert r.returncode == 0, r.stdout + r.stderr
    c = [x for x in json.load(open('/tmp/cg_add3.json'))['claims'] if x['id'] == 'c'][0]
    assert {'id': 'a', 'type': 'caveat', 'weight': 0.5} in c['depends_on'], c
    for args, why in ((['add', '--claims', p, '--id', 'a', '--statement', 'x', '-o', '/tmp/x.json'], '이미 있다'),
                      (['add', '--claims', p, '--id', 'e', '--statement', 'x', '--dep', 'zz', '-o', '/tmp/x.json'], '없는 주장'),
                      (['link', '--claims', p, 'b', 'a', '-o', '/tmp/x.json'], '이미 있다'),
                      (['link', '--claims', p, 'a', 'c', '-o', '/tmp/x.json'], '순환'),
                      (['link', '--claims', p, 'a', 'b', '--type', 'bogus', '-o', '/tmp/x.json'], 'invalid choice')):
        r = run(*args)
        assert r.returncode != 0 and why in (r.stdout + r.stderr), (args, r.stdout, r.stderr)


def t_v169_doc_index_zero_or_out_of_range():
    """리뷰어 3: doc:tbl:N·doc:p:N 은 1부터 — 0 은 파이썬 음수 번호로 **마지막 표·문단**을 돌려줬다(조용한 오답). 이제 알린다."""
    p = '/tmp/cg_test.docx'; _mk_docx(p)
    src = CGm.DocSource(p)
    for site in ('doc:tbl:0', 'doc:tbl:2', 'doc:p:0', 'doc:p:999'):
        try:
            src.resolve(site); assert False, site                                                    # 실패 길(옛 코드): tbl:0 = 마지막 표
        except KeyError as e:
            assert '1부터' in str(e), e
    assert '0.612' in src.resolve('doc:tbl:1') and src.resolve('doc:p:1')                              # 성공 길


def t_v169_positional_claims_hint():
    """발표 3: mapgraph claims.json 처럼 파일 이름만 주면 '--claims' 를 알려 준다."""
    import subprocess
    p = '/tmp/cg_pos.json'; CGm.save_claims(p, copy.deepcopy(GRAPH))
    for cmd in ('mapgraph', 'mapreport', 'scaffold', 'gaps'):
        r = subprocess.run([sys.executable, CGm.__file__, cmd, p], capture_output=True, text=True)
        assert r.returncode == 2 and ('%s --claims %s' % (cmd, p)) in (r.stdout + r.stderr), (cmd, r.stdout, r.stderr)
    r = subprocess.run([sys.executable, CGm.__file__, 'mapgraph', '--claims', p], capture_output=True, text=True)
    assert r.returncode in (0, 1) and '--claims' not in r.stderr


def t_v169_mapcheck_pass_wording():
    """저자 4: '통과' 는 '자리가 있다' 이지 '최신' 이 아니다 — 문구로 알린다."""
    buf = io.StringIO()
    probs, _ = CGm.mapcheck(_resolve_dict({'slide:7': 'title', 'notes:7': 'note', 'notes:8': ''}), copy.deepcopy(GRAPH), stream=buf)
    assert not probs and '모든 주장의 자리에 찾는 표현이 있음 — 주장·evidence 가 최신인지는 보지 않는다(mapstale·판 올림 때 evidence 갱신)' in buf.getvalue(), buf.getvalue()
    assert '반영됨' not in buf.getvalue()


def t_v169_selfcheck_env_row():
    """발표 4: '통과' 는 환경에 기댄 판정 — selfcheck 표에 Python·pypdf·Pillow·python-pptx 판을 적는다."""
    import platform
    buf = io.StringIO()
    CGm.selfcheck(os.path.dirname(os.path.abspath(CGm.__file__)), stream=buf)
    t = buf.getvalue()
    row = [l for l in t.splitlines() if l.startswith('| 환경 |')]
    assert len(row) == 1 and 'Python %s' % platform.python_version() in row[0], t[:800]
    for lib in ('pypdf', 'Pillow', 'python-pptx'):
        assert lib + ' ' in row[0], row
    assert CGm._lib_versions({'nope-lib-xyz': 'nope'}) == 'nope 없음'                            # 없는 라이브러리는 '없음'



# ---- v16.17 (사용자 09-30 ③): 탐색적 표지 exploratory · suggest ----

def _xg():
    """가짜 그래프: mn(main) ←premise cl ←premise ev2 · mn ←support ev1 · ev4 는 main 과 안 이어짐(ev4 → x2 context)."""
    return [
        {'id': 'ev1', 'role': 'evidence', 'statement': 'E1', 'depends_on': []},
        {'id': 'ev2', 'role': 'evidence', 'statement': 'E2', 'depends_on': []},
        {'id': 'ev4', 'role': 'evidence', 'statement': 'E4', 'depends_on': []},
        {'id': 'cl', 'role': 'claim', 'statement': 'C', 'depends_on': [{'id': 'ev2', 'type': 'premise'}]},
        {'id': 'x2', 'role': 'claim', 'statement': 'X2', 'depends_on': [{'id': 'ev4', 'type': 'context'}, {'id': 'ev1', 'type': 'premise'}]},
        {'id': 'mn', 'role': 'main', 'statement': 'M', 'depends_on': [{'id': 'cl', 'type': 'premise'}, {'id': 'ev1', 'type': 'support'}]}]


def _by(g):
    return {c['id']: c for c in g}


def _must(probs, word):
    return [p for p in probs if not p.startswith('[참고]') and word in p]


def _ref(probs, word):
    return [p for p in probs if p.startswith('[참고]') and word in p]


def t_v1617_exploratory_premise_chain_must_and_reason():
    """③: 탐색적 주장이 main 의 premise 사슬에 있으면 [필수] — exploratory_reason 을 적으면 [참고]로 내리고 사유를 보인다."""
    g = _xg(); _by(g)['ev2']['exploratory'] = True
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert len(_must(probs, 'ev2')) == 1 and '탐색' in _must(probs, 'ev2')[0], probs          # 실패 길
    _by(g)['ev2']['exploratory_reason'] = '사전 계획 분석 — Methods 절'
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not _must(probs, 'ev2'), probs                                                      # 성공 길
    assert any('사전 계획 분석 — Methods 절' in p for p in _ref(probs, 'ev2')), probs
    g = _xg(); _by(g)['mn']['exploratory'] = True                                              # main 자신도 사슬
    assert _must(CGm.mapgraph(g, io.StringIO())[0], 'mn')


def t_v1617_exploratory_support_path_ref_off_chain_silent():
    """support 가 섞인 경로로만 main 에 닿으면 [참고], 이어지지 않으면 mapgraph 는 말하지 않는다."""
    g = _xg(); _by(g)['ev1']['exploratory'] = True; _by(g)['ev4']['exploratory'] = True
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not _must(probs, '탐색'), probs
    assert _ref(probs, 'ev1') and any('support' in p for p in _ref(probs, 'ev1')), probs
    assert not [p for p in probs if p.startswith('[참고] ev4') and '탐색' in p], probs


def t_v1617_exploratory_field_checks_and_case_mains():
    """exploratory 는 true/false 만([필수]) · 사유만 있고 표지가 없으면 [참고] · 증례(main 여럿)는 main 마다 사슬."""
    g = _xg(); _by(g)['cl']['exploratory'] = 'yes'
    assert _must(CGm.mapgraph(g, io.StringIO())[0], 'exploratory'), 'true/false 가 아닌 값'
    g = _xg(); _by(g)['cl']['exploratory_reason'] = '까닭'
    assert _ref(CGm.mapgraph(g, io.StringIO())[0], 'exploratory_reason')
    g = _xg(); _by(g)['cl']['exploratory'] = False; _by(g)['cl']['exploratory_reason'] = ''
    probs, _ = CGm.mapgraph(g, io.StringIO())
    assert not [p for p in probs if 'exploratory' in p or '탐색' in p], probs                  # false·빈 사유는 조용
    case = [{'id': 'f1', 'group': 'g1', 'role': 'evidence', 'statement': 'F1', 'exploratory': True, 'depends_on': []},
            {'id': 'd1', 'group': 'g1', 'role': 'main', 'statement': 'D1', 'depends_on': [{'id': 'f1', 'type': 'premise'}]},
            {'id': 'f2', 'group': 'g2', 'role': 'evidence', 'statement': 'F2', 'depends_on': []},
            {'id': 'd2', 'group': 'g2', 'role': 'main', 'statement': 'D2', 'depends_on': [{'id': 'f2', 'type': 'premise'}]}]
    probs, _ = CGm.mapgraph(case, io.StringIO(), kind='증례')
    assert len(_must(probs, 'f1')) == 1 and 'd1' in _must(probs, 'f1')[0] and not _must(probs, 'f2'), probs
    case[0]['status'] = 'superseded'                                                            # 철회한 주장은 보지 않는다
    assert not _must(CGm.mapgraph(case, io.StringIO(), kind='증례')[0], '탐색')


def t_v1617_mapreport_exploratory_section():
    """mapreport(리뷰어 보고) 끝에 탐색적 주장 절 — 사슬 안·밖과 사유."""
    g = _xg(); _by(g)['ev2'].update(exploratory=True, exploratory_reason='사전 계획 분석'); _by(g)['ev4']['exploratory'] = True
    buf = io.StringIO(); CGm.mapreport(g, buf); t = buf.getvalue()
    assert '탐색적 주장' in t, t
    l2 = [l for l in t.splitlines() if l.strip().startswith('[ev2]') and '사전 계획 분석' in l]
    l4 = [l for l in t.splitlines() if l.strip().startswith('[ev4]') and 'main 과 이어지지 않음' in l]
    assert l2 and '전제 사슬' in l2[0] and l4, t
    buf = io.StringIO(); CGm.mapreport(_xg(), buf)
    assert '탐색적 주장' not in buf.getvalue()                                                   # 없으면 절도 없다


def t_v1617_mapstale_exploratory_removed_is_change():
    """사용자 09-30: 표지를 지우는 것(true→false·칸 삭제)은 [필수] 를 사유 없이 넘는 길 — mapstale [변경]. false→true 는 [참고]."""
    site = {'doc:p:1': 'alpha text'}
    base = [{'id': 'a', 'statement': 'A', 'evidence': 'e', 'sites': ['doc:p:1'], 'keys': ['alpha'], 'exploratory': True, 'depends_on': []},
            {'id': 'b', 'statement': 'B', 'evidence': 'e', 'sites': ['doc:p:1'], 'keys': ['alpha'], 'depends_on': [{'id': 'a', 'type': 'premise'}]}]
    fr = CGm.mapfreeze(_resolve_dict(site), copy.deepcopy(base))
    for how in ('false', 'del'):
        g = copy.deepcopy(fr)
        if how == 'false':
            g[0]['exploratory'] = False
        else:
            del g[0]['exploratory']
        buf = io.StringIO(); r = CGm.mapstale(_resolve_dict(site), g, stream=buf)
        assert r['changed'] == ['a'] and '탐색' in buf.getvalue() and 'b' in [x[0] for x in r['suspect']], (how, buf.getvalue())
    buf = io.StringIO(); r = CGm.mapstale(_resolve_dict(site), copy.deepcopy(fr), stream=buf)   # 성공 길: 그대로면 조용
    assert not r['changed'] and '바뀐 것 없음' in buf.getvalue(), buf.getvalue()
    g = copy.deepcopy(fr); g[1]['exploratory'] = True                                          # false → true: [참고] 한 줄
    buf = io.StringIO(); r = CGm.mapstale(_resolve_dict(site), g, stream=buf)
    assert not r['changed'] and any(l.strip().startswith('[참고] b') and '탐색' in l for l in buf.getvalue().splitlines()), buf.getvalue()
    old = copy.deepcopy(fr); old[0]['verified'].pop('exploratory'); old[0]['exploratory'] = False   # 옛 freeze(기록 없음)는 보지 않는다
    assert not CGm.mapstale(_resolve_dict(site), old, stream=io.StringIO())['changed']


def _sg():
    """suggest 용 가짜 그래프."""
    L = lambda i: {'id': i, 'role': 'caveat', 'statement': i, 'depends_on': []}
    cav = lambda *ks: [{'id': k, 'type': 'caveat'} for k in ks]
    lit = lambda *ds: [{'kind': '문헌', 'what': d, 'at': '전체'} for d in ds]
    return [
        L('k-all'), L('k-two'),
        {'id': 'e1', 'role': 'evidence', 'statement': 'e1', 'depends_on': cav('k-all')},
        {'id': 'e2', 'role': 'evidence', 'statement': 'e2', 'depends_on': cav('k-all')},
        {'id': 'e3', 'role': 'evidence', 'statement': 'e3', 'depends_on': cav('k-all', 'k-two')},
        {'id': 'e-unused', 'role': 'evidence', 'statement': 'eu', 'depends_on': cav('k-two')},
        {'id': 'e-rb', 'role': 'evidence', 'statement': 'er', 'depends_on': []},
        {'id': 'c1', 'role': 'claim', 'statement': 'c1', 'depends_on': [{'id': 'e1', 'type': 'premise'}, {'id': 'e2', 'type': 'support'}]},
        {'id': 'c2', 'role': 'claim', 'statement': 'c2', 'depends_on': [{'id': 'e1', 'type': 'premise'}, {'id': 'e2', 'type': 'premise'}]},
        {'id': 'c3', 'role': 'claim', 'statement': 'c3', 'depends_on': [{'id': 'e1', 'type': 'support'}, {'id': 'e2', 'type': 'support'}]},
        {'id': 'c4', 'role': 'claim', 'statement': 'c4', 'depends_on': [{'id': 'e1', 'type': 'premise'}, {'id': 'e3', 'type': 'premise'}, {'id': 'c1', 'type': 'support'}]},
        {'id': 'c5', 'role': 'claim', 'statement': 'c5', 'depends_on': [{'id': 'e3', 'type': 'premise'}, {'id': 'e-rb', 'type': 'rebuttal'}]},
        {'id': 'b1', 'role': 'background', 'statement': 'b1', 'sources': lit('10.1/a', '10.1/b'), 'depends_on': []},
        {'id': 'b2', 'role': 'background', 'statement': 'b2', 'sources': lit('10.1/a', '10.1/b', '10.1/c'), 'depends_on': []},
        {'id': 'b3', 'role': 'background', 'statement': 'b3', 'sources': lit('10.1/a'), 'depends_on': []},
        {'id': 'old', 'role': 'claim', 'status': 'superseded', 'statement': 'o', 'depends_on': [{'id': 'e1', 'type': 'premise'}, {'id': 'e2', 'type': 'premise'}]},
        {'id': 'mn', 'role': 'main', 'statement': 'm', 'depends_on': [{'id': 'c2', 'type': 'premise'}, {'id': 'c4', 'type': 'premise'}, {'id': 'c5', 'type': 'support'},
                                                                        {'id': 'b1', 'type': 'context'}, {'id': 'b2', 'type': 'context'}]}]


def t_v1617_suggest_rules():
    """규칙 1 같은 근거 2개 이상·안 이어짐(묶어 한 줄, 문헌 DOI 따로) · 규칙 2 안 쓰인 evidence · 규칙 3 공통 한계 3개 이상 + main 에 직접 걸린 한계 0 머리 줄."""
    g = _sg(); before = copy.deepcopy(g)
    buf = io.StringIO(); r = CGm.suggest(g, stream=buf); t = buf.getvalue()
    assert g == before                                                                          # claims 를 바꾸지 않는다
    groups = [sorted(x['ids']) for x in r['shared']]
    assert ['c1', 'c2', 'c3'] in groups, groups                                                 # 셋이 e1·e2 를 함께 — 한 줄로 묶음
    assert not any('c4' in x for x in groups), groups                                           # c4 는 c1 에 이어짐 · 공유 1
    assert not any('old' in x for x in groups), groups                                          # 철회는 뺀다
    assert ['b1', 'b2'] in groups and not any('b3' in x for x in groups), groups                # DOI 공유 2 / 1
    assert [x['id'] for x in r['unused']] == ['e-unused'], r['unused']                         # e-rb 는 rebuttal 로 쓰임
    assert [x['id'] for x in r['caveats']] == ['k-all'], r['caveats']                          # 3 이상만(k-two 는 2)
    lines = t.splitlines()
    head = [i for i, l in enumerate(lines) if 'main 에 직접 걸린 한계 0개 — 아래 공통 한계 중 main 에 걸 것을 고른다' in l]
    k = [i for i, l in enumerate(lines) if 'k-all' in l and '[참고]' in l]
    assert head and k and head[0] < k[0], t                                                     # 머리 줄이 먼저, 줄은 그대로
    assert all(l.lstrip().startswith(('[참고]', '===', '규칙', '-', '')) for l in lines), t
    g = _sg(); _by(g)['mn']['depends_on'].append({'id': 'k-all', 'type': 'caveat'})           # main 에 한계가 걸려 있으면 머리 줄 없음
    buf = io.StringIO(); r = CGm.suggest(g, stream=buf)
    assert '직접 걸린 한계 0개' not in buf.getvalue() and r['caveats'], buf.getvalue()
    r = CGm.suggest(_sg(), stream=io.StringIO(), min_shared=3, min_caveat=2)                    # 기준 옵션
    assert not r['shared'] and sorted(x['id'] for x in r['caveats']) == ['k-all', 'k-two'], r


def t_v1617_suggest_half_chain_and_case():
    """main 사슬 절반 넘게 걸렸는데 main 에는 없는 한계 · 증례 그래프는 규칙 1 을 끈다."""
    g = [{'id': 'k', 'role': 'caveat', 'statement': 'k', 'depends_on': []},
         {'id': 'e1', 'role': 'evidence', 'statement': 'e1', 'depends_on': [{'id': 'k', 'type': 'caveat'}]},
         {'id': 'e2', 'role': 'evidence', 'statement': 'e2', 'depends_on': [{'id': 'k', 'type': 'caveat'}]},
         {'id': 'e3', 'role': 'evidence', 'statement': 'e3', 'depends_on': [{'id': 'k', 'type': 'caveat'}]},
         {'id': 'mn', 'role': 'main', 'statement': 'm', 'depends_on': [{'id': t, 'type': 'premise'} for t in ('e1', 'e2', 'e3')]}]
    r = CGm.suggest(g, stream=io.StringIO())
    assert [x['id'] for x in r['half']] == ['k'], r['half']
    g[4]['depends_on'].append({'id': 'k', 'type': 'caveat'})
    assert not CGm.suggest(g, stream=io.StringIO())['half']
    case = [{'id': 'f1', 'role': 'evidence', 'statement': 'f1', 'depends_on': []},
            {'id': 'f2', 'role': 'evidence', 'statement': 'f2', 'depends_on': []},
            {'id': 'd1', 'role': 'claim', 'statement': 'd1', 'depends_on': [{'id': 'f1', 'type': 'premise'}, {'id': 'f2', 'type': 'support'}]},
            {'id': 'd2', 'role': 'claim', 'statement': 'd2', 'depends_on': [{'id': 'f1', 'type': 'support'}, {'id': 'f2', 'type': 'premise'}]}]
    assert CGm.suggest(copy.deepcopy(case), stream=io.StringIO())['shared']
    buf = io.StringIO(); r = CGm.suggest(case, stream=buf, kind='증례')
    assert not r['shared'] and '증례' in buf.getvalue(), buf.getvalue()


def t_v1617_suggest_cli_and_add_exploratory():
    """suggest CLI: -o md · 종료 0 · claims 파일 그대로. add --exploratory 는 표지를 단다."""
    import json, tempfile, hashlib
    d = tempfile.mkdtemp()
    p = os.path.join(d, 'c.json'); json.dump({'doc': 'Fake', 'claims': _sg()}, open(p, 'w'), ensure_ascii=False)
    h = hashlib.sha256(open(p, 'rb').read()).hexdigest()
    o = os.path.join(d, 's.md')
    r = subprocess.run([sys.executable, CGm.__file__, 'suggest', '--claims', p, '-o', o], capture_output=True, text=True)
    assert r.returncode == 0 and os.path.exists(o) and 'k-all' in open(o, encoding='utf8').read(), (r.stdout, r.stderr)
    assert hashlib.sha256(open(p, 'rb').read()).hexdigest() == h
    r = subprocess.run([sys.executable, CGm.__file__, 'suggest', p], capture_output=True, text=True)
    assert r.returncode == 2 and 'suggest --claims' in r.stdout, r.stdout                     # 파일 이름만 주면 고칠 명령
    q = os.path.join(d, 'q.json')
    r = subprocess.run([sys.executable, CGm.__file__, 'add', '--claims', p, '--id', 'nx', '--statement', 'N', '--dep', 'c1:support', '--exploratory', '-o', q],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    assert _by(json.load(open(q))['claims'])['nx']['exploratory'] is True
    r = subprocess.run([sys.executable, CGm.__file__, 'add', '--claims', p, '--id', 'ny', '--statement', 'N', '-o', q], capture_output=True, text=True)
    assert r.returncode == 0 and 'exploratory' not in _by(json.load(open(q))['claims'])['ny']


def t_v1617_draw_and_focus_mark():
    """mapdraw(전체·--all-edges)·focus 상자에 '탐색' — 색은 더하지 않는다. 없으면 표시·범례도 없다."""
    g = _xg(); _by(g)['ev2']['exploratory'] = True
    for out in (CGm.mapdraw(g), CGm.mapdraw(g, all_edges=True)):
        node = [l for l in out.splitlines() if '"ev2<br/>' in l]
        assert node and '탐색' in node[0], out
        assert not [l for l in out.splitlines() if '"ev1<br/>' in l and '탐색' in l]
    assert '탐색' not in CGm.mapdraw(_xg())
    fg = CGm.focus_graph(g, ['cl'])
    assert '(탐색)' in CGm.focus_mermaid(fg) and any('탐색' in w for w, _ in CGm.focus_legend(fg))
    fg = CGm.focus_graph(_xg(), ['cl'])
    assert '탐색' not in CGm.focus_mermaid(fg) and not any('탐색' in w for w, _ in CGm.focus_legend(fg))


def t_v1617_oral_merge_keeps_exploratory():
    """구연 합친 그래프에서도 저자 표지가 따라온다 — 무대 밖 받침이 탐색적이면 oral check 가 [필수]."""
    import json, tempfile
    d = tempfile.mkdtemp()
    A = os.path.join(d, 'a.json'); O = os.path.join(d, 'o.json')
    au = _xg(); _by(au)['ev2']['exploratory'] = True
    json.dump({'doc': 'Fake', 'claims': au}, open(A, 'w'), ensure_ascii=False)
    ov = CGm.oral_init(A, deck='d'); ov['use'] = {'mn': {'sites': ['slide@300'], 'keys': ['m']}}
    json.dump(ov, open(O, 'w'), ensure_ascii=False)
    _, merged, _ = CGm.load_oral(O, A)
    assert _by(merged)['ev2'].get('exploratory') is True and _by(merged)['ev2'].get('offstage')
    r = subprocess.run([sys.executable, CGm.__file__, 'oral', 'check', '--author', A, '--oral', O], capture_output=True, text=True)
    assert r.returncode == 1 and 'ev2' in r.stdout and '탐색' in r.stdout, r.stdout



# ---- v16.18 (저자 도구회신 09-30 suggest 첫 실행) ----

def _hg(n_chain, n_on, on_main=False):
    """main 이 premise 로 e1..e{n_chain} 에 기대고, 한계 k 가 e1..e{n_on} 에 걸린 가짜 그래프."""
    g = [{'id': 'k', 'role': 'caveat', 'statement': '단면 연구라 인과를 말할 수 없다 — 추적 자료가 없고 표본이 한 기관이다', 'depends_on': []}]
    for i in range(1, n_chain + 1):
        g.append({'id': 'e%d' % i, 'role': 'evidence', 'statement': 'e%d' % i,
                  'depends_on': [{'id': 'k', 'type': 'caveat'}] if i <= n_on else []})
    g.append({'id': 'mn', 'role': 'main', 'statement': 'm', 'depends_on': [{'id': 'e%d' % i, 'type': 'premise'} for i in range(1, n_chain + 1)]
              + ([{'id': 'k', 'type': 'caveat'}] if on_main else [])})
    return g


def t_v1618_half_chain_counts_main():
    """[확인 필요] 덧줄의 사슬은 설계대로 main 을 포함해 센다(premise 만). v16.17 은 main 을 빼고 세어 저자 v10 에서 3줄(코드 실측 1줄).
    사슬 = main + 5 → 6. 걸린 3 은 절반을 넘지 않는다(3*2 = 6), 걸린 4 는 넘는다."""
    assert not CGm.suggest(_hg(5, 3), stream=io.StringIO())['half']                      # 실패 길(v16.17 은 3*2 > 5 로 냈다)
    buf = io.StringIO(); r = CGm.suggest(_hg(5, 4), stream=buf)
    assert [(x['id'], x['on'], x['of']) for x in r['half']] == [('k', 4, 6)], r['half']   # 성공 길
    assert 'main(mn) 전제 사슬 6개 중 4개' in buf.getvalue(), buf.getvalue()
    assert not CGm.suggest(_hg(5, 5, on_main=True), stream=io.StringIO())['half']        # main 에 걸리면 덧줄 없음


def t_v1618_rule3_already_on_main_and_statement():
    """이미 main 에 걸린 공통 한계는 권고 대신 '(main 에 이미 걸림)' · 규칙 3 줄마다 한계 statement 앞 40자."""
    buf = io.StringIO(); CGm.suggest(_hg(4, 3), stream=buf)                              # main 에 없음 — 권고 그대로
    line = [l for l in buf.getvalue().splitlines() if l.strip().startswith('[참고] k ')]
    assert line and '→ main 에 직접 걸기' in line[0] and '이미 걸림' not in line[0], buf.getvalue()
    assert '"단면 연구라 인과를 말할 수 없다 — 추적 자료가 없고 표본이 한 기관이…"' in line[0], line
    buf = io.StringIO(); CGm.suggest(_hg(4, 3, on_main=True), stream=buf)                # main 에 걸림 — 권고를 빼고 표시
    line = [l for l in buf.getvalue().splitlines() if l.strip().startswith('[참고] k ')]
    assert line and '(main 에 이미 걸림)' in line[0] and '직접 걸기' not in line[0], buf.getvalue()
    buf = io.StringIO(); CGm.suggest(_hg(3, 3), stream=buf)
    hl = [l for l in buf.getvalue().splitlines() if '전제 사슬' in l]
    assert hl and '"단면 연구라' in hl[0], buf.getvalue()



# ---- v16.19: 발표 S1(oral sync 간선 변화) · ④ 1판 문헌 그래프(paper.md 자리·kind 문헌) ----

def _s1_files(d, v2_deps, old_snap=False):
    import json
    A, O, A2 = [os.path.join(d, x) for x in ('a1.json', 'o.json', 'a2.json')]
    base = [{'id': 'cav-a', 'role': 'caveat', 'statement': 'CA', 'depends_on': []},
            {'id': 'cav-b', 'role': 'caveat', 'statement': 'CB', 'depends_on': []},
            {'id': 'ev1', 'role': 'evidence', 'statement': 'E1', 'depends_on': []},
            {'id': 'ev2', 'role': 'evidence', 'statement': 'E2', 'depends_on': []},
            {'id': 'mn', 'role': 'main', 'statement': 'M', 'depends_on': [{'id': 'ev1', 'type': 'premise'}, {'id': 'cav-a', 'type': 'caveat'},
                                                                            {'id': 'ev2', 'type': 'support'}]}]
    json.dump({'doc': 'Fake', 'claims': base}, open(A, 'w'), ensure_ascii=False)
    ov = CGm.oral_init(A, deck='d')
    ov['use'] = {'mn': {'sites': ['slide@300'], 'keys': ['m']}, 'cav-a': {'sites': ['slide@301'], 'keys': ['ca']}}
    if old_snap:
        for v in ov['source']['snap'].values():
            v.pop('deps', None)
    json.dump(ov, open(O, 'w'), ensure_ascii=False)
    new = copy.deepcopy(base)
    for c in new:
        if c['id'] in v2_deps:
            c['depends_on'] = v2_deps[c['id']]
    json.dump({'doc': 'Fake', 'claims': new}, open(A2, 'w'), ensure_ascii=False)
    return A, O, A2


def t_v1619_oral_sync_reports_edge_changes():
    """발표 S1(09-30): v10→v11 처럼 간선만 바뀐 판에서 oral sync 가 아무 줄도 내지 않았다 — 결론 화면에 새 한계가 걸렸는데(화면에 없음) 모름.
    쓴 주장·무대 밖 상류의 간선이 늘거나 줄면 [참고] 한 줄, 새로 걸린 쪽이 화면에 있는지도."""
    import json, tempfile
    d = tempfile.mkdtemp(prefix='cgs1_')
    v2 = {'mn': [{'id': 'ev1', 'type': 'premise'}, {'id': 'cav-a', 'type': 'caveat'}, {'id': 'cav-b', 'type': 'caveat'}],
          'ev1': [{'id': 'cav-a', 'type': 'caveat'}]}
    A, O, A2 = _s1_files(d, v2)
    am, ac = CGm.load_claims_full(A2)
    new, lines, hard = CGm.oral_sync(json.load(open(O)), am, ac, CGm._file_sha(A2), 'a2.json')
    t = '\n'.join(lines)
    assert not hard and new, (lines, hard)
    assert '[참고] 쓴 주장 mn 에 caveat cav-b 가 새로 걸림 — 화면에 없음' in t, t
    assert '[참고] 쓴 주장 mn 에서 support ev2 가 빠짐' in t, t
    assert '[참고] 무대 밖 상류 ev1 에 caveat cav-a 가 새로 걸림 — 화면 slide@301' in t, t
    assert all('deps' in v for v in new['source']['snap'].values())                              # 새 스냅숏에 간선
    A, O, A2 = _s1_files(tempfile.mkdtemp(prefix='cgs1_'), {})                                   # 성공 길: 간선 그대로면 조용
    am, ac = CGm.load_claims_full(A2)
    new, lines, hard = CGm.oral_sync(json.load(open(O)), am, ac, 'f' * 16, 'a2.json')
    assert not hard and new and not [l for l in lines if '간선' in l or '새로 걸림' in l or '빠짐' in l], lines
    A, O, A2 = _s1_files(tempfile.mkdtemp(prefix='cgs1_'), v2, old_snap=True)                   # 옛 스냅숏: 모른다고 한 줄
    am, ac = CGm.load_claims_full(A2)
    new, lines, hard = CGm.oral_sync(json.load(open(O)), am, ac, CGm._file_sha(A2), 'a2.json')
    assert not hard and any('옛 스냅숏' in l and '간선' in l for l in lines) and not any('새로 걸림' in l for l in lines), lines


_PAPER_MD = """# Fake paper title

> 원문 PDF 의 글자층(literature.py v0.8.5). 쪽 표지 `[p.인쇄쪽 · PDF 쪽]`.

[p.101 · PDF 1]

Abstract. Mean ﬂow
velocity in the fake lesion group was higher than in controls (0.81 vs 0.62).
This effect remained after adjustment.
Smith et al. reported a similar pattern.

[p.102 · PDF 2]

Methods: we measured the fake index
in two groups of volunteers.
The fake index was repeated twice.
The fake index was repeated twice.
"""


def _paper_dir(d, doi='10.9999/fake-a', sha='abc123def4567890'):
    pd = os.path.join(d, doi.replace('/', '_'))
    os.makedirs(pd, exist_ok=True)
    open(os.path.join(pd, 'paper.md'), 'w', encoding='utf8').write(_PAPER_MD)
    open(os.path.join(pd, 'meta.md'), 'w', encoding='utf8').write('<!-- lit: doi=%s pages=2 blank=0 sha=%s -->\n# Fake paper title\n' % (doi, sha))
    return pd


def t_v1619_paper_md_find_across_lines_and_ligature():
    """④ 1판: paper.md(PDF 글자층)는 한 줄이 물리적 줄이라 줄바꿈에 걸린 구절·합자(ﬂ)를 doc:find 가 못 찾았다.
    쪽 표지가 있는 md 는 locate 와 같은 규칙(쪽 표지로 나눔 · 빈 줄이나 . : 로 끝난 줄에서 문단 · NFKC · 띄어쓰기 무시)으로 읽는다."""
    import tempfile
    pd = _paper_dir(tempfile.mkdtemp(prefix='cgpp_'))
    src = CGm.DocSource(os.path.join(pd, 'paper.md'))
    t = src.resolve('doc:find:mean flow velocity in the fake lesion')                           # 줄바꿈·합자를 넘는 구절
    assert 'higher than in controls' in t and 'Abstract' in t, t
    assert 'Methods' not in t                                                                   # 다른 쪽 문단은 섞이지 않는다
    assert src.mark_of('doc:find:mean flow velocity') == '[p.101 · PDF 1]'
    for bad, why in (('doc:find:the fake lesion group was lower', '없음'), ('doc:find:the fake index was repeated', '모호')):
        try:
            src.resolve(bad); assert False, bad                                                 # 실패 길: 원문에 없는 구절(지어낸 인용)·모호
        except KeyError as e:
            assert why in str(e), e
    assert 'measured the fake index in two groups' in src.resolve('doc:sec:PDF 2')             # 쪽 표지로 절
    md = os.path.join(pd, 'plain.md')                                                           # 쪽 표지 없는 md 는 전처럼(한 줄 = 문단)
    open(md, 'w', encoding='utf8').write('# H\nline one\nline two\n')
    assert CGm.DocSource(md).resolve('doc:find:line two') == 'line two'


def _lit_claims(pd, **top):
    import json
    g = dict({'kind': '문헌', 'doc': '10.9999/fake-a', 'doi': '10.9999/fake-a', 'paper_sha': 'abc123def4567890', 'claims': [
        {'id': 'r1', 'role': 'evidence', 'statement': '가짜 병변 군의 흐름 속도가 높다(0.81 vs 0.62)', 'status': 'proposed', 'origin': 'ai',
         'sites': ['doc:find:velocity in the fake lesion group was higher'], 'keys': ['higher than in controls'], 'depends_on': []},
        {'id': 'm1', 'role': 'main', 'statement': '가짜 지표가 병변을 가른다', 'status': 'proposed', 'origin': 'ai',
         'sites': ['doc:find:This effect remained after adjustment'], 'keys': ['remained'], 'depends_on': [{'id': 'r1', 'type': 'premise'}]}]}, **top)
    p = os.path.join(pd, 'claims.json'); json.dump(g, open(p, 'w'), ensure_ascii=False)
    return p


def t_v1619_lit_graph_checks():
    """④ 1판: kind 문헌 그래프 — doi·meta.md 대조, id 에 : # 금지(나중 lit:<DOI>#<id>), 사람 이름 꼴 [참고], 판정 대기 수, 논문용 [참고] 일부 끔."""
    import json, tempfile
    pd = _paper_dir(tempfile.mkdtemp(prefix='cglg_'))
    p = _lit_claims(pd)
    meta, cl = CGm.load_claims_full(p)
    probs = CGm.lit_graph_problems(meta, cl, p)
    assert not [x for x in probs if not x.startswith('[참고]')], probs                          # 성공 길
    assert any('판정 대기' in x and '2' in x for x in probs), probs
    gp, _ = CGm.mapgraph(cl, io.StringIO(), kind='문헌')
    assert not [x for x in gp if 'caveat 이 없음' in x], gp                                       # 짝이 될 주장만 뽑으므로 끈다
    bad = [('doi 없음', dict(doi=None), 'doi'), ('meta 와 다름', dict(doi='10.9999/other', doc='10.9999/other'), 'meta.md'),
           ('id #', None, '#')]
    for name, top, word in bad:
        g = json.load(open(p))
        if top:
            g.update({k: v for k, v in top.items()})
            if top.get('doi') is None:
                g.pop('doi')
        else:
            g['claims'][0]['id'] = 'r#1'; g['claims'][1]['depends_on'] = [{'id': 'r#1', 'type': 'premise'}]
        q = os.path.join(pd, 'claims.json'); json.dump(g, open(q, 'w'), ensure_ascii=False)
        m2, c2 = CGm.load_claims_full(q)
        pr = CGm.lit_graph_problems(m2, c2, q)
        assert [x for x in pr if not x.startswith('[참고]') and word in x], (name, pr)           # 실패 길
    p = _lit_claims(pd, paper_sha='0000000000000000')
    pr = CGm.lit_graph_problems(*CGm.load_claims_full(p), p)
    assert any(x.startswith('[참고]') and 'paper_sha' in x for x in pr), pr                        # 논문 판이 바뀜
    g = json.load(open(p)); g['claims'][0]['statement'] = 'Kim et al. 이 보고한 결과(Lee, 2021)'
    g['claims'][1]['sites'] = ['doc:find:Smith et al. reported']
    json.dump(g, open(p, 'w'), ensure_ascii=False)
    pr = CGm.lit_graph_problems(*CGm.load_claims_full(p), p)
    nm = [x for x in pr if '사람 이름' in x]
    assert len(nm) == 2 and all(x.startswith('[참고]') for x in nm), pr                          # r1·m1 각각 한 줄


def t_v1619_lit_graph_cli_mapcheck_freeze_stale():
    """④ 1판: paper.md 에 mapgraph·mapcheck·mapfreeze·mapstale 이 그대로 — 원문 문단이 바뀌면 [변경]."""
    import json, tempfile
    pd = _paper_dir(tempfile.mkdtemp(prefix='cglc_'))
    p = _lit_claims(pd); md = os.path.join(pd, 'paper.md'); fz = os.path.join(pd, 'claims_f.json')
    run = lambda *a: subprocess.run([sys.executable, CGm.__file__] + list(a), capture_output=True, text=True)
    r = run('mapgraph', '--claims', p)
    assert r.returncode == 0 and '판정 대기' in r.stdout, r.stdout[-600:]
    r = run('mapcheck', md, '--claims', p)
    assert r.returncode == 0 and '모든 주장의 자리에 찾는 표현이 있음' in r.stdout, r.stdout[-600:]
    r = run('mapfreeze', md, '--claims', p, '-o', fz)
    assert r.returncode == 0 and json.load(open(fz))['kind'] == '문헌', r.stdout[-400:]           # 맨 위 칸 보존
    r = run('mapstale', md, '--claims', fz)
    assert r.returncode == 0 and '[변경]' not in r.stdout and '[경고]' not in r.stdout, r.stdout[-400:]   # doc(DOI) 와 paper.md 이름이 달라도 경고 없음
    open(md, 'w', encoding='utf8').write(_PAPER_MD.replace('This effect remained after adjustment.', 'This effect vanished after adjustment.'))
    r = run('mapstale', md, '--claims', fz)
    assert '[변경] m1' in r.stdout, r.stdout[-600:]
    g = json.load(open(p)); g.pop('doi'); json.dump(g, open(p, 'w'), ensure_ascii=False)
    assert run('mapgraph', '--claims', p).returncode == 1                                        # [필수] 면 종료 1


# ---- v16.20: ④ 2판 전 작은 고침 — 옛 쪽 표지 풀어 적기 · 쪽 머리·꼬리 줄 빼기 · 줄 끝 하이픈 낱말 ----

def _md_src(d, txt, name='paper.md'):
    p = os.path.join(d, name)
    open(p, 'w', encoding='utf8').write(txt)
    return CGm.DocSource(p)


def t_v1620_mark_of_old_marks_spelled_as_pdf():
    """④ 1판 실물(09-30): 보관소의 옛 변환은 쪽 표지가 `[p.N]`(N = PDF 쪽) — mark_of 가 그대로 돌려주면 sources.at 에 인쇄 쪽처럼 적힌다.
    옛 표지는 `PDF N`(인쇄 쪽이 붙은 `[p.N · 인쇄]` 는 `[p.인쇄 · PDF N]`)으로, 새 표지·절 표지는 그대로. doc:sec:PDF 1 이 PDF 10 을 집지 않는다."""
    import tempfile
    d = tempfile.mkdtemp(prefix='cgmk_')
    pages = ''.join('[p.%d]\n\nPage %s body text sentence here.\n\n' % (n, w) for n, w in
                    ((1, 'one'), (2, 'two'), (3, 'three'), (4, 'four'), (5, 'five'), (6, 'six'), (7, 'seven'), (8, 'eight'), (9, 'nine'), (10, 'ten')))
    src = _md_src(d, '# T\n\n' + pages.replace('[p.10]', '[p.10 · 110]'))
    assert src.mark_of('doc:find:Page three body') == 'PDF 3', src.mark_of('doc:find:Page three body')
    assert src.mark_of('doc:find:Page ten body') == '[p.110 · PDF 10]'
    assert 'Page one' in src.resolve('doc:sec:PDF 1') and 'Page ten' not in src.resolve('doc:sec:PDF 1')
    assert 'Page ten' in src.resolve('doc:sec:PDF 10')                                          # 옛 표지도 PDF 쪽으로
    try:
        src.resolve('doc:sec:PDF 11'); assert False                                             # 실패 길: 없는 쪽
    except KeyError as e:
        assert '표지 없음' in str(e), e
    new = _md_src(d, _PAPER_MD, 'new.md')                                                       # 새 표지는 그대로
    assert new.mark_of('doc:find:measured the fake index') == '[p.102 · PDF 2]'
    sec = _md_src(d, '# T\n\n[§ Methods]\n\nWe measured things.\n', 'sec.md')
    assert sec.mark_of('doc:find:We measured') == '[§ Methods]'


def _running_md(n_pages=5, head_from=2):
    """가짜 논문 — 2쪽부터 쪽 머리 한 줄, 모든 쪽 첫 줄 앞에 쪽 꼬리가 붙어 나온다(실물 PDF 글자층 모양: 줄바꿈 없이 본문에 붙음)."""
    out = ['# Fake running', '']
    first = ('alpha', 'beta', 'the', 'gamma', 'that', 'delta', 'this')                         # 쪽마다 다른 본문 첫 낱말(실물처럼)
    for n in range(1, n_pages + 1):
        out += ['[p.%d · PDF %d]' % (n + 40, n), '']
        if n >= head_from:
            out.append('Doe et al. 10.9999/fake-b')
        out.append('Fake Journal of Tests %02d fakejournal.org%s page %d starts with this sentence that' % (n, first[n - 1], n))
        out.append('continues onto the %s line of page %d.' % (('next', 'second', 'following', 'lower', 'last')[(n - 1) % 5], n))
        if n in (2, 4):
            out.append('FIGURE %d' % n)
        out.append('%s closes page %d here.' % (('Omega', 'Sigma', 'Kappa', 'Theta', 'Lambda', 'Zeta', 'Eta')[n - 1], n))
        out.append('')
    return '\n'.join(out)


def t_v1620_running_head_and_foot_removed():
    """④ 1판 실물(09-30): 15쪽 중 14쪽 맨 위에 '저자 et al. DOI' 쪽 머리 한 줄, 쪽 꼬리 '학술지 이름 쪽번호 사이트' 는 줄바꿈 없이 본문 첫 줄에 붙어 나왔다
    → 그 쪽 첫 문단이 쪽 머리·꼬리를 달고 있어 사람 이름 꼴·구절 찾기가 흐려진다. 쪽 절반 이상(3쪽 이상)에서 쪽 위·아래 세 줄 안에 되풀이되는
    줄 앞머리(숫자는 같게 본다)를 뺀다. 절반이 안 되는 줄(FIGURE)·쪽이 적은 md 는 그대로."""
    import tempfile
    d = tempfile.mkdtemp(prefix='cgrh_')
    src = _md_src(d, _running_md())
    allp = '\n'.join(src.paras)
    assert 'Doe et al' not in allp and 'fakejournal' not in allp and 'Fake Journal' not in allp, src.paras
    t = src.resolve('doc:find:page 3 starts with this sentence that continues onto')                 # 붙어 있던 본문은 남는다
    assert t.startswith('the page 3 starts'), t
    assert src.mark_of('doc:find:page 3 starts') == '[p.43 · PDF 3]'
    assert 'FIGURE 2' in allp and 'FIGURE 4' in allp                                            # 실패 길: 2/5 쪽뿐인 줄은 본문
    few = _md_src(d, _running_md(n_pages=3), 'few.md')                                          # 3쪽 — 절반 이상이어도 3쪽이면 뺀다
    assert 'Doe et al' in '\n'.join(few.paras) and 'fakejournal' not in '\n'.join(few.paras), few.paras   # 머리는 2쪽뿐 → 남음
    two = _md_src(d, _running_md(n_pages=2), 'two.md')                                          # 2쪽 — 되풀이를 판단하지 않는다
    assert 'fakejournal' in '\n'.join(two.paras)
    sec = _md_src(d, _running_md().replace('[p.', '[§ S').replace(' · PDF ', ' '), 'sec.md')     # 절 표지 md 는 쪽이 아니다 — 그대로
    assert 'fakejournal' in '\n'.join(sec.paras)


def t_v1620_line_end_hyphen_found_joined():
    """④ 1판 실물(09-30): 줄 끝 하이픈이 18곳('tri-'·'per -' 줄 끝). 찾기 사본에서만 붙인 꼴로도 찾는다 — 돌려주는 문단은 원문 그대로.
    다음 줄이 대문자로 시작하면(ISF-HFC 같은 복합어) 붙이지 않는다."""
    import tempfile
    d = tempfile.mkdtemp(prefix='cghy_')
    md = ('# T\n\n[p.1 · PDF 1]\n\nThe fake tri-\nals used a voxel-\nwise approach and the analysis was per -\nformed twice.\n'
          'Maps of ISF-\nHFC coupling were drawn­\nlater.\n')
    src = _md_src(d, md)
    for q, w in (('fake trials used', 'voxel'), ('voxelwise approach', 'voxel'), ('voxel-wise approach', 'voxel'),
                 ('analysis was performed twice', 'voxel'), ('drawnlater', 'Maps'), ('drawn later', 'Maps'), ('ISF-HFC coupling', 'Maps')):
        assert w in src.resolve('doc:find:' + q), q
    assert 'tri- als' in src.resolve('doc:find:fake trials used')                               # 원문 그대로 돌려준다
    for bad in ('ISFHFC coupling', 'fake trails used'):
        try:
            src.resolve('doc:find:' + bad); assert False, bad                                   # 실패 길: 대문자 뒤는 붙이지 않음·없는 낱말
        except KeyError as e:
            assert '없음' in str(e), e


# ---- v16.21: ④ 2판 — lit_links(우리 주장 ↔ 논문 주장 짝) · litcheck · --lit 합치기(mapgraph·impact) ----

def _lit_store(status='accepted'):
    """가짜 보관소 — 논문 둘(fake-a: 논문 그래프 있음, fake-b: 원문만). 논문 주장 r1(evidence)·m1(main)."""
    import json, tempfile
    store = tempfile.mkdtemp(prefix='cgls_')
    pa = _paper_dir(store)
    p = _lit_claims(pa)
    g = json.load(open(p))
    for c in g['claims']:
        c['status'] = status
    json.dump(g, open(p, 'w'), ensure_ascii=False)
    _paper_dir(store, doi='10.9999/fake-b')
    return store


def _ours(d, links, src_doi='10.9999/fake-a'):
    """우리 가짜 그래프 — main mn ← premise cond ← premise ev · side 는 support 로만 mn 을 받침."""
    import json
    cl = [{'id': 'ev', 'role': 'evidence', 'statement': 'E', 'depends_on': []},
          {'id': 'cond', 'role': 'claim', 'statement': 'C', 'depends_on': [{'id': 'ev', 'type': 'premise'}],
           'sources': [{'kind': '문헌', 'what': src_doi, 'at': 'PDF 1', 'verdict': '부합'}] if src_doi else []},
          {'id': 'side', 'role': 'claim', 'statement': 'S', 'depends_on': [{'id': 'ev', 'type': 'premise'}]},
          {'id': 'mn', 'role': 'main', 'statement': 'M', 'depends_on': [{'id': 'cond', 'type': 'premise'}, {'id': 'side', 'type': 'support'}]}]
    p = os.path.join(d, 'ours.json')
    json.dump({'doc': 'Fake', 'lit_links': links, 'claims': cl}, open(p, 'w'), ensure_ascii=False)
    return p


def _link(ours='cond', doi='10.9999/fake-a', theirs='m1', rel='support', verdict='부합', **kw):
    return dict(ours=ours, doi=doi, theirs=theirs, rel=rel, **({'verdict': verdict} if verdict else {}), **kw)


def t_v1621_lit_merge_success_edges():
    """④ 2판(결정 1·2): lit_links 의 논문 주장만 lit:<DOI>#<id> 로 합친다 — same·support → support 0.7, rebut → rebuttal 0.5, background → context 0.3.
    논문 주장은 자리·keys·간선 없이, 역할은 lit_role. 우리 파일은 바뀌지 않는다. 문제 없으면 빈 목록."""
    import json, tempfile
    store = _lit_store()
    d = tempfile.mkdtemp(prefix='cglm_')
    p = _ours(d, [_link(), _link(ours='mn', theirs='r1', rel='background'), _link(ours='side', theirs='r1', rel='rebut', verdict='반대 방향')])
    before = open(p, encoding='utf8').read()
    meta, cl = CGm.load_claims_full(p)
    for c in cl:
        if c['id'] in ('mn', 'side'):
            c['sources'] = [{'kind': '문헌', 'what': '10.9999/fake-a', 'at': 'PDF 1', 'verdict': '부합'}]
    merged, probs = CGm.lit_merge(meta, cl, store)
    assert probs == [], probs
    by = {c['id']: c for c in merged}
    lm, lr = 'lit:10.9999/fake-a#m1', 'lit:10.9999/fake-a#r1'
    assert by[lm]['lit'] is True and by[lm]['lit_role'] == 'main' and by[lm]['sites'] == [] and by[lm]['depends_on'] == [], by[lm]
    e = {(c['id'], up): (t, w) for c in merged for up, t, w in CGm._edges([c])[c['id']]}
    assert e[('cond', lm)] == ('support', 0.7) and e[('mn', lr)] == ('context', 0.3) and e[('side', lr)] == ('rebuttal', 0.5), e
    assert len([c for c in merged if c.get('lit')]) == 2                                          # r1 은 두 짝이어도 한 번
    assert open(p, encoding='utf8').read() == before and not any(up.startswith('lit:') for c in cl for up, _, _ in CGm._edges([c])[c['id']])
    assert CGm.lit_merge({'doc': 'x'}, cl, store) == (cl, [])                                     # lit_links 없으면 그대로


def t_v1621_litcheck_hard_errors():
    """④ 2판: [필수] — ours 없음 · rel·verdict 값 · DOI 모양 · 보관소에 없는 DOI · 논문 claims.json 없음 · 다른 논문의 그래프 · theirs 없음 · 목록 아님."""
    import json, tempfile
    store = _lit_store()
    d = tempfile.mkdtemp(prefix='cglh_')
    cases = [(_link(ours='nope'), 'ours'), (_link(rel='same-ish'), 'rel'), (_link(verdict='맞음'), 'verdict'), (_link(doi='fake'), 'DOI 모양'),
             (_link(doi='10.9999/fake-z'), '보관소에 없다'), (_link(doi='10.9999/fake-b'), 'claims.json 이 없다'), (_link(theirs='m9'), '"m9" 가 없다')]
    for link, word in cases:
        meta, cl = CGm.load_claims_full(_ours(d, [link]))
        _, probs = CGm.lit_merge(meta, cl, store)
        hard = [x for x in probs if not x.startswith('[참고]')]
        assert len(hard) == 1 and word in hard[0], (word, probs)
    pb = os.path.join(store, '10.9999_fake-b', 'claims.json')                                    # fake-b 폴더에 fake-a 그래프를 둠
    json.dump({'kind': '문헌', 'doi': '10.9999/fake-a', 'claims': [{'id': 'm1', 'statement': 'x'}]}, open(pb, 'w'))
    meta, cl = CGm.load_claims_full(_ours(d, [_link(doi='10.9999/fake-b')]))
    assert any('문헌 그래프가 아니다' in x for x in CGm.lit_merge(meta, cl, store)[1])
    assert CGm.lit_merge({'lit_links': {'a': 1}}, [], store)[1][0].startswith('lit_links 는 목록')


def t_v1621_unjudged_lit_backing_main_chain_is_hard():
    """④ 2판(결정 5): 판정 전(proposed) 논문 주장이 same·support 로 main 의 전제(premise) 사슬 안 주장을 받치면 [필수].
    성공 길: 판정됨(accepted) · 사슬 밖(support 로만 main 에 닿는 side) · rebut/background 짝은 [필수] 아님."""
    import tempfile
    d = tempfile.mkdtemp(prefix='cglu_')
    sp = _lit_store(status='proposed')
    for link, hard_expected in ((_link(), True), (_link(ours='mn', rel='same'), True), (_link(ours='side'), False),
                                (_link(rel='rebut', verdict='반대 방향'), False), (_link(rel='background'), False)):
        meta, cl = CGm.load_claims_full(_ours(d, [link]))
        _, probs = CGm.lit_merge(meta, cl, sp)
        hard = [x for x in probs if not x.startswith('[참고]')]
        assert bool(hard) == hard_expected and all('판정 전' in x and 'main(mn)' in x for x in hard), (link, probs)
    meta, cl = CGm.load_claims_full(_ours(d, [_link()]))
    assert not [x for x in CGm.lit_merge(meta, cl, _lit_store())[1] if not x.startswith('[참고]')]   # accepted 면 조용


def t_v1621_lit_notes():
    """④ 2판 [참고]: verdict 없음(by 표시) · 우리 sources 에 같은 DOI 없음 · 같은 짝 두 번 · 논문 주장 사람 이름 꼴 · 철회된 논문 주장."""
    import json, tempfile
    store = _lit_store()
    d = tempfile.mkdtemp(prefix='cgln_')
    meta, cl = CGm.load_claims_full(_ours(d, [_link(verdict=None, by='AI 제안'), _link(verdict=None, by='AI 제안')], src_doi=None))
    _, probs = CGm.lit_merge(meta, cl, store)
    assert all(x.startswith('[참고]') for x in probs), probs
    assert any('verdict 없음' in x and 'AI 제안' in x for x in probs) and any('sources 에 같은 DOI 가 없다' in x for x in probs)
    assert any('같은 짝이 두 번' in x for x in probs), probs
    pc = os.path.join(store, '10.9999_fake-a', 'claims.json')
    g = json.load(open(pc)); g['claims'][1]['statement'] = 'Kim et al. 이 보고(2021)'; g['claims'][1]['status'] = 'superseded'
    json.dump(g, open(pc, 'w'), ensure_ascii=False)
    meta, cl = CGm.load_claims_full(_ours(d, [_link()]))
    probs = CGm.lit_merge(meta, cl, store)[1]
    assert any('사람 이름 꼴' in x for x in probs) and any('철회된' in x for x in probs) and all(x.startswith('[참고]') for x in probs), probs


def t_v1621_cli_litcheck_mapgraph_impact_lit():
    """④ 2판 CLI: litcheck(종료 0/1) · mapgraph --lit(논문 주장이 순서에, [필수] 면 종료 1) · --lit 없으면 한 줄 안내 ·
    impact lit:<DOI>#<id> --lit 이 우리 하류를 보인다 · --oral 과 같이 쓰면 종료 2 · add 가 lit_links 를 지키고 판정 전이면 litcheck 종료 1."""
    import json, tempfile
    store = _lit_store()
    d = tempfile.mkdtemp(prefix='cglc_')
    p = _ours(d, [_link()])
    run = lambda *a: subprocess.run([sys.executable, CGm.__file__] + list(a), capture_output=True, text=True)
    r = run('litcheck', '--claims', p, '--store', store)
    assert r.returncode == 0 and '짝 1 · 논문 주장 1(판정 대기 0)' in r.stdout and '문제 없음' in r.stdout, r.stdout
    r = run('mapgraph', '--claims', p, '--lit', store)
    assert r.returncode == 0 and 'lit:10.9999/fake-a#m1' in r.stdout and '합침' in r.stdout, r.stdout[-800:]
    r = run('mapgraph', '--claims', p)
    assert r.returncode == 0 and 'lit_links 1개' in r.stdout and 'lit:10.9999' not in r.stdout, r.stdout[-600:]
    r = run('impact', '--claims', p, '--lit', store, 'lit:10.9999/fake-a#m1')
    assert r.returncode == 0 and 'cond' in r.stdout and 'mn' in r.stdout and '[필수]' in r.stdout, r.stdout
    r = run('mapgraph', '--claims', p, '--lit', store, '--oral', p, '--author', p)
    assert r.returncode == 2, r.stdout
    q = os.path.join(d, 'ours2.json')
    r = run('add', '--claims', p, '--id', 'nw', '--statement', 'N', '-o', q)
    assert r.returncode == 0 and json.load(open(q))['lit_links'] == [_link()], r.stdout            # 맨 위 칸 보존
    sp = _lit_store(status='proposed')                                                            # 실패 길: 판정 전 → [필수] 종료 1
    r = run('litcheck', '--claims', p, '--store', sp)
    assert r.returncode == 1 and '[필수]' in r.stdout and '판정 대기 1' in r.stdout, r.stdout
    r = run('mapgraph', '--claims', p, '--lit', sp)
    assert r.returncode == 1 and '판정 전' in r.stdout, r.stdout[-600:]
    open(os.path.join(d, 'none.json'), 'w').write(json.dumps({'claims': []}))
    r = run('litcheck', '--claims', os.path.join(d, 'none.json'), '--store', store)
    assert r.returncode == 0 and 'lit_links 없음' in r.stdout


if __name__ == '__main__':
    sys.exit(1 if run() else 0)
