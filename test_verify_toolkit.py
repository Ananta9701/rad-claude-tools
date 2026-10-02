#!/usr/bin/env python3
"""test_verify_toolkit.py — verify_toolkit.py 자동 테스트 (v1.1 신설).

    python3 test_verify_toolkit.py      # 실패 0, SKIP 0 이어야 함

실물 원고는 프로젝트에 없으므로 fixture 는 python-docx 로 만든다.
XML 레벨 검사(track changes / highlight)는 python-docx 가 못 만드는 요소라
zip 안의 document.xml 을 직접 고쳐 만든다.
"""
import io, os, re, shutil, sys, traceback, zipfile
sys.dont_write_bytecode = True   # /mnt/project 는 대화창 안에서 쓰기 가능 — __pycache__ 를 남기지 않는다 (v2.3.1)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import verify_toolkit as V
from docx import Document


def _manifest_version(fname):
    """같은 폴더의 TOOLS_MANIFEST.md 에 적힌 판. 없으면 None (Z1 이후: EXPECT_VERSION 만 맞추고 manifest 를 안 올린 사고 방지)."""
    import os, re
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'TOOLS_MANIFEST.md')
    if not os.path.exists(p):
        return None
    m = re.search(r'\| `%s` \| v([0-9.]+)' % re.escape(fname), open(p, encoding='utf8').read())
    return m.group(1) if m else None

EXPECT_VERSION = '1.3.9'
TMP = os.environ.get('VT_TMP', '/tmp/vt_test')
shutil.rmtree(TMP, ignore_errors=True)
os.makedirs(TMP, exist_ok=True)


def _docx(name, paras):
    d = Document()
    for t in paras:
        d.add_paragraph(t)
    p = os.path.join(TMP, name)
    d.save(p)
    return p


def _patch_xml(src, name, fn):
    """document.xml 을 fn 으로 바꾼 사본을 만든다."""
    out = os.path.join(TMP, name)
    zin = zipfile.ZipFile(src)
    zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.namelist():
        data = zin.read(item)
        if item == 'word/document.xml':
            data = fn(data.decode('utf-8')).encode('utf-8')
        zout.writestr(item, data)
    zin.close(); zout.close()
    return out


def _quiet(fn, *a, **k):
    buf, old = io.StringIO(), sys.stdout
    sys.stdout = buf
    try:
        return fn(*a, **k), buf.getvalue()
    finally:
        sys.stdout = old


BODY = ['Title', 'Objective: a b c d e. Key Words: x', 'INTRODUCTION',
        'First [1]. Then [2,3] and [4-6].', 'Again [1] and [7].',
        'AI Disclosure', 'none', 'REFERENCES',
        '1. Alpha A. J 2020.', '2. Beta B. J 2020.', '3. Gamma C. J 2020.', '4. Delta D. J 2020.',
        '5. Eps E. J 2020.', '6. Zeta F. J 2020.', '7. Eta G. J 2020.']


def t_version_matches_manifest():
    assert getattr(V, '__version__', None) == EXPECT_VERSION, (getattr(V, '__version__', None), EXPECT_VERSION)
    mv = _manifest_version('verify_toolkit.py')
    assert mv is None or mv == EXPECT_VERSION, ('TOOLS_MANIFEST.md 의 판', mv, '코드', EXPECT_VERSION)


def t_citations_clean():
    r, _ = _quiet(V.check_citations, _docx('c_ok.docx', BODY))
    assert r['used'] == 7 and r['listed'] == 7 and r['order_violations'] == [] and r['gaps'] == [], r


def t_citations_order_violation():
    paras = list(BODY); paras[3] = 'First [2]. Then [1,3] and [4-6].'
    r, _ = _quiet(V.check_citations, _docx('c_bad.docx', paras))
    assert r['order_violations'] == [2, 1], r


def t_citations_gap_and_count_mismatch():
    paras = list(BODY); paras[3] = 'First [1]. Then [2,3] and [5-6].'; paras[4] = 'Again [1] and [7].'
    r, _ = _quiet(V.check_citations, _docx('c_gap.docx', paras))
    assert r['gaps'] == [4] and r['used'] == 6 and r['listed'] == 7, r


def t_citations_refs_counted_after_marker_only():
    # v1.1: 본문 안의 "1. Something" 번호 목록은 REFERENCES 수에 안 들어간다
    paras = list(BODY); paras.insert(4, '1. Numbered list item in body.')
    r, _ = _quiet(V.check_citations, _docx('c_body_list.docx', paras))
    assert r['listed'] == 7, r


def t_citations_korean_memo_ignored():
    paras = list(BODY); paras.insert(4, '〔수정〕 [99] 는 예시 — 한글 메모')
    r, _ = _quiet(V.check_citations, _docx('c_kor.docx', paras))
    assert r['used'] == 7 and r['gaps'] == [], r


def t_v137_citations_en_dash_range():
    # 코드 리뷰 ⑫: [4–6](en dash, Word 자동 고침)을 인용으로 읽지 못해 4·5·6 이 빠지고, 뒤 [7] 이 순서 위반·결번으로 나왔다
    paras = list(BODY); paras[3] = 'First [1]. Then [2,3] and [4\u20136].'
    r, _ = _quiet(V.check_citations, _docx('c_endash.docx', paras))
    assert r['used'] == 7 and r['order_violations'] == [] and r['gaps'] == [], r          # 성공 길: en dash 범위를 편다
    paras[3] = 'First [1]. Then [2,3] and [5 \u2013 6].'                                   # 실패 길: 범위가 4 를 건너뛰면 결번은 그대로 잡는다
    r, _ = _quiet(V.check_citations, _docx('c_endash_gap.docx', paras))
    assert r['gaps'] == [4], r
    paras[3] = 'First [1]. Then [3,2] and [4\u20136].'                                    # 실패 길: 진짜 순서 위반은 그대로
    r, _ = _quiet(V.check_citations, _docx('c_endash_bad.docx', paras))
    assert r['order_violations'] == [3, 2], r


def t_citations_missing_marker_returns_none():
    r, out = _quiet(V.check_citations, _docx('c_nomark.docx', ['no markers here [1]']))
    assert r is None and 'CONFIG' in out


def t_pvalue_threshold_passes_actual_fails():
    p = _docx('p1.docx', ['Test P < 0.05 threshold and P = 0.02 actual.'])
    r, _ = _quiet(V.check_pvalue_format, p)
    assert r == ['P = 0.02'], r


def t_pvalue_three_decimals_ok():
    p = _docx('p2.docx', ['P = 0.023 and P = 0.001 and P < 0.001 and P<0.0001.'])
    r, _ = _quiet(V.check_pvalue_format, p)
    assert r == [], r


def t_pvalue_lt_bonferroni_not_flagged():
    # v1.2 (#8): `<` 뒤 값은 전부 임계값으로 보고 자릿수 검사 제외
    p = _docx('p3.docx', ['Bonferroni P < 0.0045 and P < 0.02; actual P = 0.02.'])
    r, out = _quiet(V.check_pvalue_format, p)
    assert r == ['P = 0.02'], r
    assert '0.0045' in out


def t_citations_first_violation_reported():
    # [1]..[3] 정상, 4번째 등장이 [5] → 첫 위반은 4번째, 그 뒤 [4] 는 연쇄
    paras = list(BODY); paras[3] = 'First [1]. Then [2,3] and [5].'; paras[4] = 'Again [4] and [6-7].'
    r, out = _quiet(V.check_citations, _docx('c_first.docx', paras))
    fb = r['first_violation']
    assert fb['nth'] == 4 and fb['found'] == 5 and fb['expected'] == 4 and '[5]' in fb['para'], fb
    assert r['order_violations'] == [5, 4] and '첫 위반' in out


def t_citations_no_violation_no_first():
    r, _ = _quiet(V.check_citations, _docx('c_ok2.docx', BODY))
    assert r['first_violation'] is None


def _docx_with_image(name, paras_before, paras_after, add_orphan=False):
    from PIL import Image
    d = Document()
    for t in paras_before:
        d.add_paragraph(t)
    img = os.path.join(TMP, name + '.png'); Image.new('RGB', (321, 123)).save(img)
    d.add_picture(img)
    for t in paras_after:
        d.add_paragraph(t)
    p = os.path.join(TMP, name + '.docx'); d.save(p)
    if not add_orphan:
        return p
    # rels 에만 있고 본문 참조 없는 이미지 추가
    out = os.path.join(TMP, name + '_orphan.docx')
    zin = zipfile.ZipFile(p); zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    orphan = os.path.join(TMP, name + '_o.png'); Image.new('RGB', (50, 60)).save(orphan)
    for item in zin.namelist():
        data = zin.read(item)
        if item == 'word/_rels/document.xml.rels':
            data = data.decode('utf-8').replace('</Relationships>',
                '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/orphan.png"/></Relationships>').encode('utf-8')
        zout.writestr(item, data)
    zout.write(orphan, 'word/media/orphan.png')
    zin.close(); zout.close()
    return out


def t_list_images_basic():
    p = _docx_with_image('img1', ['Intro.', 'Figure 2. Quantitative map.'], ['Next para.'])
    r, out = _quiet(V.list_images, p)
    assert len(r) == 1 and r[0]['pixels'] == (321, 123) and r[0]['referenced'], r
    assert r[0]['caption_before'] == 'Figure 2. Quantitative map.' and r[0]['caption_after'] == 'Next para.', r


def t_list_images_orphan():
    p = _docx_with_image('img2', ['A'], ['B'], add_orphan=True)
    r, out = _quiet(V.list_images, p)
    orphans = [x for x in r if not x['referenced']]
    assert len(orphans) == 1 and orphans[0]['rId'] == 'rId99' and orphans[0]['pixels'] == (50, 60), r
    assert '고아' in out


def t_table_notation_ambiguous_and_mixed():
    d = Document()
    d.add_paragraph('Body uses \u221222.9 and \u22123.5 (\u221222.9 to \u22123.5).')
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text = '(-22.984--3.487)'
    t.rows[0].cells[1].text = '(-16.283-2.139)'
    t.rows[0].cells[2].text = '(1.2-3.4)'
    p = os.path.join(TMP, 'tn.docx'); d.save(p)
    r, out = _quiet(V.check_table_notation, p)
    assert [a[1] for a in r['ambiguous']] == ['(-22.984--3.487)', '(-16.283-2.139)'], r
    assert r['minus_body'] == 4 and r['hyphen_neg_body'] == 0 and r['minus_table'] == 0 and r['hyphen_neg_table'] == 3, r
    assert '혼용' in out


REF_BODY = ['Title', 'Objective: a b. Key Words: x', 'INTRODUCTION',
            'Alpha [1]. Beta [2,3]. Gamma [4-6].', 'Delta [7] and [2].', 'AI Disclosure', 'x', 'REFERENCES',
            '1. R one.', '2. R two.', '3. R three.', '4. R four.', '5. R five.', '6. R six.', '7. R seven.']


def t_renumber_remove_and_reorder():
    # [3] 삭제 대상: 본문에서 먼저 지운 상태. [2] 첫 등장이 [4-6] 뒤로 옮겨진 원고
    paras = list(REF_BODY); paras[3] = 'Alpha [1]. Gamma [4-6].'; paras[4] = 'Delta [7] and [2].'
    src = _docx('rn_src.docx', paras)
    out = os.path.join(TMP, 'rn_out.docx')
    r, log = _quiet(V.renumber_references, src, remove=[3], out=out)
    assert r['ok'], (r, log)
    assert r['mapping'] == {1: 1, 2: 6, 3: None, 4: 2, 5: 3, 6: 4, 7: 5}, r['mapping']
    txt = [p.text for p in Document(out).paragraphs]
    assert txt[3] == 'Alpha [1]. Gamma [2-4].' and txt[4] == 'Delta [5] and [6].', txt[3:5]
    refs = [t for t in txt if re.match(r'^\d+\. R', t)]
    assert refs == ['1. R one.', '2. R four.', '3. R five.', '4. R six.', '5. R seven.', '6. R two.'], refs
    assert r['recheck']['order_violations'] == [] and r['recheck']['listed'] == 6, r['recheck']
    import json
    m = json.load(open(r['map_json'], encoding='utf8'))
    assert m['map']['3'] is None and m['map']['2'] == 6 and m['removed'] == [3], m


def t_renumber_compresses_runs():
    paras = list(REF_BODY); paras[3] = 'Alpha [1,2,3,5].'; paras[4] = 'Delta [4] [6,7].'
    src = _docx('rn2_src.docx', paras); out = os.path.join(TMP, 'rn2_out.docx')
    r, _ = _quiet(V.renumber_references, src, remove=[], out=out)
    txt = [p.text for p in Document(out).paragraphs]
    assert txt[3] == 'Alpha [1-4].' and txt[4] == 'Delta [5] [6,7].', txt[3:5]


def t_renumber_aborts_when_removed_still_cited():
    src = _docx('rn3_src.docx', REF_BODY)
    r, log = _quiet(V.renumber_references, src, remove=[3], out=os.path.join(TMP, 'rn3_out.docx'))
    assert not r['ok'] and any(c == 'c' for c, _, _ in r['where']), r
    assert not os.path.exists(os.path.join(TMP, 'rn3_out.docx'))


def t_renumber_aborts_on_split_run():
    src = _docx('rn4_src.docx', REF_BODY)
    split = _patch_xml(src, 'rn4_split.docx',
                       lambda s: s.replace('<w:t>Alpha [1]. Beta [2,3]. Gamma [4-6].</w:t>',
                                           '<w:t xml:space="preserve">Alpha [</w:t></w:r><w:r><w:t>1]. Beta [2,3]. Gamma [4-6].</w:t>'))
    r, log = _quiet(V.renumber_references, split, remove=[], out=os.path.join(TMP, 'rn4_out.docx'))
    assert not r['ok'] and any(c == 'a' for c, _, _ in r['where']), r
    assert '보수 모드 중단' in log


def t_renumber_cli_remove_space_and_comma():
    # v1.2.1: "--remove 13 26" 이 26 을 버리던 결함
    import subprocess
    paras = list(REF_BODY); paras[3] = 'Alpha [1]. Gamma [4-6].'; paras[4] = 'Delta [7].'   # [2][3] 미인용
    src = _docx('rn6_src.docx', paras)
    for args in (['--remove', '2', '3'], ['--remove', '2,3']):
        r = subprocess.run([sys.executable, V.__file__, 'renumber', src] + args + ['--dry-run'], capture_output=True, text=True)
        assert r.returncode == 0 and '삭제 [2, 3]' in r.stdout and '목록 7 → 5' in r.stdout, (args, r.stdout[-300:])


def t_renumber_dry_run_writes_nothing():
    src = _docx('rn5_src.docx', REF_BODY); out = os.path.join(TMP, 'rn5_out.docx')
    r, _ = _quiet(V.renumber_references, src, remove=[], out=out, dry_run=True)
    assert r['ok'] and r.get('dry_run') and not os.path.exists(out)


def t_pvalue_scans_tables():
    # v1.3 (저자 260910 #4): 표 안 P 값을 놓치던 결함
    d = Document()
    d.add_paragraph('Body has P = 0.023 only.')
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = 'Region (r/p)'
    t.rows[0].cells[1].text = 'P = 0.02'
    t.rows[1].cells[0].text = 'NAT'
    t.rows[1].cells[1].text = 'P = 0.0451'
    p = os.path.join(TMP, 'pv_tbl.docx'); d.save(p)
    r, out = _quiet(V.check_pvalue_format, p)
    assert sorted(r) == ['P = 0.02', 'P = 0.0451'], r   # 표 안 두 표기 모두 잡힘(본문 0.023 은 3자리라 통과)
    assert list(r) == r.violations


def t_pvalue_gt_counted_as_threshold():
    p = _docx('pv_gt.docx', ['Non-significant P > 0.05 and P ≥ 0.05; actual P = 0.02.'])
    r, out = _quiet(V.check_pvalue_format, p)
    assert r == ['P = 0.02'] and 'P > 0.05' in out and 'P ≥ 0.05' in out, (r, out)


def t_pvalue_lowercase_labels():
    d = Document()
    d.add_paragraph('correlation coefficients (r) and corresponding p-values were listed.')
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = 'dLES (r/p)'
    t.rows[0].cells[1].text = 'pLES (r/p)'
    p = os.path.join(TMP, 'pv_low.docx'); d.save(p)
    r, out = _quiet(V.check_pvalue_format, p)
    low = dict(r.lowercase_p)
    assert low.get('(r/p)') == 2 and low.get('p-values') == 1, r.lowercase_p
    assert r == [] and '소문자 p 라벨' in out          # 위반 목록에는 안 올림


def t_pvalue_uppercase_not_counted_as_lowercase():
    p = _docx('pv_up.docx', ['P = 0.023 and P-value reported.'])
    r, _ = _quiet(V.check_pvalue_format, p)
    assert r.lowercase_p == [], r.lowercase_p


def t_wordcount_body_and_abstract():
    r, _ = _quiet(V.check_word_count, _docx('w1.docx', BODY))
    # body = paras[i0:i1] — 마커 'INTRODUCTION'(1) + 'First [1]. Then [2,3] and [4-6].'(6) + 'Again [1] and [7].'(4) = 11
    assert r['body'] == 11 and r['abstract'] == 6, r


def t_wordcount_excludes_korean_and_numbered():
    paras = list(BODY); paras.insert(4, '한글 메모 다섯 단어 입니다'); paras.insert(4, '1. Numbered A.')
    r, _ = _quiet(V.check_word_count, _docx('w2.docx', paras))
    assert r['body'] == 11, r


def t_wordcount_excludes_tables_and_warns_on_missing_marker():
    # v1.3.1: 표는 세지 않는다(KJR). 마커 없으면 세지 않고 경고
    from docx import Document as D
    d = D()
    for t in BODY:
        d.add_paragraph(t)
    tb = d.add_table(rows=1, cols=1); tb.rows[0].cells[0].text = 'table words that must not be counted here'
    p = os.path.join(TMP, 'wc_tbl.docx'); d.save(p)
    r, _ = _quiet(V.check_word_count, p)
    assert r['body'] == 11, r
    r, out = _quiet(V.check_word_count, p, body_end='NoSuchMarker')
    assert r['body'] is None and '구간 마커를 찾지 못했습니다' in out


def t_wordcount_cli_from_to():
    import subprocess
    p = _docx('wc_cli.docx', BODY)
    r = subprocess.run([sys.executable, V.__file__, 'wordcount', p, '--from', 'INTRODUCTION', '--to', 'AI Disclosure'],
                       capture_output=True, text=True)
    assert r.returncode == 0 and "구간: 'INTRODUCTION' ~ 'AI Disclosure'" in r.stdout, r.stdout


def t_citations_ignores_tables():
    # v1.3.1: 표 각주 인용은 등장순 대상 아님(KJR §4/§5)
    from docx import Document as D
    d = D()
    for t in BODY:
        d.add_paragraph(t)
    tb = d.add_table(rows=1, cols=1); tb.rows[0].cells[0].text = 'Table footnote citing [99].'
    p = os.path.join(TMP, 'cit_tbl.docx'); d.save(p)
    r, _ = _quiet(V.check_citations, p)
    assert r['used'] == 7 and r['gaps'] == [] and r['order_violations'] == [], r


def t_highlight_scan_covers_tables():
    # scan_highlight_and_korean 은 document.xml 을 직접 읽으므로 표 안도 이미 본다(저자 요청 검증)
    from docx import Document as D
    d = D(); d.add_paragraph('clean')
    tb = d.add_table(rows=1, cols=1); tb.rows[0].cells[0].text = '표 안 한글 메모'
    p = os.path.join(TMP, 'hl_tbl.docx'); d.save(p)
    r, _ = _quiet(V.scan_highlight_and_korean, p)
    assert r['korean'] == 1, r


def t_highlight_and_korean_clean():
    r, _ = _quiet(V.scan_highlight_and_korean, _docx('h_ok.docx', ['Clean text only.']))
    assert r == {'highlight': 0, 'korean': 0}, r


def t_highlight_detected_and_stripped():
    src = _docx('h_src.docx', ['Marked text.'])
    hl = _patch_xml(src, 'h_hl.docx',
                    lambda s: s.replace('<w:t>Marked text.</w:t>',
                                        '<w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>Marked text.</w:t>'))
    r, _ = _quiet(V.scan_highlight_and_korean, hl)
    assert r['highlight'] == 1, r
    out = os.path.join(TMP, 'h_stripped.docx')
    _quiet(V.strip_all_highlight, hl, out)
    r2, _ = _quiet(V.scan_highlight_and_korean, out)
    assert r2['highlight'] == 0, r2


def t_korean_residue_detected():
    r, _ = _quiet(V.scan_highlight_and_korean, _docx('h_kor.docx', ['English', '한글 잔존']))
    assert r['korean'] == 1, r


def _with_track_changes(name):
    src = _docx(name + '_src.docx', ['Base sentence.'])
    ins = ('<w:ins w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t>added English</w:t></w:r></w:ins>'
           '<w:ins w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t>한글 지시문</w:t></w:r></w:ins>'
           '<w:del w:id="3" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>removed</w:delText></w:r></w:del>')
    return _patch_xml(src, name + '.docx', lambda s: s.replace('<w:t>Base sentence.</w:t></w:r>',
                                                              '<w:t>Base sentence.</w:t></w:r>' + ins))


def t_track_changes_scanned():
    r, out = _quiet(V.scan_track_changes, _with_track_changes('tc1'))
    assert r['ins'] == 2 and r['del'] == 1 and r['korean_ins'] == ['한글 지시문'], r
    assert '추적 변경 상태' in out


def t_track_changes_no_changes():
    r, out = _quiet(V.scan_track_changes, _docx('tc0.docx', ['Plain.']))
    assert r['ins'] == 0 and r['del'] == 0 and '추적 변경 상태' not in out


def t_accept_or_reject_changes():
    src = _with_track_changes('tc2')
    out = os.path.join(TMP, 'tc2_clean.docx')
    _quiet(V.accept_or_reject_changes, src, out)
    r, _ = _quiet(V.scan_track_changes, out)
    assert r['ins'] == 0 and r['del'] == 0, r
    txt = ' '.join(p.text for p in Document(out).paragraphs)
    assert 'added English' in txt and '한글 지시문' not in txt and 'removed' not in txt, txt


_TC = ' w:author="A" w:date="2026-01-01T00:00:00Z"'


def _with_mark_changes(name):
    """문단 표지에 붙는 스스로 닫는 추적 변경 표지(<w:ins/>·<w:del/> in rPr) + 보통 ins/del. v1.3.5 재현용(코드 리뷰 09-28)."""
    src = _docx(name + '_src.docx', ['Keep one.', 'Keep two.', 'Keep 셋.', 'Keep four.'])
    def fn(s):
        s = re.sub(r'<w:p>(<w:r><w:t>Keep one\.)', r'<w:p><w:pPr><w:rPr><w:ins w:id="10"%s/></w:rPr></w:pPr>\1' % _TC, s)
        s = re.sub(r'<w:p>(<w:r><w:t>Keep two\.)', r'<w:p><w:pPr><w:rPr><w:del w:id="11"%s/></w:rPr></w:pPr>\1' % _TC, s)
        s = s.replace('<w:t>Keep four.</w:t></w:r>', '<w:t>Keep four.</w:t></w:r>'
                      '<w:ins w:id="12"%s><w:r><w:t> new text</w:t></w:r></w:ins>'
                      '<w:del w:id="13"%s><w:r><w:delText> old text</w:delText></w:r></w:del>' % (_TC, _TC))
        assert s.count('w:id="1') == 4, s   # 표지 넷이 다 들어갔는지(fixture 자체 점검)
        return s
    return _patch_xml(src, name + '.docx', fn)


def t_v135_track_changes_self_closing_marks():
    # 코드 리뷰 09-28 [결함]: 스스로 닫는 <w:ins/>·<w:del/> 표지를 여는 태그로 잡아 다음 </w:ins>·</w:del> 까지 한 덩어리로 —
    # XML 이 깨지거나 남길 문단이 지워지고, 사이에 한글이 있으면 덩어리째 지워졌다
    import xml.dom.minidom
    src = _with_mark_changes('tc5')
    r, _ = _quiet(V.scan_track_changes, src)
    assert r['ins'] == 1 and r['del'] == 1 and r.get('marks') == 2, r        # 보통 변경만 세고 문단 표지는 따로
    out = os.path.join(TMP, 'tc5_clean.docx')
    _quiet(V.accept_or_reject_changes, src, out)
    xml.dom.minidom.parseString(zipfile.ZipFile(out).read('word/document.xml'))   # 깨지지 않았다
    txt = ' | '.join(p.text for p in Document(out).paragraphs)
    for keep in ('Keep one.', 'Keep two.', 'Keep 셋.', 'Keep four. new text'):
        assert keep in txt, (keep, txt)
    assert 'old text' not in txt, txt
    r2, _ = _quiet(V.scan_track_changes, out)
    assert r2['ins'] == 0 and r2['del'] == 0 and r2.get('marks') == 0, r2


def t_v135_track_changes_broken_input_stops():
    # 실패 길: 결과 XML 이 온전하지 않으면 파일을 쓰지 않고 멈춘다(깨진 docx 를 조용히 내지 않는다)
    src = _patch_xml(_docx('tc6_src.docx', ['Base.']), 'tc6.docx', lambda s: s.replace('</w:body>', '<w:p></w:body>'))
    out = os.path.join(TMP, 'tc6_clean.docx')
    try:
        _quiet(V.accept_or_reject_changes, src, out); assert False, '멈추지 않았다'
    except SystemExit as e:
        assert '온전하지 않다' in str(e), e
    assert not os.path.exists(out), '깨진 결과 파일이 남았다'


def t_v136_track_changes_moves_stop():
    # 사용자 09-28: 이동 표지(moveFrom·moveTo)는 처리하지 않고 멈춰 "Word 에서 직접 적용" 을 알린다
    src = _patch_xml(_docx('tc7_src.docx', ['Stay.', 'Moved here.']), 'tc7.docx', lambda s: s.replace(
        '<w:t>Stay.</w:t></w:r>', '<w:t>Stay.</w:t></w:r><w:moveFrom w:id="20"%s><w:r><w:t> gone</w:t></w:r></w:moveFrom>' % _TC).replace(
        '<w:t>Moved here.</w:t></w:r>', '<w:t>Moved here.</w:t></w:r><w:moveTo w:id="21"%s><w:r><w:t> gone</w:t></w:r></w:moveTo>' % _TC))
    r, out_txt = _quiet(V.scan_track_changes, src)
    assert r.get('moves') == 2 and '이동' in out_txt, (r, out_txt)
    out = os.path.join(TMP, 'tc7_clean.docx')
    try:
        _quiet(V.accept_or_reject_changes, src, out); assert False, '멈추지 않았다'
    except SystemExit as e:
        assert 'Word' in str(e) and '이동' in str(e), e
    assert not os.path.exists(out)
    # 성공 길: 이동이 없는 문서는 전처럼 처리
    src2 = _with_track_changes('tc8'); out2 = os.path.join(TMP, 'tc8_clean.docx')
    _quiet(V.accept_or_reject_changes, src2, out2)
    r2, _ = _quiet(V.scan_track_changes, out2)
    assert os.path.exists(out2) and r2['ins'] == 0 and r2['del'] == 0 and r2.get('moves') == 0, r2


def t_fix_zoom_bug():
    src = _docx('z_src.docx', ['x'])
    out = os.path.join(TMP, 'z_out.docx')
    zin = zipfile.ZipFile(src); zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.namelist():
        data = zin.read(item)
        if item == 'word/settings.xml':
            s = data.decode('utf-8')
            s = re.sub(r'<w:zoom[^/]*/>', '', s)
            s = s.replace('<w:settings', '<w:settings', 1)
            s = re.sub(r'(<w:settings[^>]*>)', r'\1<w:zoom w:val="bestFit"/>', s, count=1)
            data = s.encode('utf-8')
        zout.writestr(item, data)
    zin.close(); zout.close()
    fixed = out.replace('.docx', '_fixed.docx')
    _quiet(V.fix_zoom_bug, out, fixed)
    s = zipfile.ZipFile(fixed).read('word/settings.xml').decode('utf-8')
    assert 'percent' in s or 'bestFit' not in s, s[:300]


def t_dedupe_bookmarks():
    src = _docx('b_src.docx', ['x'])
    bm = ('<w:bookmarkStart w:id="0" w:name="_GoBack"/><w:bookmarkEnd w:id="0"/>'
          '<w:bookmarkStart w:id="1" w:name="_GoBack"/><w:bookmarkEnd w:id="1"/>')
    dup = _patch_xml(src, 'b_dup.docx', lambda s: s.replace('<w:body>', '<w:body>' + bm, 1))
    out = os.path.join(TMP, 'b_out.docx')
    _quiet(V.dedupe_bookmarks, dup, out)
    s = zipfile.ZipFile(out).read('word/document.xml').decode('utf-8')
    assert 'bookmarkStart' not in s and 'bookmarkEnd' not in s  # bookmark 자체를 제거하는 설계


def t_figure_spec_by_pixels():
    from PIL import Image
    p = os.path.join(TMP, 'fig.jpg')
    Image.new('RGB', (1500, 900)).save(p, dpi=(300, 300))   # 5 x 3 inch @300
    r, _ = _quiet(V.check_figure_spec, p)
    assert r['ok'] is True, r
    Image.new('RGB', (600, 600)).save(p, dpi=(72, 72))
    r, _ = _quiet(V.check_figure_spec, p)
    assert r['ok'] is False, r


def t_cli_all_runs():
    import subprocess
    p = _docx('cli.docx', BODY)
    r = subprocess.run([sys.executable, V.__file__, 'all', p], capture_output=True, text=True)
    assert r.returncode == 0 and '인용 검증' in r.stdout and 'P값 형식' in r.stdout, r.stdout[-500:] + r.stderr[-500:]


def t_v138_korean_placeholder_in_english_paragraph():
    """저자 10-01: 영어 본문 문단에 한글 자리표시 하나만 있어도 v1.3.7 은 그 문단을 경고 없이 통째로 뺐다
    (본문 단어 누락 · 그 안 인용이 사라져 뒤 번호가 순서 위반·결번으로 보임). 재현: 아래가 v1.3.7 에서 body 5 · 순서 위반 [7] · 결번 [2..6]."""
    paras = list(BODY); paras[3] = 'First [1]. Then [2,3] and [4-6]. [자리표시 — 244명·44명]'
    r, out = _quiet(V.check_word_count, _docx('w_ph.docx', paras))
    assert r['body'] == 11 and r['korean']['segments'] == 1 and r['korean']['paras'] == 0, r     # 성공 길: 괄호만 지우고 영어는 센다
    assert '한글 괄호 구간 1곳' in out, out
    r, out = _quiet(V.check_citations, _docx('c_ph.docx', paras))
    assert r['used'] == 7 and r['order_violations'] == [] and r['gaps'] == [], r
    paras[3] = 'First [1]. Then [2,3] and [4-6]. 〔수정〕 (한글 메모)'                              # 〔〕·() 도 같은 길
    r, _ = _quiet(V.check_word_count, _docx('w_ph2.docx', paras))
    assert r['body'] == 11 and r['korean']['segments'] == 2, r


def t_v138_korean_outside_brackets_dropped_with_warning():
    """실패 길: 괄호 밖에도 한글이 남는 문단은 종전대로 통째로 빼되, 이제 몇 문단·몇 낱말·그 안 인용을 알린다."""
    paras = list(BODY); paras[3] = 'First [1]. Then [2,3] and [4-6] 여기 한글 문장.'
    r, out = _quiet(V.check_word_count, _docx('w_kd.docx', paras))
    assert r['body'] == 5 and r['korean']['paras'] == 1 and r['korean']['words'] == 6, r
    assert '문단 1개를 통째로 뺐습니다' in out and '영어 낱말 6개' in out, out
    r, out = _quiet(V.check_citations, _docx('c_kd.docx', paras))
    assert r['gaps'] == [2, 3, 4, 5, 6] and r['korean']['cites'] == ['[1]', '[2,3]', '[4-6]'], r   # 빠진 결과는 그대로 — 경고로 원인을 보인다
    assert '그 안 인용 [1] [2,3] [4-6]' in out, out
    r, out = _quiet(V.check_word_count, _docx('w_clean.docx', BODY))                      # 한글 없으면 경고도 없다
    assert r['korean'] == {'segments': 0, 'paras': 0, 'words': 0, 'cites': []} and '한글' not in out, out


if __name__ == '__main__':
    names = [n for n in list(globals()) if n.startswith('t_')]
    ok = fail = 0
    for n in names:
        try:
            globals()[n]()
            ok += 1; print(f'PASS {n:40s}')
        except Exception:
            fail += 1; print(f'FAIL {n:40s}'); traceback.print_exc()
    print(f'\n통과 {ok} / 건너뜀 0 / 실패 {fail}  (전체 {ok + fail})')
    sys.exit(1 if fail else 0)
