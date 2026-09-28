"""deck_toolkit 자체 검증.

실행: python test_toolkit.py
환경변수 (선택):
  TK_DECK     4:3 샘플 덱.  없으면 python-pptx 로 최소 덱을 만들어 쓴다
  TK_DECK169  16:9 샘플 덱.  없으면 만들어 쓴다
  TK_DECK4    원고 대조용 발표 덱 (crosscheck 테스트용). 없으면 fixture 로 대체
  TK_DOCS     원고 docx 경로들 (콤마 구분).            없으면 fixture 로 대체
  * test_claim_graph.py 의 공용 테스트(개수는 CLAIM_GRAPH.md §8)를 함께 돈다 (같은 폴더에 있어야 함)

결과 의미
  PASS  검증됨
  SKIP  입력이 없어 검증하지 못함 — 결함 아님. 그러나 "통과"도 아님
  FAIL  코드 결함
"""
import io, os, re, shutil, subprocess, sys, zipfile
sys.dont_write_bytecode = True   # /mnt/project 는 대화창 안에서 쓰기 가능 — __pycache__ 를 남기지 않는다 (v2.3.1)

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import deck_toolkit as T


def _manifest_version(fname):
    """같은 폴더의 TOOLS_MANIFEST.md 에 적힌 판. 없으면 None (Z1 이후: EXPECT_VERSION 만 맞추고 manifest 를 안 올린 사고 방지)."""
    import os, re
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'TOOLS_MANIFEST.md')
    if not os.path.exists(p):
        return None
    m = re.search(r'\| `%s` \| v([0-9.]+)' % re.escape(fname), open(p, encoding='utf8').read())
    return m.group(1) if m else None

EXPECT_VERSION = '16.34'

def t_deck_version_matches_manifest():
    assert getattr(T, '__version__', None) == EXPECT_VERSION, (getattr(T, '__version__', None), EXPECT_VERSION)
    mv = _manifest_version('deck_toolkit.py')
    assert mv is None or mv == EXPECT_VERSION, ('TOOLS_MANIFEST.md 의 판', mv, '코드', EXPECT_VERSION)

TMP = os.environ.get('TK_TMP', '/tmp/tk_test')
os.makedirs(TMP, exist_ok=True)


def _fixture(name, **kw):
    p = os.path.join(TMP, name)
    if not os.path.exists(p):
        T.make_fixture(p, **kw)
    return p


SRC = os.environ.get('TK_DECK') or _fixture('fixture43.pptx')
CHEST = os.environ.get('TK_DECK169') or _fixture('fixture169.pptx', widescreen=True)
DECK4 = os.environ.get('TK_DECK4')
DOCS = [s for s in os.environ.get('TK_DOCS', '').split(',') if s]

results = []


def check(name, fn):
    try:
        fn()
        results.append(('PASS', name, ''))
    except FileNotFoundError as e:
        results.append(('SKIP', name, '입력 파일 없음: %s' % e))
    except Exception as e:
        results.append(('FAIL', name, '%s: %s' % (type(e).__name__, e)))


def need(path):
    if not path or not os.path.exists(path):
        raise FileNotFoundError(path)
    return path


def wd(n):
    p = os.path.join(TMP, 'wd_' + n)
    if os.path.exists(p):
        shutil.rmtree(p)
    return p


def out(n):
    return os.path.join(TMP, n)


def cli(*args):
    return subprocess.run([sys.executable, os.path.join(HERE, 'deck_toolkit.py')] + list(args),
                          capture_output=True, text=True, cwd=TMP)


# ---------------------------------------------------------------- 기본 I/O
def t_roundtrip():
    d = T.Deck.open(SRC, wd('rt')); o = out('rt.pptx'); d.save(o)
    a = set(zipfile.ZipFile(SRC).namelist()); b = set(zipfile.ZipFile(o).namelist())
    assert a == b, 'part 불일치: %s' % (a ^ b)
    assert T.validate(o, SRC)

def t_audit():
    d = T.Deck.open(SRC, wd('au')); probs = d.audit(io.StringIO())
    assert any('중복 사용' in p for p in probs), '이미지 중복을 못 잡음'
    assert any('한 글자' in p for p in probs), '잔재 텍스트를 못 잡음'
    assert any('동일' in p for p in probs), '노트 중복을 못 잡음'

def t_edit():
    d = T.Deck.open(SRC, wd('ed'), theme='cud')
    b = (T.Body().header('H1').line('a ', ('b', T.C.KEY), ' c')
         .gap().line(('x', T.C.FLAG)).ref('Ref A, et al. J. 2020;1(1):1-2.'))
    d.set_body(7, b); d.set_title(7, 'New title'); d.set_notes(7, ['[노트]', '', '본문'])
    o = out('ed.pptx'); d.save(o); assert T.validate(o, SRC)
    d2 = T.Deck.open(o, wd('ed2'))
    assert d2.texts(7)[0] == 'New title', d2.texts(7)[:3]
    assert 'b' in d2.texts(7) and '[노트]' in d2.notes(7)

def t_body_single_tuple():
    """DECK_SPEC §10 예시가 ('text',) 1-튜플을 쓴다. 이게 터지면 안 된다."""
    x = T.Body().line(('No cortical destruction.',)).xml('cud')
    assert 'No cortical destruction.' in x and 'srgbClr' not in x

def t_title_multiline():
    d = T.Deck.open(SRC, wd('tm')); d.set_title(7, ['Line one', 'Line two'])
    o = out('tm.pptx'); d.save(o); assert T.validate(o, SRC)
    d2 = T.Deck.open(o, wd('tm2')); assert 'Line one' in d2.texts(7) and 'Line two' in d2.texts(7)

def t_escape():
    d = T.Deck.open(SRC, wd('es'), theme='cud')
    d.set_body(7, T.Body().line('A & B < C > D "q"')); d.set_notes(7, ['a & b < c'])
    o = out('es.pptx'); d.save(o); assert T.validate(o, SRC)

def t_save_guards():
    d = T.Deck.open(SRC, wd('sg'))
    for bad in (TMP, out('x.txt')):
        try:
            d.save(bad); raise AssertionError('잘못된 출력 경로를 막지 못함: %s' % bad)
        except (IsADirectoryError, ValueError):
            pass

# ---------------------------------------------------------------- add_slide
def t_add_slide():
    d = T.Deck.open(SRC, wd('as'), theme='cud'); before = [x[0] for x in d.order()]
    n = d.add_slide(after=6, title='Radiologic findings', body=T.Body().header('X').line('y'), notes=['n1'])
    after = [x[0] for x in d.order()]
    assert after.index(n) == after.index(6) + 1, after
    assert before == [x for x in after if x != n], '기존 순서가 바뀜'
    o = out('as.pptx'); d.save(o); assert T.validate(o, SRC)
    assert T.Deck.open(o, wd('as2')).notes(n) == ['n1']

def t_add_slide_x3():
    d = T.Deck.open(SRC, wd('a3'), theme='cud')
    ns = [d.add_slide(after=a, title='F%d' % a, body=T.Body().line('z'), notes=['x']) for a in (6, 10, 12)]
    o = out('a3.pptx'); d.save(o); assert T.validate(o, SRC)
    ordr = [x[0] for x in d.order()]
    assert len(set(ordr)) == len(ordr) and all(n in ordr for n in ns)

def t_add_slide_169():
    d = T.Deck.open(need(CHEST), wd('w9'), theme='cud'); w, h = d.slide_size()
    assert abs(w / h - 16 / 9) < 0.02, '16:9 fixture 가 아님'
    n = d.add_slide(after=2, title='T', body=T.Body().line('a'), notes=['n'])
    assert 'cx="%d"' % w in open(d._slide(n), encoding='utf8').read(), '제목 너비가 슬라이드 폭과 다름'
    o = out('w9.pptx'); d.save(o); assert T.validate(o, CHEST)

# ---------------------------------------------------------------- clrMap (§14)
def _invert(d):
    mp = os.path.join(d.dir, 'ppt/slideMasters/slideMaster1.xml')
    x = open(mp, encoding='utf8').read()
    x = re.sub(r'bg1="\w+"', 'bg1="dk1"', x); x = re.sub(r'tx1="\w+"', 'tx1="lt1"', x)
    open(mp, 'w', encoding='utf8').write(x)

def t_clrmap_normal():
    d = T.Deck.open(SRC, wd('cm0'))
    assert d.white(2) == 'bg1' and d.dark(2) == 'tx1', d.clrmap(2)

def t_clrmap_inverted():
    """실제 사고: bg1=dk1 인 덱에서 흰 글씨가 검정으로, 새 슬라이드 배경이 흰색으로 렌더됨."""
    d = T.Deck.open(SRC, wd('cm1')); _invert(d)
    assert d.white(2) == 'tx1' and d.dark(2) == 'bg1', d.clrmap(2)
    d.set_body(2, T.Body().line('plain')); d.set_title(2, 'T')
    n = d.add_slide(after=2, title='New', body=T.Body().line('p'), notes=['n'])
    s2 = open(d._slide(2), encoding='utf8').read(); sn = open(d._slide(n), encoding='utf8').read()
    assert 'schemeClr val="bg1"/>' not in s2.split('<p:txBody>', 1)[1], '본문에 bg1(=검정) 잔존'
    assert '<a:schemeClr val="bg1"><a:lumMod' in sn, '새 슬라이드 배경이 어두운 쪽을 안 씀'
    o = out('cm1.pptx'); d.save(o); assert T.validate(o, SRC)

# ---------------------------------------------------------------- lint / restyle / plan
def t_lint():
    d = T.Deck.open(SRC, wd('li'), theme='cud')
    d.set_body(7, T.Body().header('H').line(' '.join(['word'] * 300)))
    probs, summ = T.lint(d, io.StringIO())
    assert any('글자량 초과' in p for p in probs), probs
    assert summ['slides'] > 0

def t_lint_ref_exempt():
    d = T.Deck.open(SRC, wd('lr'), theme='cud')
    d.set_body(7, T.Body().header('H').line('short').ref('Kim J, et al. Radiology. 2020;1(1):1-2.'))
    probs, _ = T.lint(d, io.StringIO())
    assert not any('글씨 작음' in p and '10.0' in p for p in probs), probs

def t_restyle():
    d = T.Deck.open(SRC, wd('rs'), theme='cud'); T.restyle(d, stream=io.StringIO())
    o = out('rs.pptx'); d.save(o); assert T.validate(o, SRC)
    assert 'val="FFC000"' not in open(d._slide(7), encoding='utf8').read()

def t_restyle_keeps_ref():
    d = T.Deck.open(SRC, wd('rk'), theme='cud')
    d.set_body(7, T.Body().header('H').line('a').ref('Kim J, et al. Radiology. 2020;1(1):1-2.'))
    T.restyle(d, stream=io.StringIO())
    assert 'sz="1000"' in open(d._slide(7), encoding='utf8').read(), '참고문헌 10pt 가 강제로 커짐'

def t_plan():
    o = T.plan('journal_review', ['P'], io.StringIO())
    assert o[0][0] == 'title' and o[-1][0] == 'summary'

def t_classify_conference():
    """v10 사고: 학회 덱의 결과 슬라이드가 history 로 잡혀 45단어 상한이 적용됨."""
    d = T.Deck.open(SRC, wd('cl'), theme='cud')
    d.set_title(2, 'Results'); d.set_body(2, T.Body().line(''))
    d.set_title(3, 'Limitations'); d.set_body(3, T.Body().line('Single center'))
    d.set_title(4, 'Acknowledgments'); d.set_body(4, T.Body().line('Funding: NRF'))
    d.set_title(5, 'Effect sizes separate A from B')
    d.set_body(5, T.Body().line('r = 0.84, P < 0.001; FA 0.341 vs 0.352 across the full cohort of participants studied'))
    assert T.classify(d, 2, pos=2) == 'divider', T.classify(d, 2, pos=2)
    assert T.classify(d, 3, pos=3) == 'limitation'
    assert T.classify(d, 4, pos=4) == 'closing'
    assert T.classify(d, 5, pos=5) == 'result', T.classify(d, 5, pos=5)   # fixture 5번에는 이미지가 있다
    assert T.classify(d, 1, pos=1) == 'title'
    d.set_notes(5, ['[kind: image]', '', 'n']); assert T.classify(d, 5, pos=5) == 'image', '노트 태그 우선 실패'

def t_plan_kinds():
    for k in ('quiz', 'case_review', 'journal_original', 'journal_review_article', 'conference', 'journal_review'):
        o = T.plan(k, ['X'], io.StringIO()); assert all(kd in T.DENSITY for kd, _ in o), k
    o = T.plan('conference', ['X'], io.StringIO())
    assert sum(1 for kd, _ in o if kd == 'summary') == 1, 'conference 에 Take home 이 중복'

def t_image_aspect():
    """v12: 그림이 찌그러졌는지. 실제 학회 덱에서 뇌 영상이 가로 15% 늘어나 있었다."""
    d = T.Deck.open(SRC, wd('ia'))
    sn = next(s for s in d.slide_numbers() if d.images(s))
    assert not any('slide%d' % sn in p for p in d.check_image_aspect()), '정상 비율을 오탐'
    p = d._slide(sn); x = open(p, encoding='utf8').read()
    m = re.search(r'(<p:pic>.*?<a:ext cx=")(\d+)(" cy=")(\d+)(")', x, re.S)
    x = x[:m.start()] + m.group(1) + str(int(int(m.group(2)) * 1.3)) + m.group(3) + m.group(4) + m.group(5) + x[m.end():]
    open(p, 'w', encoding='utf8').write(x)
    assert any('slide%d' % sn in p and '왜곡' in p for p in d.check_image_aspect()), '30% 왜곡을 못 잡음'

def t_near_duplicate_bodies():
    d = T.Deck.open(SRC, wd('nd'), theme='cud')
    body = T.Body().line('Same long body text that appears on two consecutive slides for staged reveal of a figure')
    d.set_body(3, body); d.set_body(4, body)
    assert any('slide3' in p and 'slide4' in p for p in d.check_near_duplicate_bodies())

def t_plan_minutes():
    buf = io.StringIO(); T.plan('conference', ['X'], buf, minutes=10)
    assert '시간 예산 10분' in buf.getvalue() and '초과' in buf.getvalue()

def t_theme():
    for th, hexv in [('cud', '56B4E9'), ('amber', 'FFC000'), ('cyan', '00B0F0')]:
        assert hexv in T.Body().header('H').xml(th), th

# ---------------------------------------------------------------- polish 계열
def t_set_fonts():
    d = T.Deck.open(SRC, wd('sf')); n, f = T.set_fonts(d, 'safe')
    x = open(d._slide(7), encoding='utf8').read()
    assert 'typeface="Arial"' in x and 'typeface="맑은 고딕"' in x and n > 0
    o = out('sf.pptx'); d.save(o); assert T.validate(o, SRC)

def t_set_fonts_edge_runs():
    """채움 없는 rPr, solidFill 뒤 effectLst, hlinkClick 앞 순서, rPr 없는 run."""
    d = T.Deck.open(SRC, wd('sfe')); p = d._slide(2); x = open(p, encoding='utf8').read()
    inject = ('<a:p><a:r><a:rPr lang="en-US" sz="1600"><a:latin typeface="Nanum"/></a:rPr><a:t>A</a:t></a:r>'
              '<a:r><a:rPr lang="en-US" sz="1600"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
              '<a:effectLst/><a:latin typeface="Nanum"/><a:hlinkClick r:id="rId1"/></a:rPr><a:t>B</a:t></a:r>'
              '<a:r><a:t>D</a:t></a:r></a:p>')
    open(p, 'w', encoding='utf8').write(x.replace('</p:txBody>', inject + '</p:txBody>', 1))
    T.set_fonts(d, 'safe'); y = open(p, encoding='utf8').read()
    assert 'Nanum' not in y, '기존 폰트 잔존'
    for m in re.finditer(r'<a:r>(<a:rPr.*?</a:rPr>)<a:t>([ABD])</a:t>', y, re.S):
        rpr = m.group(1)
        assert 'typeface="Arial"' in rpr, (m.group(2), rpr)
        if '<a:hlinkClick' in rpr:
            assert rpr.index('<a:latin') < rpr.index('<a:hlinkClick'), '스키마 순서 위반'
        if '<a:effectLst' in rpr:
            assert rpr.index('<a:effectLst') < rpr.index('<a:latin'), '스키마 순서 위반'

def t_font_profiles():
    d = T.Deck.open(SRC, wd('fp')); T.set_fonts(d, 'office')
    assert 'typeface="Calibri"' in open(d._slide(7), encoding='utf8').read()
    o = out('fp.pptx'); d.save(o); assert T.validate(o, SRC)

def t_merge_runs():
    d = T.Deck.open(SRC, wd('mr')); before = sum(len(d.texts(s)) for s in d.slide_numbers())
    joined_b = ''.join(''.join(d.texts(s)) for s in d.slide_numbers())
    T.merge_runs(d); after = sum(len(d.texts(s)) for s in d.slide_numbers())
    assert after <= before
    o = out('mr.pptx'); d.save(o); assert T.validate(o, SRC)
    d2 = T.Deck.open(o, wd('mr2'))
    assert ''.join(''.join(d2.texts(s)) for s in d2.slide_numbers()) == joined_b, '병합으로 글자가 바뀜'

def t_clean_placeholders():
    d = T.Deck.open(SRC, wd('cp')); assert T.clean_placeholders(d) >= 0
    o = out('cp.pptx'); d.save(o); assert T.validate(o, SRC)

def t_autofit():
    d = T.Deck.open(SRC, wd('af')); T.force_autofit(d)
    o = out('af.pptx'); d.save(o); assert T.validate(o, SRC)

def t_phi():
    d = T.Deck.open(SRC, wd('ph')); d.set_notes(2, ['환자번호 12345678 확인'])
    hits = T.scan_phi(d, io.StringIO()); assert any('slide2' in h for h in hits), hits

def t_phi_ignores_dates():
    d = T.Deck.open(SRC, wd('phd')); d.set_notes(2, ['촬영일 20260906 기준'])
    hits = T.scan_phi(d, io.StringIO())
    assert not any('slide2' in h and '등록번호' in h for h in hits), hits

def t_bounds():
    d = T.Deck.open(SRC, wd('bd')); assert isinstance(T.check_bounds(d, io.StringIO()), list)

def t_handout():
    d = T.Deck.open(SRC, wd('ho')); p = T.export_notes(d, out('handout.md'))
    txt = open(p, encoding='utf8').read(); assert txt.startswith('# 발표 대본') and '## 1.' in txt

def t_polish_full():
    d = T.Deck.open(SRC, wd('pf')); r = T.polish(d, 'safe', io.StringIO())
    o = out('pf.pptx'); d.save(o); assert T.validate(o, SRC)
    assert isinstance(r['bounds'], list) and isinstance(r['phi'], list) and isinstance(r['overflow'], list)

# ---------------------------------------------------------------- 교차검증 계열
def _corpus_fixture():
    """원고 docx 가 없을 때 쓰는 대체 원전 (txt)."""
    p = out('corpus.txt')
    open(p, 'w', encoding='utf8').write(
        'Table 1. IDX 0.523 (0.49-0.55). FA 0.341 vs 0.352; P = 0.617. '
        'MD decreased 11.770 to 6.151. b = 800 and 2000 s/mm2.\n')
    return [p]

def _docx_fixture():
    """TK_DOCS 가 없으면 python-docx 로 원고 모양 docx 를 만든다. SKIP 은 통과가 아니므로 의존을 없앤다."""
    p = os.path.join(TMP, 'fixture_ms.docx')
    if True:   # 매번 새로 만든다 — 지난 실행의 캐시가 남아 fixture 를 고쳐도 옛 내용으로 돌던 일(v16.20)
        from docx import Document
        d = Document()
        d.add_paragraph('RESULTS')
        d.add_paragraph('IDX was higher in pLES than dLES (0.612 vs 0.523 a.u., P < 0.001); '
                        'FA did not differ (P = 0.617). MD was 11.8 at b = 800 and 6.2 at b = 2000 s/mm2. ' * 3)
        tb = d.add_table(rows=2, cols=2); tb.cell(0, 0).text = 'IDX'; tb.cell(0, 1).text = '0.523'
        d.save(p)
    return p

def t_read_docx():
    src = DOCS[0] if DOCS else _docx_fixture()
    t = T.read_docx(need(src)); assert len(t) > 100

def t_crosscheck():
    srcs = DOCS or _corpus_fixture()
    d = T.Deck.open(SRC, wd('cc'), theme='cud')
    d.set_body(7, T.Body().line('IDX 0.523; MD 11.8 -> 6.2; FA P = 0.617; b = 2000'))
    un, tot = T.crosscheck(d, srcs, io.StringIO()); toks = [u[1] for u in un]
    assert '0.523' not in toks and '11.8' not in toks and '6.2' not in toks, toks
    assert '2000' not in toks, 'b=2000 이 연도로 버려지거나 미확인으로 남음'

def t_crosscheck_catches_bogus():
    srcs = DOCS or _corpus_fixture()
    d = T.Deck.open(SRC, wd('cb'), theme='cud')
    d.set_body(7, T.Body().line('r = 0.7777 and AUC 0.6543 and 0.8'))
    un, _ = T.crosscheck(d, srcs, io.StringIO()); toks = [u[1] for u in un]
    assert '0.7777' in toks and '0.6543' in toks, toks

def t_crosscheck_real_deck():
    """실제 발표 덱 + 원고가 있으면 그것으로, 없으면 fixture 덱에 수치 슬라이드를 넣어 docx fixture 와 대조."""
    if DECK4 and DOCS:
        d = T.Deck.open(need(DECK4), wd('cr'))
        for s in DOCS: need(s)
        un, tot = T.crosscheck(d, DOCS, io.StringIO()); assert tot > 10
        return
    d = T.Deck.open(SRC, wd('cr'), theme='cud')
    d.set_body(7, T.Body().line('IDX 0.612 vs 0.523; FA P = 0.617; MD 11.8 -> 6.2 at b = 2000'))
    un, tot = T.crosscheck(d, [_docx_fixture()], io.StringIO())
    assert tot >= 5 and not [u for u in un if u[1] in ('0.612', '0.523', '0.617', '2000')], un

def t_locate():
    srcs = DOCS or _corpus_fixture()
    h = T.locate(srcs, '0.617', '9.9999', stream=io.StringIO())
    assert '0.617' in h and '9.9999' not in h

def t_numbers_bvalue():
    n = T._numbers('b = 800 and 2000 s/mm2, published 2020, n=1065, -12.3')
    assert '2000' in n and '2020' not in n and '12.3' in n and '1065' in n, n

def t_sync_detects():
    d = T.Deck.open(SRC, wd('sy'), theme='cud')
    d.set_body(7, T.Body().line('MD decreased at b = 800 and 2000'))
    d.set_notes(7, ['diffusion did change at b = 800'])
    probs = T.check_notes_slide_sync(d, io.StringIO())
    assert any('slide7' in p and '주장' in p for p in probs), probs

def t_spoken_notes():
    lines = ['본문 문장 하나', 'ANTICIPATED QUESTIONS', 'Q1 ...']
    assert T._spoken_notes(lines) == ['본문 문장 하나']
    lines = ['Q&A 에서 자주 나오는 질문은 뒤에 있습니다', '본문 계속', '**예상 질문**', 'Q1']
    assert T._spoken_notes(lines) == lines[:2], T._spoken_notes(lines)
    lines = ['spoken', '────── ANTICIPATED QUESTIONS (not spoken) ──────', 'Q1']
    assert T._spoken_notes(lines) == ['spoken'], '실제 학회 덱의 표지 줄을 못 잡음'

# ---------------------------------------------------------------- 넘침
def t_overflow_detects():
    d = T.Deck.open(SRC, wd('od'))
    d.set_body(7, T.Body(size=1600).line(' '.join(['overflowing'] * 200)))
    assert any('slide7' in p for p in T.check_text_overflow(d, stream=io.StringIO()))

def t_overflow_clean():
    d = T.Deck.open(SRC, wd('oc')); d.set_body(7, T.Body(size=1600).line('short'))
    assert not any('slide7' in p for p in T.check_text_overflow(d, stream=io.StringIO()))

def _card_deck(name, box_h_in, anchor_mid, text, wrap_none=False, footnote=None, inset_left=(0.2, 0.2)):
    """카드(채움 사각형) + 텍스트 상자 덱. v16.5 overflow 테스트용 (도구회신 overflow검사범위)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for k, left in enumerate(inset_left):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(6), Inches(3.0)); card.name = 'Card'
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0x22, 0x33, 0x44)
        cy = 1.5 + (3.0 - box_h_in) / 2 if anchor_mid else 1.5 + 0.2
        tb = s.shapes.add_textbox(Inches(1 + left), Inches(cy), Inches(6 - 2 * left), Inches(box_h_in)); tb.name = 'Body'
        tf = tb.text_frame; tf.word_wrap = not wrap_none
        if anchor_mid:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].text = text; tf.paragraphs[0].runs[0].font.size = Pt(16)
        if footnote:
            fn = s.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(3), Inches(0.4)); fn.name = 'Footnote'
            fn.text_frame.word_wrap = False
            fn.text_frame.paragraphs[0].text = footnote; fn.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    p = os.path.join(TMP, name); prs.save(p); return p


def t_overflow_card_boundary_when_box_larger_than_card():
    # 사고 1 재현: 상자 3.3" > 카드 3.0", anchor=ctr, 텍스트가 상자엔 들어가나 카드 밖으로
    long = ' '.join(['word'] * 110)   # 16pt·5.6" 폭 → 약 3.1" 필요
    d = T.Deck.open(_card_deck('ov_card.pptx', 3.3, True, long), wd('ovc'))
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert any(p.startswith('[심각]') and '카드 "Card" 아래로' in p for p in probs), probs
    assert not any('자기 상자를 넘침' in p for p in probs), probs   # 상자 안에는 들어감

def t_overflow_box_outside_card_is_note():
    d = T.Deck.open(_card_deck('ov_box.pptx', 3.3, True, 'short text'), wd('ovb'))
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert any(p.startswith('[참고]') and '상자가 카드 "Card" 밖으로' in p for p in probs), probs
    assert not any(p.startswith('[심각]') for p in probs), probs

def t_overflow_wrap_none_horizontal():
    # 사고 2 재현: wrap=none 각주가 슬라이드 오른쪽 밖으로
    d = T.Deck.open(_card_deck('ov_wrap.pptx', 2.5, False, 'ok', footnote='footnote ' * 60), wd('ovw'))
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert any(p.startswith('[심각]') and 'Footnote' in p and '슬라이드 오른쪽 밖' in p for p in probs), probs
    d = T.Deck.open(_card_deck('ov_wrap2.pptx', 2.5, False, 'ok', footnote='footnote ' * 8), wd('ovw2'))
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert any(p.startswith('[심각]') and 'Footnote' in p and '상자 폭을 넘음' in p for p in probs), probs
    d = T.Deck.open(_card_deck('ov_wrap3.pptx', 2.5, False, 'ok', footnote='fn'), wd('ovw3'))
    assert not [p for p in T.check_text_overflow(d, stream=io.StringIO()) if 'Footnote' in p]

def t_overflow_headroom_option():
    mid = ' '.join(['word'] * 80)   # 상자 2.5" 에 약 2.3" 필요 → 85% 넘음, 자기 상자는 안 넘음
    d = T.Deck.open(_card_deck('ov_head.pptx', 2.5, False, mid), wd('ovh'))
    assert not [p for p in T.check_text_overflow(d, stream=io.StringIO()) if 'Body' in p]
    probs = T.check_text_overflow(d, headroom=0.15, stream=io.StringIO())
    assert any('여유 15% 미만' in p for p in probs), probs
    r = cli('overflow', _card_deck('ov_head.pptx', 2.5, False, mid), '--headroom', '0.15'); assert r.returncode == 0 and '여유' in r.stdout

def t_card_insets_header_band_alignment():
    """v16.5.2: 카드 폭을 꽉 채우는 헤더 띠와 들여쓴 본문의 어긋남 — v16.5.1 은 소속 판정(완전 포함)
    때문에 놓쳤다(발표 260910 §3 실물 사례)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for hdr_x, hdr_w in ((1.0, 6.05), (1.3, 5.4)):   # 1: 헤더가 카드 가장자리+살짝 초과 / 2: 본문과 정렬
        s = prs.slides.add_slide(prs.slide_layouts[6])
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(6), Inches(4)); card.name = 'Card'
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(1, 2, 3)
        h = s.shapes.add_textbox(Inches(hdr_x), Inches(1.5), Inches(hdr_w), Inches(0.5)); h.name = 'Header'
        h.text_frame.paragraphs[0].text = 'Header'; h.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        b = s.shapes.add_textbox(Inches(1.3), Inches(2.2), Inches(5.4), Inches(1.0)); b.name = 'Body'
        b.text_frame.paragraphs[0].text = 'Body'; b.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    p = os.path.join(TMP, 'ov_hdr.pptx'); prs.save(p)
    probs = T.check_card_insets(T.Deck.open(p, wd('ovhd')), stream=io.StringIO())
    hits = [x for x in probs if '좌우가 안 맞음' in x]
    assert len(hits) == 1 and 'slide1' in hits[0] and 'Header' in hits[0], probs

def t_card_insets_first_box_only_and_left_align():
    # v16.5.1: 카드 안 두 번째(요약) 상자의 y 는 위 여백에 안 들어감; 왼쪽이 다르면 정렬 [참고]
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for second_x in (0.2, 0.5):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(6), Inches(4)); card.name = 'Card'
        card.fill.solid(); card.fill.fore_color.rgb = RGBColor(1, 2, 3)
        for k, (x, y) in enumerate(((0.2, 0.2), (second_x, 3.0))):
            tb = s.shapes.add_textbox(Inches(1 + x), Inches(1.5 + y), Inches(4), Inches(0.6)); tb.name = 'T%d' % k
            tb.text_frame.paragraphs[0].text = 'text'; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    p = os.path.join(TMP, 'ov_first.pptx'); prs.save(p)
    probs = T.check_card_insets(T.Deck.open(p, wd('ovf')), stream=io.StringIO())
    assert not any('위 여백' in x for x in probs), probs           # 첫 상자 y 는 두 슬라이드 모두 0.2"
    assert any('좌우가 안 맞음' in x and 'slide2' in x for x in probs), probs
    assert not any('slide1' in x for x in probs), probs

def t_lint_card_inset_consistency():
    d = T.Deck.open(_card_deck('ov_ins.pptx', 2.5, False, 'text', inset_left=(0.2, 0.45)), wd('ovi'))
    probs = T.check_card_insets(d, stream=io.StringIO())
    assert any('왼쪽 여백이 슬라이드마다 다름' in p for p in probs), probs
    d = T.Deck.open(_card_deck('ov_ins2.pptx', 2.5, False, 'text', inset_left=(0.2, 0.25)), wd('ovi2'))
    assert not T.check_card_insets(d, stream=io.StringIO())
    lp, _ = T.lint(T.Deck.open(_card_deck('ov_ins.pptx', 2.5, False, 'text', inset_left=(0.2, 0.45)), wd('ovi3')), stream=io.StringIO())
    assert any('왼쪽 여백' in p for p in lp), lp

def t_merge_runs_preserves_text():
    """전평 대화창 9/8 보고: self-closing rPr 에서 lazy 정규식이 중간 run 을 삼켜 ', ' 가 사라졌다."""
    d = T.Deck.open(SRC, wd('mrp'))
    fp = d._slide(1); x = open(fp, encoding='utf8').read()
    inj = ('<a:p><a:r><a:rPr lang="ko-KR"/><a:t>당뇨, </a:t></a:r><a:r><a:rPr lang="ko-KR"/><a:t>심부전, </a:t></a:r>'
           '<a:r><a:rPr lang="ko-KR"/><a:t>신부전</a:t></a:r><a:r><a:rPr lang="ko-KR" b="1"/><a:t>간경화</a:t></a:r></a:p>')
    i = x.rfind('</p:txBody>'); x = x[:i] + inj + x[i:]
    open(fp, 'w', encoding='utf8').write(x)
    before = ''.join(d.texts(1))
    n = T.merge_runs(d)
    after = ''.join(d.texts(1))
    assert after == before, (before, after)
    assert '당뇨, 심부전, 신부전' in ''.join(re.findall(r'<a:t[^>]*>([^<]*)</a:t>', open(fp, encoding='utf8').read()))
    assert n >= 2, n

def t_polish_reports_text_intact():
    d = T.Deck.open(SRC, wd('pti'))
    r = T.polish(d, stream=io.StringIO())
    assert r['text_intact'] is True

def t_import_slide_roundtrip():
    """전평 대화창 요청: 다른 덱 슬라이드 복사 후 validate 통과, 순서·번호 불변, 노트 새로 생성."""
    dst = T.Deck.open(SRC, wd('imp_dst')); src = T.Deck.open(SRC, wd('imp_src'))
    n0 = len(dst.slide_numbers()); order0 = [s for s, _, _ in dst.order()]
    new = dst.import_slide(src, 2, after=1, copy_notes=True)
    order1 = [s for s, _, _ in dst.order()]
    assert len(order1) == n0 + 1 and order1[1] == new and order1[0] == order0[0], (order0, order1)
    assert ''.join(dst.texts(new)) == ''.join(src.texts(2))
    dst.save('/tmp/imp.pptx'); assert T.validate('/tmp/imp.pptx', SRC)
    d2 = T.Deck.open('/tmp/imp.pptx', wd('imp_re'))
    assert len(d2.slide_numbers()) == n0 + 1
    dst.remove_slide(new)
    assert [s for s, _, _ in dst.order()] == order0

def t_exam_theme_and_inherit_color():
    assert T.THEMES['exam'][T.C.PLAIN] == 'INHERIT'
    r = T._run('x', 'INHERIT', 1800, bold=False)
    assert 'solidFill' not in r and 'b="1"' not in r
    assert 'b="1"' in T._run('x', 'FF0000', 1800)

def t_label_screen_vs_xml():
    # 학회 덱 회신 결함 1: 출력에 화면 번호와 xml 번호를 함께
    d = T.Deck.open(SRC, wd('lbl'))
    first = [s for s, _, _ in d.order()][0]
    assert d.label(first) == '화면 1 (slide%d.xml)' % first
    assert d.relabel('slide%d: x / slide%d.xml' % (first, first)) == '화면 1 (slide%d.xml): x / slide%d.xml' % (first, first)

def t_overflow_label_over_within_tol():
    # 학회 덱 회신 결함 2: 필요 > 상자인데 4% 이내 → "상자를 넘음(허용)" 라벨, headroom 라벨 아님
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5.6), Inches(0.36)); tb.name = 'T'
    tb.text_frame.paragraphs[0].text = 'one line of text here'; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    p = os.path.join(TMP, 'ov_lbl.pptx'); prs.save(p)
    d = T.Deck.open(p, wd('ovl'))
    need = T._estimated_height(T._text_shapes(d, 1)[0]); h = T._text_shapes(d, 1)[0]['h']
    probs = T.check_text_overflow(d, headroom=0.15, stream=io.StringIO())
    if h < need <= h * 1.04:
        assert any('상자를 넘음(허용' in x for x in probs) and not any('여유' in x for x in probs), (probs, need, h)
    else:
        assert not any('상자를 넘음(허용' in x for x in probs), (probs, need, h)

def t_read_any_magic_bytes():
    # 학회 덱 회신 결함 3: 이름만 docx 인 텍스트 파일
    p = os.path.join(TMP, 'fake.docx'); open(p, 'w', encoding='utf8').write('plain text 0.523')
    assert '0.523' in T.read_any(p)

def t_audit_orphan_image_separated():
    # 학회 덱 회신 결함 4: rels 에만 있는 이미지는 media 칸이 아니라 [참고]
    d = T.Deck.open(SRC, wd('orph'))
    sn = [s for s, _, _ in d.order()][0]
    rp = d._slide_rels(sn); r = open(rp, encoding='utf8').read()
    r = r.replace('</Relationships>', '<Relationship Id="rId77" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/ghost.png"/></Relationships>')
    open(rp, 'w', encoding='utf8').write(r)
    assert 'ghost.png' in d.orphan_images(sn) and 'ghost.png' not in d.images(sn)
    probs = d.audit(io.StringIO())
    assert any('배치되지 않은 이미지 rel' in x and 'ghost.png' in x for x in probs), probs

def t_lint_density_label_and_kind_warning():
    d = T.Deck.open(SRC, wd('lk'))
    buf = io.StringIO(); probs, _ = T.lint(d, stream=buf)
    assert '밀도 기준' in buf.getvalue() and '청중이 화면 글자를 읽는 데' not in buf.getvalue()
    assert any('[kind:] 태그 없음' in x for x in probs), probs

def t_crosscheck_notes():
    srcs = DOCS or _corpus_fixture()
    d = T.Deck.open(SRC, wd('ccn'), theme='cud')
    d.set_body(7, T.Body().line('IDX 0.523')); d.set_notes(7, ['expected question: about 0.7777 and 0.523'])
    un, tot = T.crosscheck(d, srcs, io.StringIO()); assert not [u for u in un if u[2] == '노트']
    un, tot2 = T.crosscheck(d, srcs, io.StringIO(), notes=True)
    assert tot2 > tot and any(u[1] == '0.7777' and u[2] == '노트' for u in un), un

def t_sync_notes_only_numbers_note():
    d = T.Deck.open(SRC, wd('syn'), theme='cud')
    d.set_body(7, T.Body().line('MD decreased')); d.set_notes(7, ['MD decreased; AUC about 0.73 and 0.80'])
    buf = io.StringIO(); T.check_notes_slide_sync(d, buf)
    assert '노트에만 있는 수치' in buf.getvalue(), buf.getvalue()

def t_settext_rules():
    # 학회 덱 회신 요청 8
    d = T.Deck.open(SRC, wd('stx'), theme='cud')
    d.set_body(7, T.Body().line('alpha beta gamma').line('alpha again'))
    r = d.settext(7, 'alpha', 'ALPHA', dry_run=True)
    assert not r['ok'] and r['count'] == 2, r                       # 2회 → 거부
    r = d.settext(7, 'zeta', 'Z'); assert not r['ok'] and r['count'] == 0
    r = d.settext(7, 'beta gamma', 'BETA', dry_run=True)
    assert r['ok'] and r['dry_run'] and 'beta gamma' in ''.join(d.texts(7))   # dry-run 은 안 바꿈
    r = d.settext(7, 'beta gamma', 'BETA'); assert r['ok'] and 'alpha BETA' in ''.join(d.texts(7))
    r = d.settext(7, 'alpha again', '', delete_para=True); assert r['ok'] and 'again' not in ''.join(d.texts(7))
    d.set_notes(7, ['note line one', 'note line two'])
    r = d.settext(7, 'line two', 'LINE 2', notes=True); assert r['ok'] and any('LINE 2' in t for t in d.notes(7))
    d.save('/tmp/stx.pptx'); assert T.validate('/tmp/stx.pptx', SRC)
    r = cli('settext', SRC, '--slide', '7', '--old', 'nonexistent', '--new', 'x', '--dry-run'); assert r.returncode == 1 and '거부' in r.stdout

def t_strip_color_and_fix_title_box():
    d = T.Deck.open(SRC, wd('sc'), theme='cud')
    fp = d._slide(7); x = open(fp, encoding='utf8').read()
    i = x.rfind('</p:txBody>')
    x = x[:i] + '<a:p><a:r><a:rPr lang="en-US"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:rPr><a:t>answer</a:t></a:r></a:p>' + x[i:]
    open(fp, 'w', encoding='utf8').write(x)
    assert d.strip_color(7, 'FF0000') == 1 and 'FF0000' not in open(fp, encoding='utf8').read()
    assert 'answer' in ''.join(d.texts(7))
    ok = d.fix_title_box(7, 1505129)
    x = open(fp, encoding='utf8').read()
    if ok:
        assert 'cy="1505129"' in x and '<a:normAutofit' in x
    d.save('/tmp/sc.pptx'); assert T.validate('/tmp/sc.pptx', SRC)

def t_group_frame_transform():
    """v16.6: 그룹 안 텍스트 상자는 chOff/chExt 좌표계 — 그룹 프레임으로 변환해야 카드·슬라이드 경계 비교가 맞다."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    card = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2)); card.name = 'GCard'
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(1, 2, 3)
    tb = grp.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(3.6), Inches(0.5)); tb.name = 'GText'
    tb.text_frame.paragraphs[0].text = 'grouped text'; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    # 그룹 프레임을 자식 공간과 다르게 옮기고 2배로 키운다
    grp.left, grp.top, grp.width, grp.height = Inches(5), Inches(3), Inches(8), Inches(4)
    p = os.path.join(TMP, 'grp.pptx'); prs.save(p)
    d = T.Deck.open(p, wd('grp'))
    sh = next(x for x in T._text_shapes(d, 1) if x['name'] == 'GText')
    c = next(x for x in T._fill_shapes(d, 1) if x['name'] == 'GCard')
    assert sh['in_group'] and abs(sh['x'] - int(Inches(5) + (Inches(1.2) - Inches(1)) * 2)) < 2000, sh
    assert abs(c['x'] - Inches(5)) < 2000 and abs(c['w'] - Inches(8)) < 2000, c
    assert not [x for x in T.check_text_overflow(d, stream=io.StringIO()) if x.startswith('[심각]')]

def t_normautofit_scale_warning():
    d = T.Deck.open(SRC, wd('naf'))
    fp = d._slide(7); x = open(fp, encoding='utf8').read()
    x = x.replace('<a:bodyPr/>', '<a:bodyPr><a:normAutofit fontScale="70000" lnSpcReduction="20000"/></a:bodyPr>', 1)
    open(fp, 'w', encoding='utf8').write(x)
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert any('70%' in p and '자동 축소' in p for p in probs), probs

def t_overflow_font_path():
    import glob
    fonts = glob.glob('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    assert fonts, 'DejaVuSans.ttf 없음 — 컨테이너 폰트 확인'
    assert T._est_lines_font('short', 16, 5 * 914400, fonts[0]) == 1
    assert T._est_lines_font(' '.join(['word'] * 60), 16, 3 * 914400, fonts[0]) > 3
    d = T.Deck.open(SRC, wd('ofp'))
    a = T.check_text_overflow(d, stream=io.StringIO()); b = T.check_text_overflow(d, stream=io.StringIO(), font_path=fonts[0])
    assert isinstance(a, list) and isinstance(b, list)
    r = cli('overflow', SRC, '--font-path', fonts[0]); assert r.returncode == 0 and 'DejaVuSans' in r.stdout

def t_audit_residue_per_paragraph():
    # 전평 수용검사 2-1: 색 강조로 쪼개진 run 'A '·'.' 은 잔재가 아니다; 문단 전체가 한 글자면 잔재
    d = T.Deck.open(SRC, wd('res'))
    sn = [s for s, _, _ in d.order()][2]
    fp = d._slide(sn); x = open(fp, encoding='utf8').read(); i = x.rfind('</p:txBody>')
    inj = '<a:p><a:r><a:rPr lang="en-US"/><a:t>A </a:t></a:r><a:r><a:rPr lang="en-US" b="1"/><a:t>bold</a:t></a:r><a:r><a:rPr lang="en-US"/><a:t>.</a:t></a:r></a:p>'
    open(fp, 'w', encoding='utf8').write(x[:i] + inj + x[i:])
    assert not [p for p in d.audit(io.StringIO()) if '한 글자 잔재' in p and 'slide%d' % sn in p]
    x = open(fp, encoding='utf8').read(); i = x.rfind('</p:txBody>')
    open(fp, 'w', encoding='utf8').write(x[:i] + '<a:p><a:r><a:rPr lang="en-US"/><a:t>B</a:t></a:r></a:p>' + x[i:])
    assert [p for p in d.audit(io.StringIO()) if '한 글자 잔재' in p and 'slide%d' % sn in p]

def t_overflow_autofit_scaled_downgrades():
    # 전평 수용검사 2-2: fontScale 이 있는(실제로 줄인) 상자의 [심각]은 [참고]로
    long = ' '.join(['word'] * 110)
    p = _card_deck('ov_af.pptx', 3.3, True, long)
    d = T.Deck.open(p, wd('ovaf'))
    assert any(x.startswith('[심각]') for x in T.check_text_overflow(d, stream=io.StringIO()))
    for sn in d.slide_numbers():
        fp = d._slide(sn); x = open(fp, encoding='utf8').read()
        x = x.replace('<a:spAutoFit/>', '<a:normAutofit fontScale="62500"/>')
        open(fp, 'w', encoding='utf8').write(x)
    probs = T.check_text_overflow(d, stream=io.StringIO())
    assert not any(x.startswith('[심각]') for x in probs) and any(x.startswith('[자동맞춤 의존]') and 'Google Slides' in x for x in probs), probs   # v16.10 W1

def t_purge_orphans():
    # 전평 수용검사 2-3: remove_slide 뒤 남은 파일·노트·미디어 정리, validate 통과
    d = T.Deck.open(SRC, wd('purge'))
    n0 = len(d.slide_numbers()); order0 = [s for s, _, _ in d.order()]
    removed = order0[-1]; d.remove_slide(removed)
    assert removed in d.slide_numbers()          # 파일은 남아 있음
    r = d.purge_orphans()
    assert r['slides'] == 1 and removed not in d.slide_numbers() and len(d.slide_numbers()) == n0 - 1, r
    assert 'slide%d.xml' % removed not in open(os.path.join(d.dir, '[Content_Types].xml'), encoding='utf8').read()
    d.save('/tmp/purge.pptx'); assert T.validate('/tmp/purge.pptx')
    d2 = T.Deck.open('/tmp/purge.pptx', wd('purge2')); assert len(d2.order()) == n0 - 1
    r = cli('purge', SRC, '-o', '/tmp/purge_cli.pptx'); assert r.returncode == 0 and '제거: 슬라이드 0' in r.stdout

def t_sync_korean_claim_terms():
    d = T.Deck.open(SRC, wd('syk'), theme='cud')
    d.set_body(7, T.Body().line('Increased mesenteric fat attenuation')); d.set_notes(7, ['장간막 지방 음영 증가 소견'])
    probs = T.check_notes_slide_sync(d, io.StringIO())
    assert not [p for p in probs if 'increase' in p], probs

def t_cutoff_widened_and_nearmiss():
    # MSK 회신 N1
    assert T._is_cutoff_line('─────────  [참고 — NOT SPOKEN]  ─────────')
    assert T._is_cutoff_line('───────── NOT SPOKEN · 참고 ─────────')
    assert not T._is_cutoff_line('이 내용은 Q&A 에서 자주 나오는 질문이라 대본에 넣었다')   # 본문 속 Q&A 는 절단 아님
    assert T._cutoff_status(['a', '━━━━━━━━ 뒷부분 참고용 ━━━━━━━━', 'b'])[0] == 'nearmiss'
    d = T.Deck.open(SRC, wd('nm'), theme='cud')
    d.set_notes(7, ['spoken', '━━━━━━━━ 뒷부분 참고용 ━━━━━━━━', 'tip'])
    probs, _ = T.lint(d, stream=io.StringIO())
    assert any('절단되지 않음' in x for x in probs), probs

def t_set_notes_with_tips_and_split():
    d = T.Deck.open(SRC, wd('tips'), theme='cud')
    d.set_notes(7, ['one', 'two'], tips=['t1', 't2'])
    sc, tp = d.split_notes(7)
    assert sc == ['one', 'two'] and tp == ['t1', 't2'], (sc, tp)
    assert T._spoken_notes(d.notes(7)) == ['one', 'two']

def t_classify_review_with_image_and_basis():
    d = T.Deck.open(SRC, wd('rv'), theme='cud')
    d.set_title(7, ['PAES review']); d.set_body(7, T.Body().line('text ' * 200))
    fp = d._slide(7); x = open(fp, encoding='utf8').read()
    # 그림이 있다고 가정: images() 가 비어 있어도 제목 키워드가 review 를 먼저 잡는지
    assert T.classify(d, 7) == 'review'
    k, why = T.classify_why(d, 7); assert k == 'review' and 'review' in why
    probs, _ = T.lint(d, stream=io.StringIO())
    assert any('종류 판정 근거' in p for p in probs if 'slide7' in p), probs

def t_move_slide():
    d = T.Deck.open(SRC, wd('mv'))
    o = [s for s, _, _ in d.order()]
    d.move_slide(o[2], after=o[0]); o2 = [s for s, _, _ in d.order()]
    assert o2[:3] == [o[0], o[2], o[1]] and sorted(o2) == sorted(o), (o, o2)
    d.move_slide(o[2], after=0); assert [s for s, _, _ in d.order()][0] == o[2]
    d.save('/tmp/mv.pptx'); assert T.validate('/tmp/mv.pptx', SRC)

def t_sldid_sites_and_conversion():
    d = T.Deck.open(SRC, wd('sid'), theme='cud')
    o = [s for s, _, _ in d.order()]
    sid = d.sld_id(o[3]); assert d.slide_by_id(sid) == o[3]
    assert T._site_text(d, 'slide@%d' % sid) == T._site_text(d, 'slide:%d' % o[3])
    cl = [{'id': 'a', 'statement': 's', 'evidence': 'e', 'sites': ['slide:%d' % o[3], 'notes:%d' % o[3], 'doc:find:x'], 'keys': ['slide'], 'depends_on': []}]
    n = T.sites_to_sldid(d, cl); assert n == 2 and cl[0]['sites'][0] == 'slide@%d' % sid and cl[0]['sites'][2] == 'doc:find:x'
    # 순서 ≠ 파일 번호 덱에서 파일 번호 사이트 경고
    d.move_slide(o[3], after=0)
    buf = io.StringIO(); T._order_warning(d, [{'sites': ['slide:%d' % o[3]]}], buf)
    assert '화면 순서 ≠ 파일 번호' in buf.getvalue()
    T._order_warning(d, cl, buf2 := io.StringIO()); assert buf2.getvalue() == ''   # sldId 사이트만이면 경고 없음

def t_clear_pictures_and_import_without_pictures():
    d = T.Deck.open(SRC, wd('cp'))
    with_pic = next((s for s in d.slide_numbers() if d.images(s)), None)
    if with_pic is None:
        return
    img = d.images(with_pic)[0]
    src = T.Deck.open(SRC, wd('cp_src'))
    new = d.import_slide(src, with_pic, after=with_pic, pictures=False)
    assert d.images(new) == [] and '<p:pic>' not in open(d._slide(new), encoding='utf8').read()
    n = d.clear_pictures(with_pic); assert n >= 1 and d.images(with_pic) == []
    d.save('/tmp/cp.pptx'); assert T.validate('/tmp/cp.pptx')

def t_replace_insert_delete_paragraph():
    d = T.Deck.open(SRC, wd('rp'), theme='cud')
    d.set_body(7, T.Body().line('alpha one').line('beta two').line('gamma three'))
    old = d.replace_paragraph(7, 'beta', [('MHG anatomy', T.C.HEAD), ': plain text'], size=1600)
    assert old == 'beta two' and 'MHG anatomy' in ''.join(d.texts(7)) and 'beta two' not in ''.join(d.texts(7))
    d.insert_after(7, 'gamma', ['inserted line'])
    t = ''.join(d.texts(7)); assert t.index('gamma three') < t.index('inserted line')
    try:
        d.replace_paragraph(7, 'zzz', ['x']); assert False
    except ValueError as e:
        assert '0회' in str(e)
    d.delete_paragraph(7, 'alpha'); assert 'alpha' not in ''.join(d.texts(7))
    d.save('/tmp/rp.pptx'); assert T.validate('/tmp/rp.pptx', SRC)

def t_set_cite_and_image_grid():
    d = T.Deck.open(SRC, wd('cite'), theme='cud')
    d.set_cite(7, ['Macedo AJR 2003;181:1259', 'Hai AJR 2008'])
    d.set_cite(7, ['only one'])                                  # 있으면 교체
    x = open(d._slide(7), encoding='utf8').read()
    assert x.count('name="CiteBox"') == 1 and 'only one' in x and 'Macedo' not in x and 'wrap="square"' in x
    n = d.image_grid(7, [('Lt', 'axial CT'), ('Rt', 'axial CT'), ('Cor', 'MIP')], cols=2)
    x = open(d._slide(7), encoding='utf8').read()
    assert n == 3 and x.count('ImagePlaceholder') == 3 and x.count('SeqLabel') == 3 and '[넣을 영상] MIP' in x
    ids = re.findall(r'<p:cNvPr id="(\d+)"', x); assert len(ids) == len(set(ids)), '도형 id 중복'
    d.save('/tmp/cite.pptx'); assert T.validate('/tmp/cite.pptx', SRC)
    assert not [p for p in T.check_text_overflow(d, stream=io.StringIO()) if 'CiteBox' in p and p.startswith('[심각]')]

def t_audit_residue_table_dash_and_bullet_para():
    # MSK 수용검사 P1
    d = T.Deck.open(SRC, wd('res2'))
    sn = [s for s, _, _ in d.order()][2]
    fp = d._slide(sn); x = open(fp, encoding='utf8').read(); i = x.rindex('</p:spTree>')
    tbl = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="900" name="T"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
           '<p:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
           '<a:tbl><a:tblPr/><a:tblGrid><a:gridCol w="914400"/></a:tblGrid><a:tr h="300000"><a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>—</a:t></a:r></a:p></a:txBody><a:tcPr/></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    x = x[:i] + tbl + x[i:]
    i = x.rfind('</p:txBody>')
    x = x[:i] + '<a:p><a:r><a:rPr lang="ko-KR"/><a:t>-</a:t></a:r><a:r><a:rPr lang="ko-KR"/><a:t> 수술기록: 제거함</a:t></a:r></a:p>' + x[i:]
    open(fp, 'w', encoding='utf8').write(x)
    probs = [p for p in d.audit(io.StringIO()) if 'slide%d' % sn in p]
    assert not any('한 글자 잔재' in p for p in probs) and not any('빈 불릿' in p for p in probs), probs

def t_sync_mc_options_and_narrow_korean():
    d = T.Deck.open(SRC, wd('mc'), theme='cud')
    d.set_body(7, T.Body().line('㉡ 좌측보다 우측으로 발생하는 빈도가 높다').line('투과성 증가'))
    d.set_notes(7, ['정답 설명. 투과성이 늘어난다.'])
    probs = T.check_notes_slide_sync(d, io.StringIO())
    assert not any('높' in p for p in probs) and any('증가' in p for p in probs), probs   # 보기 줄 제외, 본문 '증가' 는 맞는 지적

def t_clear_pictures_with_labels():
    d = T.Deck.open(SRC, wd('cpl'))
    sn = next((s for s in d.slide_numbers() if d.images(s)), None)
    if sn is None:
        return
    fp = d._slide(sn); x = open(fp, encoding='utf8').read()
    g = re.search(r'<p:pic>.*?<a:off x="(-?\d+)" y="(-?\d+)"/>\s*<a:ext cx="(\d+)" cy="(\d+)"', x, re.S)
    bx, by, bw, bh = (int(v) for v in g.groups())
    lab = ('<p:sp><p:nvSpPr><p:cNvPr id="901" name="SeqLabel 9"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="300000" cy="200000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
           '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>T1 TSE SAG</a:t></a:r></a:p></p:txBody></p:sp>' % (bx + 10000, by + 10000))
    far = lab.replace('id="901" name="SeqLabel 9"', 'id="902" name="FarBox"').replace('T1 TSE SAG', 'far away').replace('<a:off x="%d" y="%d"/>' % (bx + 10000, by + 10000), '<a:off x="%d" y="%d"/>' % (bx + bw + 2000000, by))
    i = x.rindex('</p:spTree>'); open(fp, 'w', encoding='utf8').write(x[:i] + lab + far + x[i:])
    d.clear_pictures(sn, labels=True)
    t = ''.join(d.texts(sn)); assert 'T1 TSE SAG' not in t and 'far away' in t, t
    d.save('/tmp/cpl.pptx'); assert T.validate('/tmp/cpl.pptx')

def t_move_slide_after_pos():
    d = T.Deck.open(SRC, wd('mvp'))
    o = [s for s, _, _ in d.order()]
    d.move_slide(o[4], after_pos=1)
    assert [s for s, _, _ in d.order()][1] == o[4]

def t_lint_caption_font_exempt():
    d = T.Deck.open(SRC, wd('capf'), theme='cud')
    fp = d._slide(7); x = open(fp, encoding='utf8').read(); i = x.rindex('</p:spTree>')
    cap = ('<p:sp><p:nvSpPr><p:cNvPr id="903" name="LitCaption 3"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2000000" cy="300000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
           '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>Wright RG 2004 Fig 6 55M Doppler</a:t></a:r></a:p></p:txBody></p:sp>')
    open(fp, 'w', encoding='utf8').write(x[:i] + cap + x[i:])
    probs, _ = T.lint(d, stream=io.StringIO())
    assert not any('slide7' in p and '글씨 작음' in p for p in probs), probs

def t_lint_table_and_placeholder_font_rules():
    # MSK 수용검사 Y1
    d = T.Deck.open(SRC, wd('tf'), theme='cud')
    fp = d._slide(7); x = open(fp, encoding='utf8').read(); i = x.rindex('</p:spTree>')
    tbl = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="910" name="T"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
           '<p:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
           '<a:tbl><a:tblPr/><a:tblGrid><a:gridCol w="914400"/></a:tblGrid><a:tr h="300000"><a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>cell text</a:t></a:r></a:p></a:txBody><a:tcPr/></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    ph = ('<p:sp><p:nvSpPr><p:cNvPr id="911" name="ImagePlaceholder 1"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2000000" cy="2000000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
          '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="ko-KR" sz="1100"/><a:t>[넣을 영상] axial CT</a:t></a:r></a:p></p:txBody></p:sp>')
    open(fp, 'w', encoding='utf8').write(x[:i] + tbl + ph + x[i:])
    probs, _ = T.lint(d, stream=io.StringIO())
    mine = [p for p in probs if 'slide7' in p]
    assert not any('본문 글씨 작음' in p for p in mine), mine
    assert any('표 글씨 11pt' in p for p in mine), mine

def t_import_slide_margin_passthrough():
    d = T.Deck.open(SRC, wd('imm')); src = T.Deck.open(SRC, wd('imm_src'))
    sn = next((s for s in src.slide_numbers() if src.images(s)), None)
    if sn is None:
        return
    fp = src._slide(sn); x = open(fp, encoding='utf8').read()
    g = re.search(r'<p:pic>.*?<a:off x="(-?\d+)" y="(-?\d+)"/>\s*<a:ext cx="(\d+)" cy="(\d+)"', x, re.S)
    bx, by, bw, bh = (int(v) for v in g.groups())
    side = ('<p:sp><p:nvSpPr><p:cNvPr id="912" name="SideLabel"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="300000" cy="200000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>T1 TSE SAG</a:t></a:r></a:p></p:txBody></p:sp>' % (bx + bw + 500000, by))
    i = x.rindex('</p:spTree>'); open(fp, 'w', encoding='utf8').write(x[:i] + side + x[i:])
    n1 = d.import_slide(src, sn, after=sn, pictures=False)
    assert 'T1 TSE SAG' in ''.join(d.texts(n1))                    # 기본 margin: 옆 라벨 남음
    n2 = d.import_slide(src, sn, after=sn, pictures=False, margin=1.0)
    assert 'T1 TSE SAG' not in ''.join(d.texts(n2))                # margin 키우면 제거

def t_lint_deck_kind_english_only():
    d = T.Deck.open(SRC, wd('dk'), theme='cud')
    d.set_body(7, T.Body().line('Popliteal artery 의 lateral'))
    probs, _ = T.lint(d, stream=io.StringIO())
    assert any('덱 종류 미지정' in p for p in probs)
    probs, _ = T.lint(d, stream=io.StringIO(), deck_kind_override='case review')
    assert any('slide7.xml' in p and '본문에 한글' in p and p.startswith('[참고] 화면 ') for p in probs), probs
    first = [s for s, _, _ in d.order()][0]; d.set_notes(first, ['[deck: 전평 풀이]', 'x'])
    probs, _ = T.lint(d, stream=io.StringIO())
    assert not any('한글' in p for p in probs) and not any('미지정' in p for p in probs)
    r = cli('lint', SRC, '--deck-kind', 'quiz review'); assert r.returncode == 0

def t_mapcheck_sldid_screen_label():
    d = T.Deck.open(SRC, wd('sidl'), theme='cud')
    o = [s for s, _, _ in d.order()]; sid = d.sld_id(o[3])
    cl = [{'id': 'a', 'statement': 's', 'evidence': 'e', 'sites': ['slide@%d' % sid], 'keys': ['zzz-not-there'], 'depends_on': []}]
    buf = io.StringIO(); probs, _ = T.mapcheck(d, cl, buf)
    assert 'slide@%d (화면 4)' % sid in buf.getvalue() and any('(화면 4)' in p for p in probs), buf.getvalue()

def _memo_notes(d, sn):
    """원작자 메모 흉내: 여러 run(굵게·색·하이퍼링크) 문단 + & < > + sldNum 자리."""
    b = d._notes_body(sn)
    if not b:
        d.set_notes(sn, ['']); b = d._notes_body(sn)
    p, x, a, e = b
    memo = ('<a:p><a:r><a:rPr lang="ko-KR" b="1"/><a:t>CT </a:t></a:r><a:r><a:rPr lang="ko-KR"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:rPr>'
            '<a:t>-&gt; 조영 &amp; 비교</a:t></a:r><a:r><a:rPr lang="en-US"><a:hlinkClick r:id="rId9"/></a:rPr><a:t>link</a:t></a:r></a:p>'
            '<a:p><a:r><a:rPr lang="ko-KR"/><a:t>두번째 메모</a:t></a:r></a:p>')
    x = x[:a] + memo + x[e:]
    sld = ('<p:sp><p:nvSpPr><p:cNvPr id="9" name="SN"/><p:cNvSpPr/><p:nvPr><p:ph type="sldNum" idx="5"/></p:nvPr></p:nvSpPr><p:spPr/>'
           '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:fld id="{1}" type="slidenum"><a:t>61</a:t></a:fld></a:p></p:txBody></p:sp>')
    i = x.rindex('</p:spTree>'); x = x[:i] + sld + x[i:]
    open(p, 'w', encoding='utf8').write(x)
    rp = os.path.join(d.dir, 'ppt/notesSlides/_rels/notesSlide%d.xml.rels' % d.notes_no(sn))
    r = open(rp, encoding='utf8').read().replace('</Relationships>', '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.org" TargetMode="External"/></Relationships>')
    open(rp, 'w', encoding='utf8').write(r)
    return memo

def t_notes_paragraph_level_unescaped_no_sldnum():
    # 전평 H&N 회신 N1
    d = T.Deck.open(SRC, wd('n1'))
    sn = [s for s, _, _ in d.order()][2]
    _memo_notes(d, sn)
    assert d.notes(sn) == ['CT -> 조영 & 비교link', '두번째 메모'], d.notes(sn)       # 문단 단위, unescape, sldNum '61' 없음

def t_set_notes_roundtrip_preserves_protected_memo():
    d = T.Deck.open(SRC, wd('n2'))
    sn = [s for s, _, _ in d.order()][2]
    memo = _memo_notes(d, sn)
    assert d.protect_memo(sn) and not d.protect_memo(sn)
    d.set_notes(sn, ['대본 첫 줄', 'second'], tips=['참고 한 줄'])
    x = open(d._notes_path(sn), encoding='utf8').read()
    assert memo in x and '>61<' in x                                               # 메모 바이트·sldNum 그대로
    t = d.notes(sn)
    assert t.index('대본 첫 줄') < t.index(T.NOTES_SEP) < t.index(T.NOTES_SEP_MEMO) < t.index('두번째 메모'), t   # 대본 → 참고 → 기존 메모
    assert T._spoken_notes(d.notes(sn)) == ['대본 첫 줄', 'second']
    d.save('/tmp/n2.pptx'); assert T.validate('/tmp/n2.pptx', SRC)

def t_note_paragraph_edit_api_guards_memo():
    d = T.Deck.open(SRC, wd('n3'))
    sn = [s for s, _, _ in d.order()][2]
    memo = _memo_notes(d, sn); d.protect_memo(sn); d.set_notes(sn, ['alpha line', 'beta line'])
    d.replace_note_paragraph(sn, 'beta', 'BETA line'); d.insert_note_after(sn, 'alpha', ['inserted'])
    t = d.notes(sn); assert t[:3] == ['alpha line', 'inserted', 'BETA line'], t
    try:
        d.replace_note_paragraph(sn, '두번째 메모', 'x'); assert False
    except ValueError as e:
        assert '기존 메모' in str(e)
    assert memo in open(d._notes_path(sn), encoding='utf8').read()

def t_copy_note_paragraphs_remaps_hyperlink_and_diff():
    # H&N 보충: 원본에서 메모 문단을 XML 그대로 복구 + 원본 대비 비교
    orig = T.Deck.open(SRC, wd('n4o')); sn = [s for s, _, _ in orig.order()][2]
    memo = _memo_notes(orig, sn); orig.save('/tmp/n4_orig.pptx')
    ed = T.Deck.open('/tmp/n4_orig.pptx', wd('n4e'))
    ed.set_notes(sn, ['대본만'])                                                   # 보호 없이 덮어씀 → 메모 소실
    r = T.diff_decks('/tmp/n4_orig.pptx', ed, io.StringIO())
    pos = ed.screen_no(sn); assert r['notes'][pos] == (0, 0, 2), r
    n = ed.copy_note_paragraphs(orig, sn, sn)
    assert n == 2
    x = open(ed._notes_path(sn), encoding='utf8').read()
    rels = open(os.path.join(ed.dir, 'ppt/notesSlides/_rels/notesSlide%d.xml.rels' % ed.notes_no(sn)), encoding='utf8').read()
    rid = re.search(r'hlinkClick r:id="(rId\d+)"', x).group(1)
    assert 'Id="%s"' % rid in rels and 'example.org' in rels
    ed.save('/tmp/n4_ed.pptx'); assert T.validate('/tmp/n4_ed.pptx', '/tmp/n4_orig.pptx')
    r = T.diff_decks('/tmp/n4_orig.pptx', '/tmp/n4_ed.pptx', io.StringIO())
    k, f, g = r['notes'][pos]; assert g == 0 and k + f == 2, r['notes'][pos]     # rId 가 바뀐 문단은 '서식 변경'으로 잡힌다
    assert r['sldnum_lost'] == []                                                  # set_notes 는 body 만 바꾸므로 sldNum 자리는 남는다

def t_diff_pairs_by_order_not_path():
    d = T.Deck.open(SRC, wd('n5')); o = [s for s, _, _ in d.order()]
    d.move_slide(o[3], after=0); d.save('/tmp/n5.pptx')
    r = T.diff_decks(SRC, '/tmp/n5.pptx', io.StringIO())
    assert 1 in r['slides_changed'] and len(r['slides_changed']) <= 5, r['slides_changed']

def t_note_substitution_lint_and_safe_replace():
    d = T.Deck.open(SRC, wd('n6'), theme='cud')
    d.set_body(7, T.Body().line('vessel enhancement pattern'))
    d.set_notes(7, ['구ecchymosis 가 있고 막tongue 소견, 과vessel성 병변, vesselenhancement 확인. CT상 정상, enhance되는 결절'])
    p = T.check_note_substitution(d)
    hit = ' '.join(x for x in p if 'slide7' in x)
    for w in ('구ecchymosis', '막tongue', 'vessel성'):
        assert w in hit, (w, hit)
    assert 'vesselenhancement' not in hit          # v16.17: 영어 두 단어 붙음 규칙 (c) 제거 — 오탐이 잦아 치환 쪽 규칙으로 막는다
    assert 'CT상' not in hit and 'enhance되' not in hit, hit
    # 발표 N1: 접두사+어근인 실제 단어는 잡지 않는다 (intra-, epi-, trans- 가 덱 안에 따로 쓰였어도)
    d.set_body(8, T.Body().line('intra axial, trans arterial, epi glottic, cranial formation glottis'))
    d.set_notes(8, ['intracranial extension, epiglottic fold, epiglottis, malignant transformation 입니다.'])
    assert not [x for x in T.check_note_substitution(d) if 'slide8' in x], T.check_note_substitution(d)
    out, skipped = T.safe_replace_terms('구멍이 있고 혀가 막혀 과혈관성, 혈관 조영', {'멍': 'ecchymosis', '혀': 'tongue', '혈관': 'vessel'})
    assert '구멍' in out and '막혀' in out and '과혈관성' in out and 'vessel 조영' in out and 'tongue가' in out, out

def t_phi_by_audience():
    # 2026-09-24 사용자: 내부 발표는 환자 정보를 남긴다, 외부 발표만 지운다
    d = T.Deck.open(SRC, wd('phi_a')); d.set_notes(2, ['환자번호 12345678 확인'])
    first = [s for s, _, _ in d.order()][0]
    assert T.deck_audience(d) is None
    buf = io.StringIO(); hits = T.scan_phi(d, buf); assert hits and '[참고]' in buf.getvalue()          # 미지정: [참고]
    d.set_notes(first, ['[deck: case review]']); assert T.deck_audience(d) == 'internal'
    buf = io.StringIO(); assert T.scan_phi(d, buf, audience='internal') == [] and '내부 발표' in buf.getvalue()
    d.set_notes(first, ['[deck: 학회 구연 · 외부]']); assert T.deck_audience(d) == 'external'
    buf = io.StringIO(); assert T.scan_phi(d, buf, audience='external') and '[!]' in buf.getvalue()
    r = T.polish(T.Deck.open(SRC, wd('phi_b')), stream=io.StringIO(), audience='internal'); assert r['phi'] == []

def t_import_slide_warns_on_theme_mismatch():
    # 작업규약 §4.7 확인: 테마가 다른 덱에서 가져오면 목적지 모양을 입는다 — 알린다
    dst = T.Deck.open(SRC, wd('imw_d')); src = T.Deck.open(SRC, wd('imw_s'))
    dst.import_slide(src, 2, after=1); assert dst.import_warnings == []
    import glob
    for tp in glob.glob(os.path.join(src.dir, 'ppt/theme/theme*.xml')):
        x = open(tp, encoding='utf8').read()
        x = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', r'\g<1>Lato', x, 1)
        open(tp, 'w', encoding='utf8').write(x)
    dst.import_slide(src, 2, after=1)
    assert dst.import_warnings and '원본 서식 유지' in dst.import_warnings[0], dst.import_warnings

def _title_decks(content_top=None):
    """베이스(전평 템플릿 흉내: 회색 띠 = 제목 채움, 28pt·맑은 고딕, 1줄 1117331 / 2줄 1505129) 와
    새 연도 덱(40pt·Arial·1줄 높이 띠에 2줄 제목). content_top 을 주면 새 연도 슬라이드 제목 바로 아래(인치)에 본문 상자를 둔다."""
    from pptx import Presentation
    from pptx.util import Pt, Emu, Inches
    from pptx.dml.color import RGBColor
    def deck(path, specs, top=None):
        prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for text, sz, font, h in specs:
            s = prs.slides.add_slide(prs.slide_layouts[5])
            t = s.shapes.title; t.left, t.top, t.width, t.height = Emu(0), Emu(0), Inches(13.333), Emu(h)
            t.fill.solid(); t.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
            t.text_frame.text = text
            for r in t.text_frame.paragraphs[0].runs:
                r.font.size = Pt(sz); r.font.name = font
            if top is not None:
                b = s.shapes.add_textbox(Inches(0.5), Inches(top), Inches(6), Inches(1)); b.name = 'Body'; b.text_frame.text = 'answer choices'
        prs.save(path); return path
    one, two = 1117331, 1505129
    long_t = 'Q. ' + 'long question title words ' * 4
    base = deck(os.path.join(TMP, 'tb_base.pptx'), [('Q. 24-%02d Colon' % k, 28, '맑은 고딕', one) for k in range(1, 7)] +
                [(long_t + '(24-07)', 28, '맑은 고딕', two), (long_t + '(24-08)', 28, '맑은 고딕', two)])
    new = deck(os.path.join(TMP, 'tb_2025_%s.pptx' % content_top), [(long_t + '(25-11)', 40, 'Arial', 603504), ('Q. 25-12 Rectum', 40, 'Arial', 603504)], content_top)
    return base, new

def t_title_template_learn_check_conform():
    # 2026-09-24 사용자: 템플릿 안 지킨 새 연도 덱을 옮겨 오면 제목이 띠를 넘친다 (도구로든 손 복사로든)
    base, new = _title_decks()
    B = T.Deck.open(base, wd('tb_b')); N = T.Deck.open(new, wd('tb_n'))
    prof = T.title_profile(B, like=1)
    assert prof['sz'] == 2800 and prof['latin'] == '맑은 고딕' and set(prof['h_by_lines'].values()) == {1117331, 1505129} and prof['fill'], prof
    last = [s for s, _, _ in B.order()][-1]
    n1 = B.import_slide(N, 1, after=last); n2 = B.import_slide(N, 2, after=n1)
    scr = [B.screen_no(n1), B.screen_no(n2)]
    bad = T.check_title_template(B, prof, screens=scr, stream=io.StringIO())
    assert any('띠를 넘친다' in x for x in bad[n1]) and any('40pt' in x for x in bad[n1]) and any('Arial' in x for x in bad[n1]), bad
    assert any('1줄 높이' in x for x in bad[n2]), bad                                  # 한 줄 제목: 규격 1줄 띠로
    base_before = {sn: open(B._slide(sn), encoding='utf8').read() for sn in B.slide_numbers() if sn not in (n1, n2)}
    for sn in (n1, n2):
        ch = T.conform_title(B, sn, prof); assert ch and not ch[0].startswith('[!]'), ch
    after = T.check_title_template(B, prof, screens=scr, stream=io.StringIO())
    assert not any(x.startswith('[!]') for v in after.values() for x in v), after
    x = open(B._slide(n1), encoding='utf8').read()
    assert 'cy="1505129"' in x and 'sz="2800"' in x and '맑은 고딕' in x and 'Arial' not in x and '25-11' in x
    assert all(open(B._slide(sn), encoding='utf8').read() == v for sn, v in base_before.items())   # 베이스 슬라이드는 그대로
    B.save('/tmp/tb_merged.pptx'); assert T.validate('/tmp/tb_merged.pptx', base)

def t_conform_title_fits_font_when_band_would_cover_content():
    # 실물 전평 덱 재현(2026-09-24): 새 연도 슬라이드는 얇은 띠 바로 아래 본문이 있어 규격 띠로 키우면 본문을 덮는다 →
    # 띠 자리는 그대로, 글자를 줄여 띠 안에 넣는다(§0-A-2). 본문은 옮기지 않는다
    base, new = _title_decks(content_top=0.8)
    B = T.Deck.open(base, wd('tb_fb')); N = T.Deck.open(new, wd('tb_fn'))
    prof = T.title_profile(B, like=1)
    n1 = B.import_slide(N, 1, after=[s for s, _, _ in B.order()][-1])
    body_before = re.search(r'<p:sp>(?:(?!</p:sp>).)*?name="Body".*?</p:sp>', open(B._slide(n1), encoding='utf8').read(), re.S).group(0)
    ch = T.conform_title(B, n1, prof)
    # v16.27 (발표 근골격): 줄이는 하한이 lint 제목 최소 24pt — 이 fixture 는 24pt 밑으로 줄여야 들어가므로 바꾸지 않고 [!]
    assert ch[0].startswith('[!]') and '겹쳐' in ch[0] and '24pt 밑으로' in ch[0], ch
    x = open(B._slide(n1), encoding='utf8').read()
    assert 'cy="603504"' in x and body_before in x                        # 띠 높이·본문 그대로
    B.save('/tmp/tb_fit.pptx'); assert T.validate('/tmp/tb_fit.pptx', base)

def t_import_slide_maps_layout_by_name_and_keeps_positions():
    # 실물 전평 덱(마스터 7개) 재현: 레이아웃 파일 이름이 같아도 뜻이 달라 본문이 그림 위로 올라갔다
    base, new = _title_decks()
    B = T.Deck.open(base, wd('tb_lm')); N = T.Deck.open(new, wd('tb_ln'))
    src_lay = N.layout_of(1)
    n1 = B.import_slide(N, 1, after=1, title_profile=T.title_profile(B, like=1))
    assert T._layout_name(B, B.layout_of(n1)) == T._layout_name(N, src_lay)
    x = open(B._slide(n1), encoding='utf8').read()
    assert all(T._GEO.search(m.group(0)) for m in re.finditer(r'<p:sp>(?:(?!</p:sp>).)*?<p:ph\b.*?</p:sp>', x, re.S))   # 상속 위치를 적어 넣음

def t_titles_cli():
    base, new = _title_decks()
    C = T.Deck.open(base, wd('tb_c')); C.import_slide(T.Deck.open(new, wd('tb_n3')), 1, after=[s for s, _, _ in C.order()][-1]); C.save('/tmp/tb_hand.pptx')
    r = cli('titles', '/tmp/tb_hand.pptx'); assert r.returncode == 1 and '띠를 넘친다' in r.stdout, r.stdout
    r = cli('titles', '/tmp/tb_hand.pptx', '--apply'); assert r.returncode != 0
    r = cli('titles', '/tmp/tb_hand.pptx', '--like', '1', '--screens', '9', '--apply', '-o', '/tmp/tb_hand_fixed.pptx')
    assert r.returncode == 0 and 'verify --original: 통과' in r.stdout, r.stdout

def t_notes_on_deck_without_notes_master():
    # v16.9 에서 발견: 노트가 하나도 없던 덱에 노트를 만들면 없는 notesMaster1.xml 을 가리켜 파일이 깨졌다
    base, _ = _title_decks()
    B = T.Deck.open(base, wd('nm0')); assert not os.path.exists(os.path.join(B.dir, 'ppt/notesMasters/notesMaster1.xml'))
    sn = [s for s, _, _ in B.order()][0]
    B.set_notes(sn, ['script'], tips=['tip'])
    assert B.notes(sn)[0] == 'script'
    B.save('/tmp/nm0.pptx'); assert T.validate('/tmp/nm0.pptx', base)

def t_import_slide_keeps_body_font_size():
    # 영상의학 회신 09-24: 받는 덱 마스터의 큰 본문 글자를 입어 25-11 정답 보기 ㉣ 가 그림 뒤로 가려졌다 → 원천 크기를 적어 넣는다
    src = T.Deck.open(SRC, wd('bs_s')); dst = T.Deck.open(SRC, wd('bs_d'))
    import glob
    for mp in glob.glob(os.path.join(dst.dir, 'ppt/slideMasters/slideMaster*.xml')):
        x = open(mp, encoding='utf8').read()
        x = re.sub(r'(<p:bodyStyle>.*?<a:lvl1pPr\b.*?<a:defRPr\b[^>]*\bsz=")\d+', r'\g<1>4400', x, 1, flags=re.S)
        open(mp, 'w', encoding='utf8').write(x)
    sn = next(s for s in src.slide_numbers() if re.search(r'<p:ph\b[^>]*idx="1"', open(src._slide(s), encoding='utf8').read()) and src.texts(s))
    want = T._body_level_sizes(src, sn, 'idx="1"').get(0)
    n = dst.import_slide(src, sn, after=1)
    x = open(dst._slide(n), encoding='utf8').read()
    body = re.search(r'<p:sp>(?:(?!</p:sp>).)*?<p:ph\b[^>]*idx="1".*?</p:sp>', x, re.S).group(0)
    assert want and want != 4400 and 'sz="%d"' % want in body and 'sz="4400"' not in body, (want, body[:400])
    assert dst.import_warnings and '크기' in dst.import_warnings[-1]

def t_diff_with_screen_map_and_text_match():
    # 영상의학 회신 09-24 §4: 순서를 바꾸고 다른 덱을 섞은 편집본(H&N v3)은 매핑표로 짝짓는다
    A = T.Deck.open(SRC, wd('dm_a')); sn = [s for s, _, _ in A.order()][2]
    _memo_notes(A, sn); A.save('/tmp/dm_orig.pptx')
    base, new = _title_decks()
    E = T.Deck.open('/tmp/dm_orig.pptx', wd('dm_e')); o = [s for s, _, _ in E.order()]
    E.move_slide(o[2], after=0)                                   # 원본 화면 3 → 편집 화면 1
    n = E.import_slide(T.Deck.open(new, wd('dm_y')), 2, after=o[0])   # 새 연도 덱 화면 2 → 편집 화면 3
    E.set_body(o[5], T.Body().line('changed body'))               # 의도하지 않은 본문 변경
    E.save('/tmp/dm_ed.pptx')
    ob = [s for s, _, _ in E.order()]
    mp = os.path.join(TMP, 'dm_map.md')
    rows = ['| 편집 | 원본 | 제목 | 의도한 수정 |', '|---|---|---|---|', '| 1 | 원본 3 | | |', '| 2 | 원본 1 | | |',
            '| 3 | 새 덱 2 | | 제목 줄임 |', '| %d | 원본 6 | | |' % (ob.index(o[5]) + 1)]
    open(mp, 'w', encoding='utf8').write('\n'.join(rows))
    m = T.parse_screen_map(mp); assert m[2] == (3, '새 덱', 2, '제목 줄임'), m
    r = T.diff_decks('/tmp/dm_orig.pptx', '/tmp/dm_ed.pptx', io.StringIO(), mapping=m, sources={'새 덱': new})
    assert r['notes'][1] == (2, 0, 0) and r['unmatched'] == [] and r['unintended'] == [ob.index(o[5]) + 1], r
    r = T.diff_decks('/tmp/dm_orig.pptx', '/tmp/dm_ed.pptx', io.StringIO(), match_text=True)
    assert r['notes'].get(1) == (2, 0, 0) and 3 in r['unmatched'], r
    rr = cli('diff', '/tmp/dm_orig.pptx', '/tmp/dm_ed.pptx', '--map', mp, '--src', '새 덱=' + new)
    assert '매핑표로 짝지음' in rr.stdout and rr.returncode == 1, rr.stdout   # 의도하지 않은 변경 1곳

def t_restore_memo_old_separator():
    # 영상의학 회신 §5: v1 대본 덱(옛 표지 '────── 기존 메모 ──────' + .text 로 덮어써 서식 잃은 메모) 되살리기
    A = T.Deck.open(SRC, wd('rm_a')); sn = [s for s, _, _ in A.order()][2]
    memo = _memo_notes(A, sn); A.save('/tmp/rm_orig.pptx')
    E = T.Deck.open('/tmp/rm_orig.pptx', wd('rm_e'))
    E.set_notes(sn, ['대본 한 줄', '', '────── 기존 메모 ──────', 'CT -> 조영 & 비교link', '두번째 메모'])
    assert T._is_memo_sep('────── 기존 메모 ──────') and T._spoken_notes(E.notes(sn)) == ['대본 한 줄']
    E.save('/tmp/rm_ed.pptx')
    r = T.diff_decks('/tmp/rm_orig.pptx', '/tmp/rm_ed.pptx', io.StringIO()); pos = E.screen_no(sn)
    assert r['notes'][pos][0] == 0, r['notes'][pos]                                    # 서식 잃음
    rr = cli('restore-memo', '/tmp/rm_ed.pptx', '--original', '/tmp/rm_orig.pptx', '-o', '/tmp/rm_fixed.pptx')
    assert rr.returncode == 0 and '사라짐 0' in rr.stdout and 'verify --original: 통과' in rr.stdout, rr.stdout
    F = T.Deck.open('/tmp/rm_fixed.pptx', wd('rm_f'))
    assert F.notes(sn)[0] == '대본 한 줄' and '────── 기존 메모 ──────' in F.notes(sn)
    r = T.diff_decks('/tmp/rm_orig.pptx', '/tmp/rm_fixed.pptx', io.StringIO())
    k, f, g = r['notes'][pos]; assert g == 0 and k >= 1, r['notes'][pos]

def t_normalize_notes_moves_tags_and_separator():
    # 사용자 09-24: 영상의학 v1 대본 덱 — 낭독 부분의 태그를 참고로, 옛 표지를 새 표지로. 기존 메모 바이트는 그대로
    d = T.Deck.open(SRC, wd('nn')); sn = [s for s, _, _ in d.order()][2]
    memo = _memo_notes(d, sn)                                  # 원작자 메모(서식 있음)
    p, x, a, e = d._notes_body(sn)
    head = T.notes_xml(['첫 문장이다. 두번째 문장이다 [검증: Som 5th].', '[색인] H&N_index.md 12쪽', '셋째 문장.', '', '────── 기존 메모 ──────'])
    open(p, 'w', encoding='utf8').write(x[:a] + head + x[a:])
    r = T.normalize_notes(d, sn)
    assert r['moved'] == 2 and r['sep'] == 1, r
    t = d.notes(sn)
    spoken = T._spoken_notes(t)
    assert spoken == ['첫 문장이다. 두번째 문장이다.', '셋째 문장.'], spoken
    assert t.index(T.NOTES_SEP) < t.index('- [검증: Som 5th] 두번째 문장이다.') < t.index(T.NOTES_SEP_MEMO), t
    assert '- [색인] H&N_index.md 12쪽' in t and '────── 기존 메모 ──────' not in t
    assert memo in open(d._notes_path(sn), encoding='utf8').read()          # 메모 바이트 그대로
    assert T.normalize_notes(d, sn) == {'moved': 0, 'sep': 0, 'samples': []}   # 두 번 돌려도 그대로
    d.save('/tmp/nn.pptx'); assert T.validate('/tmp/nn.pptx', SRC)
    r = cli('normalize-notes', '/tmp/nn.pptx', '--dry-run'); assert r.returncode == 0 and '낭독 부분에 남은 태그: 없음' in r.stdout, r.stdout

def t_restore_memo_match_text_multi_source():
    # LGI 처럼 원본이 없을 때: 처음 받은 두 덱(베이스·새 연도)에서 슬라이드 글로 원천 화면을 찾아 메모를 되살린다
    A = T.Deck.open(SRC, wd('mt_a')); sn = [s for s, _, _ in A.order()][2]; _memo_notes(A, sn); A.save('/tmp/mt_a.pptx')
    base, new = _title_decks()
    E = T.Deck.open('/tmp/mt_a.pptx', wd('mt_e'))
    E.import_slide(T.Deck.open(new, wd('mt_y')), 1, after=[s for s, _, _ in E.order()][-1])
    E.set_notes(sn, ['대본', '', '────── 기존 메모 ──────', 'plain memo'])
    E.save('/tmp/mt_e.pptx')
    pairs, miss = T.match_sources_by_text(T.Deck.open('/tmp/mt_e.pptx', wd('mt_e2')), [T.Deck.open('/tmp/mt_a.pptx', wd('mt_a2')), T.Deck.open(new, wd('mt_y2'))])
    assert not miss and len(pairs) == len(E.order()), (miss, len(pairs))
    r = cli('restore-memo', '/tmp/mt_e.pptx', '--original', '/tmp/mt_a.pptx', '--src', new, '--match-text', '-o', '/tmp/mt_fixed.pptx')
    assert r.returncode == 0 and '짝 없음 0화면' in r.stdout and 'verify --original: 통과' in r.stdout, r.stdout
    F = T.Deck.open('/tmp/mt_fixed.pptx', wd('mt_f')); assert 'CT -> 조영 & 비교link' in F.notes(sn) and 'plain memo' not in F.notes(sn)

def t_apply_fixes_table():
    d = T.Deck.open(SRC, wd('af'), theme='cud')
    d.set_body(7, T.Body().line('Parathyoid adenoma').line('tympanic membran rupture').line('dup word dup word'))
    d.save('/tmp/af.pptx')
    scr = T.Deck.open('/tmp/af.pptx', wd('af2')).screen_no(7)
    fx = os.path.join(TMP, 'fixes.md')
    open(fx, 'w', encoding='utf8').write('## 본문 수정\n| 화면 | 원문 | 수정문 | 근거 |\n|---|---|---|---|\n'
        '| %d | Parathyoid | Parathyroid | 오기 |\n| %d | membran rupture | membrane rupture | 오기 |\n| %d | dup word | x | |\n' % (scr, scr, scr))
    assert T.parse_fixes(fx)[0] == (scr, 'Parathyoid', 'Parathyroid')
    r = cli('apply-fixes', '/tmp/af.pptx', '--fixes', fx, '-o', '/tmp/af_fixed.pptx')
    assert r.returncode == 1 and '적용 2 / 거부 1' in r.stdout and '매치 2회' in r.stdout and 'verify --original: 통과' in r.stdout, r.stdout
    t = ' '.join(T.Deck.open('/tmp/af_fixed.pptx', wd('af3')).texts(7))
    assert 'Parathyroid adenoma' in t and 'tympanic membrane rupture' in t

def t_v1692_temp_cleanup_protect_memo_verify_eomi_memo_only():
    import tempfile, glob
    # Y1: diff 가 경로로 연 덱의 추출 폴더를 지운다
    before = set(glob.glob(os.path.join(tempfile.gettempdir(), 'dk_*')))
    T.diff_decks(SRC, SRC, io.StringIO())
    assert set(glob.glob(os.path.join(tempfile.gettempdir(), 'dk_*'))) == before
    # Y2: 이미 3부 구조인 노트는 protect_memo 가 건드리지 않는다
    d = T.Deck.open(SRC, wd('y2')); sn = [s for s, _, _ in d.order()][2]
    d.set_notes(sn, ['우리 대본'], tips=['참고'])
    assert d.protect_memo(sn) is False and T.NOTES_SEP_MEMO not in d.notes(sn)
    # Y3: verify 통과 출력
    r = cli('verify', SRC); assert r.returncode == 0 and 'verify: 통과' in r.stdout, r.stdout
    # X2: '영어 동사 + 됩니다' 는 치환 흔적이 아니다
    d.set_notes(7, ['병변이 enhance됩니다. resorption됩니다. deposition됐다.'])
    assert not [x for x in T.check_note_substitution(d) if 'slide7' in x]
    # memo_only: 판끼리 비교에서 대본을 다시 써도 메모 구역만 센다
    A = T.Deck.open(SRC, wd('mo_a')); s2 = [s for s, _, _ in A.order()][2]; _memo_notes(A, s2); A.protect_memo(s2)
    A.set_notes(s2, ['옛 대본'], tips=['옛 참고']); A.save('/tmp/mo_v3.pptx')
    B = T.Deck.open('/tmp/mo_v3.pptx', wd('mo_b')); B.set_notes(s2, ['새 대본'], tips=['새 참고']); B.save('/tmp/mo_v5.pptx')
    pos = B.screen_no(s2)
    assert T.diff_decks('/tmp/mo_v3.pptx', '/tmp/mo_v5.pptx', io.StringIO())['notes'][pos][2] >= 2
    assert T.diff_decks('/tmp/mo_v3.pptx', '/tmp/mo_v5.pptx', io.StringIO(), memo_only=True)['notes'][pos] == (2, 0, 0)
    # W4 (v16.10): 기존 메모 표지가 없는 화면은 memo_only 에서 세지 않는다
    other = [s for s, _, _ in B.order()][4]; B.set_notes(other, ['대본만']); B.save('/tmp/mo_v6.pptx')
    r = T.diff_decks('/tmp/mo_v5.pptx', '/tmp/mo_v6.pptx', io.StringIO(), memo_only=True)
    assert B.screen_no(other) not in r['notes'], r['notes']

def t_v1610_open_tmp_numcol_bake_bg_memo_only():
    import tempfile
    # X1: 기본 풀기 자리는 임시 폴더(읽기 전용 폴더의 파일도 연다)
    d = T.Deck.open(SRC); assert d.dir.startswith(tempfile.gettempdir()), d.dir
    ro = tempfile.mkdtemp(); shutil.copy(SRC, ro); os.chmod(ro, 0o555)
    try:
        src_ro = os.path.join(ro, os.path.basename(SRC))
        if not os.access(ro, os.W_OK):                     # root 로 돌면 chmod 가 안 먹으므로 그때는 건너뛴다
            assert os.path.dirname(T._default_out(src_ro, '_x')) != ro
            r = cli('lint', src_ro); assert r.returncode == 0 and 'Read-only' not in r.stderr, r.stderr[-300:]
    finally:
        os.chmod(ro, 0o755)
    # W2: 두 단 상자는 필요 높이를 단 수로 나눈다
    sh = {'w': 6 * 914400, 'paras': [{'text': 'item %d short' % k, 'sz': 18, 'aft': 0} for k in range(30)], 'numCol': 1, 'spcCol': 0}   # 짧은 항목 30줄(교육목표 목록)
    h1 = T._estimated_height(sh); h2 = T._estimated_height(dict(sh, numCol=2))
    assert h2 < h1 * 0.75, (h1, h2)
    # W1-3: bake_autofit — 비율을 실제 크기·줄 간격으로
    d = T.Deck.open(SRC, wd('bake'), theme='cud')
    d.set_body(7, T.Body().line('alpha one').line('beta two'))
    fp = d._slide(7); x = open(fp, encoding='utf8').read()
    b = re.search(r'<p:sp>(?:(?!</p:sp>).)*?<p:ph\b[^>]*idx="1".*?</p:sp>', x, re.S).group(0)
    b2 = re.sub(r'<a:bodyPr([^>]*)/>', r'<a:bodyPr\1><a:normAutofit fontScale="70000" lnSpcReduction="20000"/></a:bodyPr>', b, 1)
    open(fp, 'w', encoding='utf8').write(x.replace(b, b2))
    szs = [int(v) for v in re.findall(r'<a:rPr\b[^>]*\bsz="(\d+)"', b2)]
    assert d.bake_autofit(7) == 1
    y = open(fp, encoding='utf8').read(); bb = re.search(r'<p:sp>(?:(?!</p:sp>).)*?<p:ph\b[^>]*idx="1".*?</p:sp>', y, re.S).group(0)
    assert 'fontScale' not in bb and 'lnSpcReduction' not in bb and '<a:normAutofit/>' in bb
    assert int(re.findall(r'<a:rPr\b[^>]*\bsz="(\d+)"', bb)[0]) == int(szs[0] * 0.7), (szs, bb[:300])
    assert '<a:lnSpc><a:spcPct val="80000"/></a:lnSpc>' in bb
    d.save('/tmp/bake.pptx'); assert T.validate('/tmp/bake.pptx', SRC)
    assert T.no_autofit_copy('/tmp/bake.pptx', '/tmp/bake_na.pptx') >= 1
    # W3: 배경
    assert d.set_background(7, rgb='D2F6F6') is False and 'srgbClr val="D2F6F6"' in open(fp, encoding='utf8').read()
    assert d.set_background(7) is True and '<p:bg>' not in open(fp, encoding='utf8').read()
    d.save('/tmp/bg.pptx'); assert T.validate('/tmp/bg.pptx', SRC)

def t_v1610_import_drops_bg_on_theme_mismatch_and_ambiguous_match():
    import glob
    src = T.Deck.open(SRC, wd('zb_s')); dst = T.Deck.open(SRC, wd('zb_d'))
    for tp in glob.glob(os.path.join(src.dir, 'ppt/theme/theme*.xml')):
        tx = open(tp, encoding='utf8').read(); open(tp, 'w', encoding='utf8').write(re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', r'\g<1>Lato', tx, 1))
    src.set_background(2, scheme='accent4')
    n1 = dst.import_slide(src, 2, after=1); assert '<p:bg>' not in open(dst._slide(n1), encoding='utf8').read()
    n2 = dst.import_slide(src, 2, after=1, keep_bg=True); assert 'accent4' in open(dst._slide(n2), encoding='utf8').read()
    # Z3: 글 없는 화면이 여럿이면 짝짓지 않고 '모호'
    A = T.Deck.open(SRC, wd('zb_a')); o = [s for s, _, _ in A.order()]
    for s in o[3:6]:
        A.set_body(s, T.Body().line(' ')); A.set_title(s, [' '])
    A.save('/tmp/zb_a.pptx')
    r = T.diff_decks('/tmp/zb_a.pptx', '/tmp/zb_a.pptx', io.StringIO(), match_text=True)
    assert set(r['ambiguous']) >= {4, 5, 6} and r['unmatched'] == [], r

def t_v1610_set_body_like_notation():
    d = T.Deck.open(SRC, wd('bl'), theme='cud')
    fp = d._slide(7); x = open(fp, encoding='utf8').read()
    tplb = ('<a:p><a:pPr marL="0" indent="0"><a:buNone/></a:pPr><a:r><a:rPr lang="en-US" sz="1800" b="1"><a:latin typeface="Arial"/></a:rPr><a:t>Head</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"><a:buFont typeface="Arial"/><a:buChar char="•"/></a:pPr><a:r><a:rPr lang="en-US" sz="1600"/><a:t>item</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"><a:tabLst><a:tab pos="4000000" algn="l"/></a:tabLst></a:pPr><a:r><a:rPr lang="en-US" sz="1600"/><a:t>item\t</a:t></a:r>'
            '<a:r><a:rPr lang="en-US" sz="1600"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill></a:rPr><a:t>25-11</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="2"/><a:r><a:rPr lang="en-US" sz="1400"/><a:t>sub</a:t></a:r></a:p>')
    p, dd, a, e = d._body_span(7); open(p, 'w', encoding='utf8').write(dd[:a] + tplb + dd[e:])
    lines = ['L0 **Salivary gland**', 'L1 Sialolithiasis', 'L1 **Pleomorphic adenoma** ⇥ {r:[짤]25-11}, 22-14', 'L2 MRI: T2 bright']
    assert T.parse_body_notation(lines)[2] == {'lvl': 1, 'runs': [('Pleomorphic adenoma', True, False), ('\t', False, False), ('[짤]25-11', False, True), (', 22-14', False, False)], 'tab': True}
    assert d.set_body_like(7, lines) == 4
    y = open(fp, encoding='utf8').read()
    ps = re.findall(r'<a:p>.*?</a:p>', y[d._body_span(7)[2]:d._body_span(7)[3]], re.S)
    assert 'marL="0"' in ps[0] and 'b="1"' in ps[0] and 'Arial' in ps[0]
    assert 'lvl="1"' in ps[1] and 'buChar' in ps[1] and 'b="1"' not in ps[1]
    assert 'tabLst' in ps[2] and 'val="C00000"' in ps[2] and '\t' in ps[2] and re.search(r'b="1"[^>]*>(?:(?!</a:r>).)*Pleomorphic', ps[2], re.S)
    assert 'lvl="2"' in ps[3]
    d.save('/tmp/bl.pptx'); assert T.validate('/tmp/bl.pptx', SRC)

def t_v1611_replace_paragraph_like_and_body_notation():
    d = T.Deck.open(SRC, wd('rpl'), theme='cud')
    fp = d._slide(7)
    body = ('<a:p><a:pPr marL="0"><a:buNone/></a:pPr><a:r><a:rPr lang="en-US" sz="1800" b="1" err="1"/><a:t>Head</a:t></a:r><a:r><a:rPr lang="en-US" sz="1800" b="1"/><a:t> </a:t></a:r><a:r><a:rPr lang="en-US" sz="1800" b="0"/><a:t>tail</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"><a:tabLst><a:tab pos="4000000" algn="l"/></a:tabLst></a:pPr><a:r><a:rPr lang="en-US" sz="1600" b="1"/><a:t>Genioglossus m.</a:t></a:r>'
            '<a:r><a:rPr lang="en-US" sz="1600" b="0"/><a:t>\t</a:t></a:r><a:r><a:rPr lang="en-US" sz="1600" b="0"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:rPr><a:t>[짤]24-15</a:t></a:r>'
            '<a:r><a:rPr lang="en-US" sz="1600" b="0"/><a:t>, 19-14 &amp; x</a:t></a:r><a:endParaRPr lang="en-US" sz="1600"/></a:p>'
            '<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1600" b="1"/><a:t>Mylohyoid m.</a:t></a:r><a:r><a:rPr lang="en-US" sz="1600" b="0"/><a:t> 19-14</a:t></a:r></a:p>')
    p, x, a, e = d._body_span(7); open(p, 'w', encoding='utf8').write(x[:a] + body + x[e:])
    # V2: 본문 → 표기 (굵은 공백 run 은 서식 없이, 탭 ⇥, escape 해제)
    nt = d.body_notation(7)
    assert nt == ['L0 **Head** tail', 'L1 **Genioglossus m.**⇥ {r:[짤]24-15}, 19-14 & x', 'L1 **Mylohyoid m.** 19-14'], nt
    assert '****' not in ''.join(nt)
    # 왕복: 표기 → 파서 → 같은 글
    rt = T.parse_body_notation(nt)
    assert ''.join(t for t, _, _ in rt[1]['runs']) == 'Genioglossus m.\t[짤]24-15, 19-14 & x'
    # V1: 한 문단만 바꾸고 나머지는 바이트 그대로
    before = open(fp, encoding='utf8').read()
    keep0 = re.findall(r'<a:p>.*?</a:p>', before[a:], re.S)[0]; keep2 = re.findall(r'<a:p>.*?</a:p>', before[a:], re.S)[2]
    old = d.replace_paragraph_like(7, 'Genioglossus', 'L1 **Floor of mouth muscles (M-H-G)** ⇥ {r:[짤]24-15}, 19-14, 16-13')
    assert old.startswith('Genioglossus')
    after = open(fp, encoding='utf8').read()
    assert keep0 in after and keep2 in after and 'err="1"' in after                     # 다른 문단 바이트 그대로
    pp = [q for q in re.findall(r'<a:p>.*?</a:p>', after, re.S) if 'Floor of mouth' in q][0]
    assert 'tabLst' in pp and 'b="1"' in pp and 'val="FF0000"' in pp and '<a:endParaRPr' in pp and '\t' in pp
    assert re.search(r'b="0"[^>]*>(?:(?!</a:r>).)*16-13', pp, re.S)                   # 보통 run 은 원래 b="0" 그대로
    d.delete_paragraph(7, 'Mylohyoid')
    assert d.body_notation(7)[1] == 'L1 **Floor of mouth muscles (M-H-G)**⇥ {r:[짤]24-15}, 19-14, 16-13', d.body_notation(7)
    d.save('/tmp/rpl.pptx'); assert T.validate('/tmp/rpl.pptx', SRC)

def t_v1612_split_notes_excludes_memo_roundtrip():
    # 발표 U1 (급함): split_notes 가 메모 구역을 참고로 돌려줘 set_notes(n, *split_notes(n)) 가 메모를 두 번 넣었다
    d = T.Deck.open(SRC, wd('u1')); sn = [s for s, _, _ in d.order()][2]
    memo = _memo_notes(d, sn); d.protect_memo(sn)
    before = open(d._notes_path(sn), encoding='utf8').read()
    assert d.split_notes(sn) == ([], []) and d.notes_sections(sn)[2] == ['CT -> 조영 & 비교link', '두번째 메모']
    d.set_notes(sn, *d.split_notes(sn))                                     # 메모만 있는 노트: 무변경
    assert open(d._notes_path(sn), encoding='utf8').read() == before
    d.set_notes(sn, ['대본'], tips=['참고 한 줄'])
    before = open(d._notes_path(sn), encoding='utf8').read()
    assert d.split_notes(sn) == (['대본'], ['참고 한 줄'])
    d.set_notes(sn, *d.split_notes(sn))                                     # 대본·참고·메모: 무변경
    assert open(d._notes_path(sn), encoding='utf8').read() == before and before.count('두번째 메모') == 1
    d.set_notes(sn, ['새 대본'], d.split_notes(sn)[1])                       # 대본만 바꾸고 참고는 그대로
    assert d.notes(sn).count('두번째 메모') == 1 and d.notes_sections(sn) == (['새 대본'], ['참고 한 줄'], ['CT -> 조영 & 비교link', '두번째 메모'])

def t_v1612_rpr_with_nested_scheme_fill():
    r = '<a:rPr lang="ko-KR" b="0"><a:solidFill><a:schemeClr val="tx1"><a:lumMod val="75000"/><a:lumOff val="25000"/></a:schemeClr></a:solidFill><a:latin typeface="+mn-lt"/></a:rPr>'
    out = T._rpr_with(r, True, '<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>')
    assert out.count('<a:solidFill>') == out.count('</a:solidFill>') == 1 and 'lumMod' not in out and 'latin' in out and 'b="1"' in out, out
    out2 = T._rpr_with(r, False, None)
    assert 'lumMod' in out2 and out2.count('<a:solidFill>') == out2.count('</a:solidFill>') == 1 and 'latin' in out2, out2   # 원래 글자색 유지
    red = '<a:rPr lang="ko-KR"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:rPr>'
    assert 'FF0000' not in T._rpr_with(red, False, None)

def t_v1612_height_model_lnspc_spcbef_and_eomi():
    base = {'w': 6 * 914400, 'numCol': 1, 'spcCol': 0}
    p0 = [{'text': 'line %d' % k, 'sz': 18, 'aft': 0} for k in range(10)]
    h0 = T._estimated_height(dict(base, paras=p0))
    h1 = T._estimated_height(dict(base, paras=[dict(q, bef=6) for q in p0]))
    h2 = T._estimated_height(dict(base, paras=[dict(q, lnspc=0.9) for q in p0]))
    assert h1 > h0 > h2
    d = T.Deck.open(SRC, wd('u3'), theme='cud')
    d.set_notes(7, ['inferior canaliculus였습니다. schwannoma이었다.'])
    assert not [x for x in T.check_note_substitution(d) if 'slide7' in x]
    assert isinstance(T.theme_fonts_missing(SRC), list)

def t_v1613_overflow_uses_theme_font_file_and_model_tolerance():
    import glob
    d = T.Deck.open(SRC, wd('tf'))
    # 테마 본문 글꼴을 컨테이너에 있는 DejaVu Sans 로 → 자동으로 그 파일을 쓴다
    for tp in glob.glob(os.path.join(d.dir, 'ppt/theme/theme*.xml')):
        x = open(tp, encoding='utf8').read()
        x = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*', r'\g<1>DejaVu Sans', x, 1)
        open(tp, 'w', encoding='utf8').write(x)
    sn = [s for s, _, _ in d.order()][0]
    fpth = T._theme_body_font_file(d, sn)
    assert fpth and 'DejaVuSans' in fpth, fpth
    buf = io.StringIO(); T.check_text_overflow(d, stream=buf)
    assert 'DejaVuSans' in buf.getvalue() and '덱 테마 글꼴 자동' in buf.getvalue(), buf.getvalue()[-400:]
    # 글꼴 파일이 없으면 모델로 계산하고, 상자 높이 5% 이내의 슬라이드 밖 넘침은 [참고]
    e = T.Deck.open(SRC, wd('tf2'))
    for tp in glob.glob(os.path.join(e.dir, 'ppt/theme/theme*.xml')):
        x = open(tp, encoding='utf8').read()
        x = re.sub(r'<a:minorFont><a:latin typeface="[^"]*"', '<a:minorFont><a:latin typeface="NoSuchFont Zz"', x, 1)
        x = re.sub(r'(<a:minorFont>.*?<a:ea typeface=")[^"]*', r'\g<1>NoSuchFont Zz', x, 1, flags=re.S)
        open(tp, 'w', encoding='utf8').write(x)
    s0 = [s for s, _, _ in e.order()][0]
    assert T._theme_body_font_file(e, s0) is None
    from pptx.util import Emu
    fp = e._slide(7); x = open(fp, encoding='utf8').read()
    W, H = e.slide_size()
    box = ('<p:sp><p:nvSpPr><p:cNvPr id="990" name="LowBox"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="457200" y="%d"/>'
           '<a:ext cx="4000000" cy="4000000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr wrap="square"/><a:lstStyle/>'
           + ''.join('<a:p><a:r><a:rPr lang="en-US" sz="1800"/><a:t>item %d</a:t></a:r></a:p>' % k for k in range(12)) + '</p:txBody></p:sp>')
    sh = None
    for top in range(H - 4000000, H, 20000):          # 모델 넘침이 상자 높이의 1–5% 가 되는 자리를 찾는다
        y = x[:x.rindex('</p:spTree>')] + box % top + x[x.rindex('</p:spTree>'):]
        open(fp, 'w', encoding='utf8').write(y)
        sh = [q for q in T._text_shapes(e, 7) if q['name'] == 'LowBox'][0]
        over = top + T._estimated_height(sh) - H
        if 0.01 * sh['h'] < over < 0.05 * sh['h']:
            break
    probs = T.check_text_overflow(e, stream=io.StringIO())
    mine = [p for p in probs if 'LowBox' in p and '슬라이드 밖' in p]
    assert mine and all(p.startswith('[참고]') and '글꼴 폭 모델' in p for p in mine), mine

def t_v1614_layout_insert_like_spacing_bold():
    # S1: 같은 덱 복제는 원천 레이아웃 유지, layout= 명시, set_layout 은 자리 고정
    d = T.Deck.open(SRC, wd('s1')); s = T.Deck.open(SRC, wd('s1s'))
    o = [x for x, _, _ in d.order()]
    other = next(f for f in sorted(os.listdir(os.path.join(s.dir, 'ppt/slideLayouts'))) if f.endswith('.xml') and f != d.layout_of(o[0]))
    src_no = o[2]; s.set_layout(src_no, other)                      # 원천 슬라이드를 다른 레이아웃으로
    n = d.import_slide(s, src_no, after=o[0])
    assert d.layout_of(n) == s.layout_of(src_no), (d.layout_of(n), s.layout_of(src_no))
    n2 = d.import_slide(s, src_no, after=o[0], layout=d.layout_of(o[0])); assert d.layout_of(n2) == d.layout_of(o[0])
    old = d.set_layout(n, d.layout_of(o[0])); assert old == s.layout_of(src_no) and d.layout_of(n) == d.layout_of(o[0])
    assert all(T._GEO.search(m.group(0)) for m in re.finditer(r'<p:sp>(?:(?!</p:sp>).)*?<p:ph\b(?![^>]*type="(?:dt|ftr|sldNum)").*?</p:sp>', open(d._slide(n), encoding='utf8').read(), re.S))
    d.save('/tmp/s1.pptx'); assert T.validate('/tmp/s1.pptx', SRC)
    # T1: 이웃 문단 서식으로 넣기 — 다른 문단 바이트 그대로
    e = T.Deck.open(SRC, wd('t1'))
    body = ('<a:p><a:pPr lvl="1"><a:spcBef><a:spcPts val="600"/></a:spcBef><a:buChar char="-"/></a:pPr><a:r><a:rPr lang="en-US" sz="1600" b="1"/><a:t>Item A</a:t></a:r><a:r><a:rPr lang="en-US" sz="1600"/><a:t> 22-01</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="2"><a:buNone/></a:pPr><a:r><a:rPr lang="en-US" sz="1200" i="1"/><a:t>* footnote</a:t></a:r></a:p>')
    p, x, a, b = e._body_span(7); open(p, 'w', encoding='utf8').write(x[:a] + body + x[b:])
    before = open(p, encoding='utf8').read()
    assert e.insert_paragraph_like(7, 'L2 † 2017 전평, 번호 미상') == '* footnote'        # 끝에, 같은 수준(L2) 문단 서식
    assert e.insert_paragraph_like(7, 'L1 **Item B** 23-02', after_key='Item A') == 'Item A 22-01'
    after = open(p, encoding='utf8').read()
    ps = re.findall(r'<a:p>.*?</a:p>', after[e._body_span(7)[2]:e._body_span(7)[3]], re.S)
    assert [T.html.unescape(''.join(T._AT.findall(q))) for q in ps] == ['Item A 22-01', 'Item B 23-02', '* footnote', '† 2017 전평, 번호 미상']
    assert 'spcBef' in ps[1] and 'b="1"' in ps[1] and 'i="1"' in ps[3] and ps[0] in before and ps[2] in before
    # T2: 문단 간격 — pPr 자식 순서(lnSpc → spcBef → spcAft → 글머리표)
    assert e.set_paragraph_spacing(7, before=3, line=90) == 4
    y = open(p, encoding='utf8').read()
    ps = re.findall(r'<a:p>.*?</a:p>', y[e._body_span(7)[2]:e._body_span(7)[3]], re.S)
    for q in ps:
        ppr = re.search(r'<a:pPr\b.*?</a:pPr>', q, re.S).group(0)
        assert '<a:spcPts val="300"/>' in ppr and '<a:spcPct val="90000"/>' in ppr and ppr.count('spcBef>') == 2, ppr
        assert ppr.index('lnSpc') < ppr.index('spcBef') and (('buChar' not in ppr) or ppr.index('spcBef') < ppr.index('buChar'))
    e.save('/tmp/t2.pptx'); assert T.validate('/tmp/t2.pptx', SRC)
    # 굵은 run 은 굵은 글꼴로 — DejaVu Bold 는 Regular 보다 넓다
    reg = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    assert T._bold_sibling(reg) and 'Bold' in T._bold_sibling(reg)
    txt = 'Glomus jugulare paraganglioma with bone erosion'
    w = int(3.62 * 914400)
    assert T._est_lines_font_runs([(txt, True)], 18, w, reg, T._bold_sibling(reg)) >= T._est_lines_font_runs([(txt, False)], 18, w, reg, T._bold_sibling(reg))

def t_v1615_lead_space_whole_strict_sldnum_eomi():
    # R1: 수준 표시 뒤 첫 공백만 구분자 — 들여쓰기 대신 쓴 앞 공백을 지키고, body_notation → replace_paragraph_like 왕복에서 글이 같다
    assert T.parse_body_notation(['L1     Extramural tumor depth 22-01'])[0]['runs'] == [('    Extramural tumor depth 22-01', False, False)]
    d = T.Deck.open(SRC, wd('r1'))
    p, x, a, b = d._body_span(7)
    body = ('<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1600"/><a:t>    Extramural tumor depth 22-01</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1600" b="1"/><a:t>Denonvillier</a:t></a:r><a:r><a:rPr lang="en-US" sz="1600"/><a:t> fascia; Denonvilliers</a:t></a:r></a:p>')
    open(p, 'w', encoding='utf8').write(x[:a] + body + x[b:])
    line = d.body_notation(7)[0]; assert line == 'L1     Extramural tumor depth 22-01', line
    d.replace_paragraph_like(7, 'Extramural', line)
    assert d.para_texts(7)[0] == '    Extramural tumor depth 22-01' or '    Extramural tumor depth 22-01' in d.para_texts(7)
    # R2: 조각 전체 / R5: strict
    assert not d.settext(7, 'Denonvillier', 'Denonvilliers')['ok']
    try:
        d.settext(7, 'Denonvillier', 'Denonvilliers', strict=True); assert False
    except ValueError as e:
        assert '조각 전체' in str(e)
    assert d.settext(7, 'Denonvillier', 'Denonvilliers', whole=True)['ok']
    assert 'Denonvilliers fascia; Denonvilliers' in ''.join(d.para_texts(7))
    # R3: 새 노트에 노트 마스터의 sldNum 자리
    base, _ = _title_decks()
    B = T.Deck.open(base, wd('r3'))
    B.set_notes([s for s, _, _ in B.order()][0], ['x'])                      # 노트 마스터가 없던 덱 — 만든 최소 마스터엔 sldNum 없음
    e = T.Deck.open(SRC, wd('r3b'))
    nm = os.path.join(e.dir, 'ppt/notesMasters/notesMaster1.xml')
    if os.path.exists(nm) and 'type="sldNum"' in open(nm, encoding='utf8').read():
        n = e.import_slide(T.Deck.open(SRC, wd('r3s')), 2, after=1)
        assert 'type="sldNum"' in open(e._notes_path(n), encoding='utf8').read()
        e.save('/tmp/r3.pptx'); assert T.validate('/tmp/r3.pptx', SRC)
    # R4
    e.set_notes(7, ['tumor뿐 아니라 MRF째로 침범'])
    assert not [q for q in T.check_note_substitution(e) if 'slide7' in q]

def t_v1616_memo_sep_variant_rid_target_chem_prefix():
    # P1: '────── 기존 노트 (작성자) ──────' 도 메모 표지 — restore-memo 뒤 메모 표지 하나, 사본 없음
    A = T.Deck.open(SRC, wd('p1a')); sn = [s for s, _, _ in A.order()][2]; memo = _memo_notes(A, sn); A.save('/tmp/p1_orig.pptx')
    E = T.Deck.open('/tmp/p1_orig.pptx', wd('p1e'))
    E.set_notes(sn, ['대본 한 줄', '', '────── 기존 노트 (김작성) ──────', 'CT -> 조영 & 비교link', '두번째 메모'])
    assert T._is_memo_sep('────── 기존 노트 (김작성) ──────') and T._spoken_notes(E.notes(sn)) == ['대본 한 줄']
    E.save('/tmp/p1_ed.pptx')
    r = cli('restore-memo', '/tmp/p1_ed.pptx', '--original', '/tmp/p1_orig.pptx', '-o', '/tmp/p1_fixed.pptx')
    assert r.returncode == 0, r.stdout
    F = T.Deck.open('/tmp/p1_fixed.pptx', wd('p1f')); t = F.notes(sn)
    assert t.count('두번째 메모') == 1 and sum(1 for l in t if T._is_memo_sep(l)) == 1 and F.notes_sections(sn)[0] == ['대본 한 줄'], t
    # P2: 링크 rId 번호만 다른 문단은 '보존' — 복사로 새 rId 가 붙어도
    rr = T.diff_decks('/tmp/p1_orig.pptx', '/tmp/p1_fixed.pptx', io.StringIO())
    assert rr['notes'][F.screen_no(sn)] == (2, 0, 0), rr['notes'][F.screen_no(sn)]
    # normalize-notes 도 이 표지를 참고로 바꾸지 않는다
    G = T.Deck.open('/tmp/p1_ed.pptx', wd('p1g')); T.normalize_notes(G, sn)
    assert G.notes_sections(sn)[1] == [] and G.notes_sections(sn)[2] == ['CT -> 조영 & 비교link', '두번째 메모'], G.notes_sections(sn)
    # P3: 화학 결합형 앞부분은 두 단어 붙음이 아니다
    d = T.Deck.open(SRC, wd('p3'), theme='cud')
    d.set_body(8, T.Body().line('methyl groups and cellulose fibers'))
    d.set_notes(8, ['장 정결제로 PEG, methylcellulose 를 씁니다'])
    assert not [q for q in T.check_note_substitution(d) if 'slide8' in q], T.check_note_substitution(d)

def t_v1617_hidden_zone_eomi_keep_format():
    # H1
    d = T.Deck.open(SRC, wd('h1')); o = [s for s, _, _ in d.order()]
    assert not d.is_hidden(o[3]) and d.set_hidden(o[3]) is False and d.is_hidden(o[3]) and d.hidden_slides() == [4]
    assert d.set_hidden(o[3], False) is True and not d.is_hidden(o[3]) and d.hidden_slides() == []
    d.set_hidden(o[3]); d.save('/tmp/h1.pptx'); assert T.validate('/tmp/h1.pptx', SRC)
    # H2: 노트의 메모 구역만
    sn = o[2]; _memo_notes(d, sn); d.protect_memo(sn)
    d.set_notes(sn, ['대본'], tips=['참고에도 두번째 메모 라는 말'])
    assert not d.settext(sn, '두번째 메모', '두번째 메모(고침)', notes=True)['ok']            # 2회 매치
    assert d.settext(sn, '두번째 메모', '두번째 메모(고침)', notes=True, zone='memo', strict=True)['ok']
    sec = d.notes_sections(sn); assert sec[2][-1] == '두번째 메모(고침)' and sec[1] == ['참고에도 두번째 메모 라는 말'], sec
    try:
        d.settext(sn, 'x', 'y', notes=True, zone='bogus'); assert False
    except ValueError:
        pass
    # H3
    d.set_notes(7, ['APHE여야 합니다. cholestasis 소견.'])
    assert not [q for q in T.check_note_substitution(d) if 'slide7' in q]
    # H4: 새 줄에 ⇥·** 가 빠져도 원래 앞 탭·공백·굵은 조각 유지
    e = T.Deck.open(SRC, wd('h4'))
    p, x, a, b = e._body_span(7)
    body = ('<a:p><a:pPr lvl="2"/><a:r><a:rPr lang="en-US" sz="1400"/><a:t>\t  </a:t></a:r><a:r><a:rPr lang="en-US" sz="1400" b="1"/><a:t>Hepatic adenoma</a:t></a:r>'
            '<a:r><a:rPr lang="en-US" sz="1400"/><a:t> 22-03</a:t></a:r></a:p>')
    open(p, 'w', encoding='utf8').write(x[:a] + body + x[b:])
    cur = e.body_notation(7)[0]; assert cur.startswith('L2 ⇥') and '**Hepatic adenoma**' in cur, cur
    e.replace_paragraph_like(7, 'Hepatic adenoma', 'L2 Hepatic adenoma 22-03, 25-05', keep_format=True)
    now = e.body_notation(7)[0]
    assert now.startswith('L2 ⇥') and '**Hepatic adenoma**' in now and '25-05' in now, now
    assert T._merge_format('L1 **A** x', 'L1 A y') == 'L1 **A** y'

def t_v1618_delete_shape():
    d = T.Deck.open(SRC, wd('c1'))
    x = open(d._slide(7), encoding='utf8').read()
    box = ('<p:sp><p:nvSpPr><p:cNvPr id="901" name="직사각형 3"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></a:xfrm>'
           '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>Less common features</a:t></a:r></a:p></p:txBody></p:sp>')
    open(d._slide(7), 'w', encoding='utf8').write(x.replace('</p:spTree>', box + box.replace('901', '902').replace('직사각형 3', '직사각형 4').replace('Less common', 'Smoking related') + '</p:spTree>', 1))
    try:
        d.delete_shape(7, '직사각형 3', must_contain='Smoking'); assert False
    except ValueError as e:
        assert '없음' in str(e)
    assert d.delete_shape(7, '직사각형 3', must_contain='Less common') == 'Less common features'
    y = open(d._slide(7), encoding='utf8').read()
    assert 'Less common' not in y and 'Smoking related' in y
    try:
        d.delete_shape(7, '직사각형 3'); assert False
    except ValueError as e:
        assert '0개' in str(e)
    # 그림 도형을 지우면 그 그림 관계도 빠진다
    pic_slide = next(s for s in d.slide_numbers() if '<p:pic>' in open(d._slide(s), encoding='utf8').read())
    px = open(d._slide(pic_slide), encoding='utf8').read()
    pm = re.search(r'<p:pic>.*?</p:pic>', px, re.S).group(0)
    nm = re.search(r'name="([^"]*)"', pm).group(1); rid = re.search(r'r:embed="(rId\d+)"', pm).group(1)
    if px.count('name="%s"' % nm) == 1 and px.count('r:embed="%s"' % rid) == 1:
        d.delete_shape(pic_slide, T.html.unescape(nm))
        assert 'Id="%s"' % rid not in open(d._slide_rels(pic_slide), encoding='utf8').read()
    d.save('/tmp/c1.pptx'); assert T.validate('/tmp/c1.pptx', SRC)

def t_v1626_widen_label():
    d = T.Deck.open(SRC, wd('k6'))
    W, _ = d.slide_size(); E = T.EMU_IN
    def box(i, nm, txt, x, w, algn='l', fill=False):
        f = '<a:solidFill><a:srgbClr val="FFFF00"/></a:solidFill>' if fill else '<a:noFill/>'
        return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="100"/>'
                '<a:ext cx="%d" cy="300000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s</p:spPr><p:txBody><a:bodyPr wrap="none">'
                '<a:spAutoFit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="%s"/><a:r><a:rPr lang="ko-KR" sz="1800"/><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>'
                % (i, nm, x, w, f, algn, txt))
    x = open(d._slide(7), encoding='utf8').read()
    boxes = (box(951, '이름표 1', '(R3 가나다)', W - int(1.5 * E), int(1.2 * E), 'r') + box(952, '이름표 2', '(R1 라마바)', int(0.5 * E), int(1.0 * E)) +
             box(953, '이름표 3', '(R2 사아자)', int(3 * E), int(1.0 * E), fill=True) + box(954, '이름표 4', '(R4 차카타)', 0, int(1.0 * E), 'r') +
             box(955, '본문 글', 'Not a label', int(5 * E), int(1.0 * E)))
    open(d._slide(7), 'w', encoding='utf8').write(x.replace('</p:spTree>', boxes + '</p:spTree>'))
    r = d.widen_label(7, pattern=r'\(R\d [^)]*\)', min_width_in=2.0)
    got = {nm: what for nm, _, _, _, what in r}
    assert got == {'이름표 1': '바꿈', '이름표 2': '바꿈', '이름표 3': '건너뜀: 채우기 있음', '이름표 4': '건너뜀: 슬라이드 밖으로 나감'}, got
    y = open(d._slide(7), encoding='utf8').read()
    s1 = re.search(r'name="이름표 1".*?<a:off x="(\d+)" y="100"/><a:ext cx="(\d+)"', y, re.S)
    assert int(s1.group(1)) + int(s1.group(2)) == W - int(1.5 * E) + int(1.2 * E) and int(s1.group(2)) == int(round(2.0 * E))   # 오른쪽 끝 고정
    s2 = re.search(r'name="이름표 2".*?<a:off x="(\d+)" y="100"/>', y, re.S)
    assert int(s2.group(1)) == int(0.5 * E)                                                                                          # 왼쪽 끝 고정
    assert '(R3 가나다)' in y and 'sz="1800"' in y
    d.save('/tmp/k6.pptx'); assert T.validate('/tmp/k6.pptx', SRC)


def t_v1627_fit_corner_boxes_and_protect_memo():
    d = T.Deck.open(SRC, wd('k7'))
    W, H = d.slide_size(); E = T.EMU_IN
    def box(i, nm, txt, x, y, w, h, fill=False, sz=1000):
        f = '<a:solidFill><a:srgbClr val="FFFF00"/></a:solidFill>' if fill else '<a:noFill/>'
        return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="%d"/>'
                '<a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s</p:spPr><p:txBody><a:bodyPr wrap="none">'
                '<a:spAutoFit/></a:bodyPr><a:lstStyle/><a:p><a:r><a:rPr lang="ko-KR" sz="%d"/><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>'
                % (i, nm, x, y, w, h, f, sz, txt))
    long_txt = 'Textbook of Radiology, Ch.25 Trauma, X. Soft tissue injury (cited)'
    w0, h0 = int(3.11 * E), int(0.27 * E)
    boxes = (box(961, '인용 구석', long_txt, W - w0, H - h0, w0, h0) +                          # 오른쪽·아래에 붙음
             box(962, '가운데 글', long_txt, int(2 * E), int(2 * E), int(1 * E), int(0.3 * E)) +  # 가장자리 아님
             box(963, '채운 구석', long_txt, 0, 0, int(1 * E), int(0.3 * E), fill=True) +
             box(964, '이웃 상자', 'x', W - int(4.3 * E), H - int(0.25 * E), int(0.5 * E), int(0.2 * E)))
    x = open(d._slide(7), encoding='utf8').read()
    open(d._slide(7), 'w', encoding='utf8').write(x.replace('</p:spTree>', boxes + '</p:spTree>'))
    r = {q['name']: q for q in d.fit_corner_boxes(7)}
    assert set(r) == {'인용 구석'}, r                                              # 가운데·채운 상자는 대상 아님
    q = r['인용 구석']; ox, oy, ow, oh = q['old']; nx, ny, nw, nh = q['new']
    assert q['what'] == '바꿈' and nw > ow and abs((nx + nw) - (ox + ow)) < 0.02 and abs((ny + nh) - (oy + oh)) < 0.02, q   # 오른쪽·아래 끝 고정
    assert '이웃 상자' in q['overlap'], q
    y = open(d._slide(7), encoding='utf8').read()
    assert long_txt in y and 'sz="1000"' in y and 'wrap="none"' in y
    d.save('/tmp/k7.pptx'); assert T.validate('/tmp/k7.pptx', SRC)
    # protect-memo CLI
    d2 = T.Deck.open(SRC, wd('p2')); order = [sn for sn, _, _ in d2.order() if sn]
    d2.set_notes(order[1], ['원작자 노트 한 줄']); d2.set_notes(order[2], ['대본'], ['· 참고']); d2.save('/tmp/p2_in.pptx')
    r = subprocess.run([sys.executable, os.path.join(HERE, 'deck_toolkit.py'), 'protect-memo', '/tmp/p2_in.pptx', '-o', '/tmp/p2_out.pptx', '--screens', '2-3'],
                       capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    assert r.returncode == 0 and '감싼 화면 1' in r.stdout, (r.stdout, r.stderr[-300:])
    d3 = T.Deck.open('/tmp/p2_out.pptx', wd('p3'))
    assert d3.notes_sections(order[1])[2] == ['원작자 노트 한 줄'] and d3.notes_sections(order[2])[0] == ['대본']
    assert T.MIN_FIT_TITLE_PT == T.MIN_TITLE_PT == 24


def t_v1628_title_need_height():
    # 발표 K8(근골격): 가져온 제목의 안쪽 여백 위·아래 0.39"(360000) 인데 띠를 0.66" 로 맞췄다 — Google Slides 에서 글이 띠 밖으로
    from pptx import Presentation
    from pptx.util import Pt, Emu, Inches
    from pptx.dml.color import RGBColor
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    for k, h in enumerate([1117331] * 6 + [605908] * 3):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        t = sl.shapes.title; t.left, t.top, t.width, t.height = Emu(0), Emu(0), Inches(10), Emu(h)
        t.fill.solid(); t.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        t.text_frame.text = 'Q. 25-%02d Short title' % k
        t.text_frame.margin_top = t.text_frame.margin_bottom = Emu(360000)
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(28)
    path = os.path.join(TMP, 'k8.pptx'); prs.save(path)
    B = T.Deck.open(path, wd('k8'))
    prof = T.title_profile(B, like=1)
    assert prof['h_by_lines'].get(1) == 1117331 and prof['excluded'] == 3, prof          # 0.66" 띠 셋은 배우지 않는다
    order = [x for x, _, _ in B.order() if x]
    bad = dict(prof, h_by_lines={1: 605908})                                               # 오염된 규격을 받았어도
    ch = T.conform_title(B, order[7], bad)
    assert ch and '필요 높이로' in ch[0], ch
    cy = int(re.search(r'<a:ext cx="\d+" cy="(\d+)"', T._title_info(B, order[7])['seg']).group(1))
    assert cy >= T._title_need_h(360000, 360000, 1, 2800) - 2, cy
    r = T.raise_title_band(B, order[8])                                                    # Google 안전: 0.66" → 필요 높이
    assert r and r[0].startswith('띠 0.66"'), r
    assert T.raise_title_band(B, order[0]) == []                                          # 1.22" 는 그대로
    B.save('/tmp/k8o.pptx'); assert T.validate('/tmp/k8o.pptx', path)


def t_v1629_shrink_bottom_inset_and_joint_profile():
    from pptx import Presentation
    from pptx.util import Pt, Emu, Inches
    from pptx.dml.color import RGBColor
    def mk(path, specs, body_top=None):
        prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        for h, ins in specs:
            sl = prs.slides.add_slide(prs.slide_layouts[5])
            t = sl.shapes.title; t.left, t.top, t.width, t.height = Emu(0), Emu(0), Inches(10), Emu(h)
            t.fill.solid(); t.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
            t.text_frame.text = 'Q. short title'; t.text_frame.margin_top = t.text_frame.margin_bottom = Emu(ins)
            for r in t.text_frame.paragraphs[0].runs:
                r.font.size = Pt(28)
            if body_top is not None:
                b = sl.shapes.add_textbox(Inches(0.5), Inches(body_top), Inches(6), Inches(1)); b.name = 'Body'; b.text_frame.text = 'content'
        prs.save(path); return path
    # K9: 0.66" 띠, 여백 0.39"×2, 본문이 1.08" 에서 시작 — 그냥은 못 키우고, 아래 여백 0.1" 로 줄이면 들어간다
    p = mk(os.path.join(TMP, 'k9.pptx'), [(605908, 360000)], body_top=1.08)
    B = T.Deck.open(p, wd('k9')); sn = [x for x, _, _ in B.order() if x][0]
    assert T.raise_title_band(B, sn)[0].startswith('[!]')
    r = T.raise_title_band(B, sn, shrink_bottom=0.1)
    assert r and '아래 여백 0.39" → 0.10"' in r[0], r
    i = T._title_info(B, sn)
    assert i['bIns'] == 91440 and i['tIns'] == 360000 and i['y'] + i['h'] <= int(1.08 * T.EMU_IN), i
    p2 = mk(os.path.join(TMP, 'k9b.pptx'), [(605908, 360000)], body_top=0.8)
    B2 = T.Deck.open(p2, wd('k9b')); sn2 = [x for x, _, _ in B2.order() if x][0]
    assert T.raise_title_band(B2, sn2, shrink_bottom=0.1)[0].startswith('[!] 아래 여백을')      # 0.1" 로도 겹침 — 그대로
    # K8-b: 높이와 여백을 같은 제목에서 — 여백 큰 제목(1.22") 셋 + 여백 작은 제목(0.66") 여섯이면 여백 규격은 작은 것, 높이도 그 제목들의 0.66"
    p3 = mk(os.path.join(TMP, 'k8b.pptx'), [(1117331, 360000)] * 3 + [(605908, 45720)] * 6)
    B3 = T.Deck.open(p3, wd('k8b')); prof = T.title_profile(B3, like=1)
    # v16.30: 기준 제목(화면 1, 여백 0.39")의 무리에서만 배운다 — 수가 적어도 뒤집히지 않는다
    assert prof['ins'][2] == 360000 and prof['h_by_lines'][1] == 1117331 and prof['groups'] == {0.39: 3, 0.05: 6}, prof
    prof_b = T.title_profile(B3, like=4)
    assert prof_b['ins'][2] == 45720 and prof_b['h_by_lines'][1] == 605908, prof_b
    p4 = mk(os.path.join(TMP, 'k8c.pptx'), [(1117331, 360000)] * 6 + [(605908, 45720)] * 3)
    prof4 = T.title_profile(T.Deck.open(p4, wd('k8c')), like=1)
    assert prof4['ins'][2] == 360000 and prof4['h_by_lines'][1] == 1117331, prof4              # 섞지 않는다
    r = subprocess.run([sys.executable, os.path.join(HERE, 'deck_toolkit.py'), 'titles', p4, '--like', '1'], capture_output=True, text=True,
                       env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    assert '1줄 제목 6개' in r.stdout and '안쪽 여백 규격' in r.stdout and '제목 무리' in r.stdout, r.stdout[-600:]


def t_v1630_balance_and_own_size():
    from pptx import Presentation
    from pptx.util import Pt, Emu, Inches
    from pptx.dml.color import RGBColor
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    def add(h, ins, sz, text, body_top=None):
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        t = sl.shapes.title; t.left, t.top, t.width, t.height = Emu(0), Emu(0), Inches(10), Emu(h)
        t.fill.solid(); t.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        t.text_frame.text = text; t.text_frame.margin_top = t.text_frame.margin_bottom = Emu(ins)
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(sz)
        if body_top is not None:
            b = sl.shapes.add_textbox(Inches(0.5), Inches(body_top), Inches(6), Inches(1)); b.name = 'Body'; b.text_frame.text = 'content'
    for _ in range(3):
        add(1117331, 360000, 28, 'Q. reference title')                         # 기준 무리(여백 0.39", 1줄 1.22")
    long_t = 'Q. a title long enough to wrap only at the larger size xx'        # 28pt 로 세면 2줄, 23pt 면 1줄
    add(605908, 45720, 23, long_t, body_top=1.08)                                # 4: 가져온 무리, 본문이 1.08" 에서
    add(605908, 45720, 23, 'Q. no body below')                                   # 5: 아래 내용 없음
    add(605908, 45720, 23, 'Q. body too close', body_top=0.45)                   # 6: 본문이 너무 가까움
    path = os.path.join(TMP, 'k10.pptx'); prs.save(path)
    B = T.Deck.open(path, wd('k10')); order = [x for x, _, _ in B.order() if x]
    i4 = T._title_info(B, order[3])
    assert T._title_lines(i4, 2800) == 2 and T._title_lines(i4, 2300) == 1, (T._title_lines(i4, 2800), T._title_lines(i4, 2300))
    prof = T.title_profile(B, like=4)
    assert prof['excluded'] == 0, prof['excluded_list']                          # 화면 6·7: 자기 크기(23pt)로 1줄 — 빼지 않는다
    ref = T.title_profile(B, like=1)
    r4 = T.balance_title_band(B, order[3], ref)
    i4 = T._title_info(B, order[3])
    assert r4 and '아래 "Body" 까지' in r4[0] and i4['tIns'] == i4['bIns'] and i4['y'] + i4['h'] == int(1.08 * T.EMU_IN) - int(0.1 * T.EMU_IN), (r4, i4['h'], i4['tIns'])
    r5 = T.balance_title_band(B, order[4], ref)
    i5 = T._title_info(B, order[4])
    assert r5 and '(규격)' in r5[0] and i5['h'] == 1117331 and i5['tIns'] == i5['bIns'], (r5, i5['h'])
    assert T.balance_title_band(B, order[5], ref)[0].startswith('[!]')                  # 0.45" − 0.1" 에 23pt 글이 안 들어감
    assert T.balance_title_band(B, order[0], ref) == []                                # 기준 제목은 이미 규격·균형
    assert '23' in r4[0] and 'content' in open(B._slide(order[3]), encoding='utf8').read()
    B.save('/tmp/k10o.pptx'); assert T.validate('/tmp/k10o.pptx', path)


def t_v1631_adopt_house_look():
    d = T.Deck.open(SRC, wd('k12')); order = [x for x, _, _ in d.order() if x]
    sn = order[6]; E = T.EMU_IN
    x = open(d._slide(sn), encoding='utf8').read()
    x = re.sub(r'<p:sp>(?:(?!<p:sp>).)*?type="(?:title|ctrTitle)"(?:(?!<p:sp>).)*?</p:sp>', '', x, flags=re.S)   # 제목 자리 표시자 없음(가져온 꼴)
    def tb(i, nm, y, runs, fill=None):
        f = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill if fill else '<a:noFill/>'
        rr = ''.join('<a:r><a:rPr lang="en-US" sz="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:rPr><a:t>%s</a:t></a:r>' % r for r in runs)
        return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="%d"/>'
                '<a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s</p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>'
                '<a:p>%s</a:p></p:txBody></p:sp>' % (i, nm, int(0.5 * E), int(y * E), int(8 * E), int(0.8 * E), f, rr))
    x = x.replace('</p:spTree>', tb(971, 'TextBox 1', 0.2, [(4000, 'FFFFFF', 'Endoleak')]) +
                  tb(972, 'TextBox 2', 2.0, [(1800, 'FFFFFF', 'Type I '), (1800, 'FF0000', 'red emphasis')]) +
                  tb(973, 'Dark box', 4.0, [(1800, 'FFFFFF', 'white on dark')], fill='1F1F1F') + '</p:spTree>')
    open(d._slide(sn), 'w', encoding='utf8').write(x)
    r = T.adopt_house_look(d, sn)
    assert any('제목 글상자 "Endoleak"' in c for c in r) and any('밝은 글자색 1곳' in c for c in r), r
    y = open(d._slide(sn), encoding='utf8').read()
    ti = T._title_info(d, sn)
    assert ti and ti['kind'] == 'ph' and ''.join(ti['paras']) == 'Endoleak' and 'name="TextBox 1"' not in y
    b2 = re.search(r'name="TextBox 2".*?</p:sp>', y, re.S).group(0)
    assert 'FFFFFF' not in b2 and 'FF0000' in b2                              # 흰 글자는 테마색으로, 빨강은 그대로
    assert 'FFFFFF' in re.search(r'name="Dark box".*?</p:sp>', y, re.S).group(0)   # 어두운 상자 안 흰 글자는 그대로
    assert T.adopt_house_look(d, order[1]) == [] or all(c.startswith('[참고]') or '밝은' in c for c in T.adopt_house_look(d, order[1], dry_run=True))
    d.save('/tmp/k12.pptx'); assert T.validate('/tmp/k12.pptx', SRC)


def t_v1632_inherited_title_diff_boxes_bake_pts():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    # K13: 위치·크기·채움을 물려받는 제목(슬라이드에 xfrm 없음) — 적어 넣고 보정, 레이아웃은 그대로
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    for txt in ('Prostate artery embolization outcomes and complications in a long title form here', 'Short'):
        sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text_frame.text = txt
    path = os.path.join(TMP, 'k13.pptx'); prs.save(path)
    d = T.Deck.open(path, wd('k13')); order = [x for x, _, _ in d.order() if x]
    lay = os.path.join(d.dir, 'ppt/slideLayouts', d.layout_of(order[0])); lx = open(lay, encoding='utf8').read()
    lx = re.sub(r'(<p:ph type="title"/></p:nvPr></p:nvSpPr>)<p:spPr/>', r'\1<p:spPr><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill></p:spPr>', lx, 1)
    open(lay, 'w', encoding='utf8').write(lx)
    i0 = T._title_info(d, order[0])
    assert i0['inherit'] and i0['fill'] and 'D9D9D9' in i0['fill'], i0['fill']            # 채움도 레이아웃에서
    r = T.raise_title_band(d, order[0])
    assert r and r[0].startswith('물려받던 제목 — 적어 넣음 · 띠 1.25"'), r
    x0 = open(d._slide(order[0]), encoding='utf8').read()
    assert '<a:xfrm>' in re.search(r'<p:sp>(?:(?!<p:sp>).)*?type="title".*?</p:sp>', x0, re.S).group(0)
    assert open(lay, encoding='utf8').read() == lx                                        # 레이아웃은 고치지 않는다
    rb = T.balance_title_band(d, order[1], {'h_by_lines': {1: int(1.22 * T.EMU_IN)}, 'ins': (91440, 91440, 45720, 45720), 'est_sz': 4400})
    assert rb and rb[0].startswith('물려받던 제목 — 적어 넣음 · ') and T._title_info(d, order[1])['tIns'] == T._title_info(d, order[1])['bIns'], rb
    d.save('/tmp/k13o.pptx'); assert T.validate('/tmp/k13o.pptx', path)
    # D1: 상자 순서만 바뀌고 한 상자 글 조각이 합쳐진 화면은 '의도하지 않은 글 변경' 이 아니다
    A = T.Deck.open(SRC, wd('d1a')); sa = [x for x, _, _ in A.order() if x][6]
    def tbx(i, nm, runs):
        return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/>'
                '<a:ext cx="100" cy="100"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p>%s</a:p></p:txBody></p:sp>'
                % (i, nm, ''.join('<a:r><a:rPr lang="en-US"/><a:t>%s</a:t></a:r>' % t for t in runs)))
    xa = open(A._slide(sa), encoding='utf8').read()
    open(A._slide(sa), 'w', encoding='utf8').write(xa.replace('</p:spTree>', tbx(981, 'L', ['(R4 name)']) + tbx(982, 'T', ['complications - ', 'Endoleak']) + '</p:spTree>'))
    A.save('/tmp/d1a.pptx')
    B = T.Deck.open('/tmp/d1a.pptx', wd('d1b')); xb = open(B._slide(sa), encoding='utf8').read()
    open(B._slide(sa), 'w', encoding='utf8').write(xb.replace(tbx(981, 'L', ['(R4 name)']) + tbx(982, 'T', ['complications - ', 'Endoleak']),
                                                              tbx(982, 'T', ['complications - Endoleak']) + tbx(981, 'L', ['(R4 name)'])))
    B.save('/tmp/d1b.pptx')
    res = T.diff_decks('/tmp/d1a.pptx', '/tmp/d1b.pptx', stream=io.StringIO())
    assert res['unintended'] == [] and 7 in res['slides_changed'], res
    C = T.Deck.open('/tmp/d1a.pptx', wd('d1c')); xc = open(C._slide(sa), encoding='utf8').read()
    open(C._slide(sa), 'w', encoding='utf8').write(xc.replace('Endoleak', 'Endoleak type II')); C.save('/tmp/d1c.pptx')
    assert T.diff_decks('/tmp/d1a.pptx', '/tmp/d1c.pptx', stream=io.StringIO())['unintended'] == [7]
    # D2: pt 로 고정된 줄 간격 문단 — bake 가 lnSpc 를 하나 더 넣지 않는다
    E = T.Deck.open(SRC, wd('d2')); se = [x for x, _, _ in E.order() if x][6]
    para = ('<a:p><a:pPr marL="0" lvl="0" indent="0"><a:lnSpc><a:spcPts val="2700"/></a:lnSpc><a:spcBef><a:spcPts val="1200"/></a:spcBef><a:buNone/></a:pPr>'
            '<a:r><a:rPr lang="en-US" sz="1800"/><a:t>google export line</a:t></a:r></a:p>')
    box = ('<p:sp><p:nvSpPr><p:cNvPr id="991" name="Google Shape;439;p42"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="100" y="100"/>'
           '<a:ext cx="2560320" cy="1371600"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr anchor="b">'
           '<a:normAutofit fontScale="55000" lnSpcReduction="20000"/></a:bodyPr><a:lstStyle/>%s</p:txBody></p:sp>' % (para * 6))
    xe = open(E._slide(se), encoding='utf8').read(); open(E._slide(se), 'w', encoding='utf8').write(xe.replace('</p:spTree>', box + '</p:spTree>'))
    assert E.bake_autofit(se, shape='Google Shape;439;p42') == 1
    q = re.search(r'name="Google Shape;439;p42".*?</p:sp>', open(E._slide(se), encoding='utf8').read(), re.S).group(0)
    assert q.count('<a:lnSpc>') == 6 and q.count('spcPts val="2160"') == 6 and 'sz="990"' in q, q[:600]
    E.save('/tmp/d2o.pptx'); assert T.validate('/tmp/d2o.pptx', SRC)


def t_v1632_title_block_k14():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text_frame.text = 'Endoleak type II'   # 제목 위치는 마스터에서 물려받음
    b = sl.shapes.add_textbox(Inches(0.5), Inches(0.43), Inches(9), Inches(6.9)); b.name = 'TextBox 4'
    tf = b.text_frame; tf.text = 'type II green'; tf.paragraphs[0].runs[0].font.size = Pt(20)
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x92, 0xD0, 0x50)
    p2 = tf.add_paragraph(); p2.text = 'red keep'; p2.runs[0].font.size = Pt(20); p2.runs[0].font.color.rgb = RGBColor(0xFF, 0, 0)
    n = sl.shapes.add_textbox(Inches(8.5), Inches(7.1), Inches(1.4), Inches(0.3)); n.name = 'Label'; n.text_frame.text = '(R4 name)'
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.0), Inches(9.5), Inches(1.0)); ln.name = 'Rule 18'
    path = os.path.join(TMP, 'k14.pptx'); prs.save(path)
    d = T.Deck.open(path, wd('k14')); sn = [x for x, _, _ in d.order() if x][0]
    assert T.adopt_house_look(d, sn, recolor=['92d050']) == ['지정한 색 92D050 1곳을 지워 테마 글자색으로']
    dry = T.title_block(d, sn, band_h=1.22, dry_run=True)                                     # 기본: 내리지도 지우지도 않고 알림
    assert any('--drop-title-rule' in c for c in dry) and any('--push-content' in c for c in dry), dry
    r = T.title_block(d, sn, band_h=1.22, drop_rule=True, push=True)
    assert r[0].startswith('물려받던 제목 — 적어 넣음 · 띠 1.25" → 1.22"') and any('옛 제목 밑줄 "Rule 18"' in c for c in r), r
    x = open(d._slide(sn), encoding='utf8').read()
    i = T._title_info(d, sn)
    ty, tcy = (int(v) for v in re.search(r'name="TextBox 4".*?<a:off x="\d+" y="(\d+)"/><a:ext cx="\d+" cy="(\d+)"', x, re.S).groups())
    assert ty == i['y'] + int(1.22 * T.EMU_IN) + int(0.1 * T.EMU_IN) and ty + tcy == int(7.5 * T.EMU_IN), (ty, tcy)   # 띠 아래 + 0.1", 슬라이드 안
    assert 'sz="2000"' not in re.search(r'name="TextBox 4".*?</p:sp>', x, re.S).group(0)            # 넘침만큼 글자 비율로
    assert 'Rule 18' not in x and '92D050' not in x and 'FF0000' in x
    assert re.search(r'name="Label".*?<a:off x="\d+" y="(\d+)"', x, re.S).group(1) == str(int(7.1 * T.EMU_IN))   # 아래 이름표는 그대로
    assert i['tIns'] == i['bIns']
    d.save('/tmp/k14o.pptx'); assert T.validate('/tmp/k14o.pptx', path)


def t_v1633_colors_band_fit_layout():
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    E = T.EMU_IN
    # K15-1 개정: 이름 색(prstClr white)·검정 → 테마 글자색, 자동 지우기 끔(노랑은 남김), 낮은 대비 강조색만 같은 계열로 진하게
    d = T.Deck.open(SRC, wd('k151')); sn = [x for x, _, _ in d.order() if x][6]
    runs = ''.join('<a:r><a:rPr lang="en-US" sz="1800"><a:solidFill>%s</a:solidFill></a:rPr><a:t>%s</a:t></a:r>' % c for c in (
        ('<a:prstClr val="white"/>', 'W'), ('<a:srgbClr val="000000"/>', 'K'), ('<a:srgbClr val="FFFF00"/>', 'Y'), ('<a:srgbClr val="FF0000"/>', 'R')))
    box = ('<p:sp><p:nvSpPr><p:cNvPr id="999" name="Colors"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="100" y="3000000"/>'
           '<a:ext cx="3000000" cy="500000"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p>%s</a:p></p:txBody></p:sp>' % runs)
    x = open(d._slide(sn), encoding='utf8').read(); open(d._slide(sn), 'w', encoding='utf8').write(x.replace('</p:spTree>', box + '</p:spTree>'))
    r = T.adopt_house_look(d, sn, recolor=['white', 'black'], auto_light=False, darken=True)
    y = re.search(r'name="Colors".*?</p:sp>', open(d._slide(sn), encoding='utf8').read(), re.S).group(0)
    cols = re.findall(r'srgbClr val="(\w{6})"', y)
    assert 'prstClr' not in y and '000000' not in cols and 'FF0000' in cols and 'FFFF00' not in cols, (r, cols)
    new_y = re.findall(r'srgbClr val="(\w{6})"', y)
    assert len(new_y) == 2 and T._contrast([c for c in new_y if c != 'FF0000'][0], 'FFFFFF') >= 3.0, (new_y, r)
    assert any('배경 대비가 낮은 FFFF00 →' in c for c in r), r
    assert T._delta_e([c for c in new_y if c != 'FF0000'][0], 'FF0000') >= 20
    # K15-3: 따로 그린 띠 도형 + 위치를 물려받는 제목·본문 자리 표시자
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[1]); sl.shapes.title.text_frame.text = 'Prostate artery embolization'
    sl.placeholders[1].text_frame.text = 'body text'
    band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(-18288), Inches(10), Inches(0.62)); band.name = '제목 1'
    band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    path = os.path.join(TMP, 'k153.pptx'); prs.save(path)
    B = T.Deck.open(path, wd('k153')); sb = [x for x, _, _ in B.order() if x][0]
    r = T.title_block(B, sb, band_h=1.8, push=True)                 # 본문(마스터에서 1.6" 물려받음)이 띠 아래 + 0.1" 보다 위가 되게
    assert r[0].startswith('띠 도형 "제목 1" 을 띠로') and any('물려받던 위치 → 적어 넣음' in c for c in r), r
    xb = open(B._slide(sb), encoding='utf8').read()
    bnd = re.search(r'name="제목 1".*?<a:off x="(-?\d+)" y="(-?\d+)"/><a:ext cx="\d+" cy="(\d+)"', xb, re.S).groups()
    ti = T._title_info(B, sb)
    assert int(bnd[2]) == ti['h'] == int(1.8 * E) and int(bnd[1]) == ti['y'] == -18288, (bnd, ti['y'], ti['h'])
    body = re.search(r'<p:sp>(?:(?!<p:sp>).)*?idx="1"(?:(?!<p:sp>).)*?</p:sp>', xb, re.S).group(0)
    assert int(re.search(r'<a:off x="-?\d+" y="(-?\d+)"', body).group(1)) == -18288 + int(1.8 * E) + int(0.1 * E), body[:300]
    B.save('/tmp/k153o.pptx'); assert T.validate('/tmp/k153o.pptx', path)
    # K16 1단계: 밀집 화면 — 제목 띠 촘촘히, 그림+주석 비율, 본문 글과 안 겹침, 이름표 그대로
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text_frame.text = 'Endoleak type II'
    b = sl.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(5.0), Inches(5.5)); b.name = 'Body'
    tf = b.text_frame; tf.word_wrap = True; tf.text = 'Type II endoleak from lumbar or IMA branches'
    for k in range(5):
        pp = tf.add_paragraph(); pp.text = 'point %d with some explanatory words here' % k
    for pp in tf.paragraphs:
        for rr in pp.runs:
            rr.font.size = Pt(20)
    img = os.path.join(TMP, 'k16.png'); Image.new('RGB', (400, 300), 'gray').save(img)
    pic = sl.shapes.add_picture(img, Inches(4.5), Inches(1.6), Inches(5.0), Inches(3.75)); pic.name = 'Pic 1'
    arr = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.0), Inches(2.5), Inches(0.5), Inches(0.3)); arr.name = 'Arrow 1'
    lab = sl.shapes.add_textbox(Inches(0.2), Inches(7.1), Inches(1.4), Inches(0.3)); lab.name = 'Label'; lab.text_frame.text = '(R4 name)'
    path = os.path.join(TMP, 'k16.pptx'); prs.save(path)
    F = T.Deck.open(path, wd('k16')); sf = [x for x, _, _ in F.order() if x][0]
    r = T.fit_layout(F, sf)
    assert r and not r[0].startswith('[!]') and any(c.startswith('그림 "Pic 1"') and '주석·설명 1 같이' in c for c in r), r
    xf = open(F._slide(sf), encoding='utf8').read()
    g = lambda nm: tuple(int(v) for v in re.search(r'name="%s".*?<a:off x="(-?\d+)" y="(-?\d+)"/><a:ext cx="(\d+)" cy="(\d+)"' % re.escape(nm), xf, re.S).groups())
    ti = T._title_info(F, sf); pg, ag, lg, bg = g('Pic 1'), g('Arrow 1'), g('Label'), g('Body')
    assert ti["y"] == 0 and ti["h"] < int(1.0 * E), ti["h"]                                   # 촘촘한 띠(44pt 글 + 0.08"×2), 위에 붙음
    assert pg[1] >= ti['h'] and pg[0] + pg[2] <= int(10 * E) and pg[1] + pg[3] <= int(7.5 * E)
    assert pg[0] >= bg[0] + bg[2] or pg[1] >= bg[1] + int(0.5 * E), (pg, bg)                   # 본문과 옆으로 갈라섬(글과 안 겹침)
    assert pg[0] <= ag[0] and ag[0] + ag[2] <= pg[0] + pg[2] and pg[1] <= ag[1] <= pg[1] + pg[3], (pg, ag)   # 주석은 그림 안에 그대로
    assert lg == (int(0.2 * E), int(7.1 * E), int(1.4 * E), int(0.3 * E))                     # 이름표 그대로
    F.save('/tmp/k16o.pptx'); assert T.validate('/tmp/k16o.pptx', path)
    many = T.Deck.open(path, wd('k16b')); sm = [x for x, _, _ in many.order() if x][0]
    xm = open(many._slide(sm), encoding='utf8').read()
    xm = xm.replace('point 0 with some explanatory words here', ' '.join(['very long text'] * 400))
    open(many._slide(sm), 'w', encoding='utf8').write(xm)
    assert T.fit_layout(many, sm)[0].startswith('[!]') and open(many._slide(sm), encoding='utf8').read() == xm   # 안 되면 바꾸지 않는다


def t_v1634_fit_layout_stage2():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image
    E = T.EMU_IN
    img = os.path.join(TMP, 'k17.png'); Image.new('RGB', (400, 300), 'gray').save(img)
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    def slide(title, body_pt, lab_xy):
        sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text_frame.text = title
        b = sl.shapes.add_textbox(Inches(0.3), Inches(1.0), Inches(6.5), Inches(3.23)); b.name = 'TextBox 4'
        tf = b.text_frame; tf.word_wrap = True; tf.text = 'Type II endoleak — retrograde flow'
        for k in range(3):
            pp = tf.add_paragraph(); pp.text = 'short point %d' % k
        for pp in tf.paragraphs:
            for rr in pp.runs:
                rr.font.size = Pt(body_pt)
        b.text_frame.auto_size = True                                  # spAutoFit
        pic = sl.shapes.add_picture(img, Inches(7.06), Inches(1.2), Inches(2.8), Inches(2.1)); pic.name = 'Pic 1'
        lab = sl.shapes.add_textbox(Inches(lab_xy[0]), Inches(lab_xy[1]), Inches(1.4), Inches(0.3)); lab.name = 'Label'; lab.text_frame.text = '(R4 name)'
    slide('Endograft complications - Endoleak', 24, (0.2, 7.1))         # 1: 기준(이름표 좌하단, 본문 24pt)
    slide('Endograft complications - Endoleak', 14, (8.0, 3.0))         # 2: 같은 제목, 이름표가 그림 근처
    slide('Endograft complications - Endoleak', 14, (8.0, 3.0))         # 3
    path = os.path.join(TMP, 'k17.pptx'); prs.save(path)
    out = os.path.join(TMP, 'k17o.pptx')
    r = subprocess.run([sys.executable, os.path.join(HERE, 'deck_toolkit.py'), 'fit-layout', path, '-o', out, '--screens', '2-3', '--like', '1',
                        '--drop-repeat-titles', 'all'], capture_output=True, text=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1'))
    assert r.returncode == 0 and r.stdout.count('제목 뺌') == 2, (r.stdout[-900:], r.stderr[-400:])   # S1: 같은 제목 묶음 둘째 장부터
    D = T.Deck.open(out, wd('k17')); order = [x for x, _, _ in D.order() if x]
    for sn in order[1:]:
        x = open(D._slide(sn), encoding='utf8').read()
        assert not re.search(r'type="title"', x)                                              # 제목·띠 뺌
        g = lambda nm: tuple(int(v) for v in re.search(r'name="%s".*?<a:off x="(-?\d+)" y="(-?\d+)"/><a:ext cx="(\d+)" cy="(\d+)"' % nm, x, re.S).groups())
        bx, by, bw, bh = g('TextBox 4'); px, py, pw, ph = g('Pic 1'); lx, ly, lw, lh = g('Label')
        assert x.index('name="Label"') > x.index('name="Pic 1"')                             # S4: 이름표 맨 앞
        assert (lx, ly) == (int(0.2 * E), int(7.1 * E)), (lx, ly)                              # 기준 화면의 이름표 자리
        assert bx + bw <= px and by < int(0.3 * E), (bx, bw, px, by)                          # 옆 배치: 글 왼쪽·그림 오른쪽, 제목 뺀 자리까지 위로
        assert bh < int(3.23 * E) and 'sz="1400"' not in x and max(int(v) for v in re.findall(r'sz="(\d+)"', re.search(r'name="TextBox 4".*?</p:sp>', x, re.S).group(0))) <= 2400
        assert not T._inter((bx, by, bw, bh), (px, py, pw, ph)) and not T._inter((lx, ly, lw, lh), (px, py, pw, ph))
    assert T.validate(out, path)


def t_v1619_da_after_vowel():
    d = T.Deck.open(SRC, wd('b1'), theme='cud')
    d.set_notes(7, ['정답은 adenocarcinoma다. 이 병변은 MRI다.'])
    assert not [q for q in T.check_note_substitution(d) if 'slide7' in q]
    d.set_notes(8, ['이것은 segment다. 가장 중요한 point다.'])   # v16.23: 읽은 소리 기준(세그먼트다·포인트다) — 잡지 않는다
    assert not [q for q in T.check_note_substitution(d) if 'slide8' in q]
    d.set_notes(8, ['Wirsung관 확장'])                              # 한글 명사가 붙은 오염은 계속 잡는다
    assert [q for q in T.check_note_substitution(d) if 'slide8' in q]

def t_v1625_merge_lead_and_insert_tab_template():
    assert T._merge_format('L2 ⇥   US: a', 'L2 ⇥  US: b') == 'L2 ⇥   US: b'          # D9: 원래 공백 수
    assert T._merge_format('L1     Extra 22-01', 'L1 Extra 22-02') == 'L1     Extra 22-02'
    assert T._merge_format('L1 **A 22-11** x', 'L1 A 22-11 y', skip_bold=lambda s: bool(re.search(r'\d\d-\d\d', s))) == 'L1 A 22-11 y'
    e = T.Deck.open(SRC, wd('i3'))
    p, x, a, b = e._body_span(7)
    body = ('<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1600" b="1"/><a:t>Item A</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1600"><a:highlight><a:srgbClr val="FFFF00"/></a:highlight></a:rPr><a:t>\tExam 24-14</a:t></a:r></a:p>'
            '<a:p><a:pPr lvl="1"/><a:r><a:rPr lang="en-US" sz="1200" i="1"/><a:t>* note</a:t></a:r></a:p>')
    open(p, 'w', encoding='utf8').write(x[:a] + body + x[b:])
    e.insert_paragraph_like(7, 'L1 ⇥ Exam 25-01', after_key='* note')     # 앞 문단은 각주지만 틀은 탭 문단
    y = open(p, encoding='utf8').read()
    new = [q for q in re.findall(r'<a:p>.*?</a:p>', y, re.S) if 'Exam 25-01' in q][0]
    assert 'highlight' in new and 'i="1"' not in new, new

def t_title_box_not_placeholder():
    base, _ = _title_decks()
    B = T.Deck.open(base, wd('tb_box')); prof = T.title_profile(B, like=1)
    sn = [s for s, _, _ in B.order()][0]
    fp = B._slide(sn); x = open(fp, encoding='utf8').read()
    x = re.sub(r'<p:ph type="title"\s*/>', '', x, 1); x = x.replace('cy="1117331"', 'cy="548640"', 1)
    open(fp, 'w', encoding='utf8').write(x)
    bad = T.check_title_template(B, prof, screens=[1], stream=io.StringIO())
    assert any('글상자' in i for i in bad.get(sn, [])), bad
    ch = T.conform_title(B, sn, prof); assert ch[0].startswith('[참고] 글상자'), ch
    ch = T.conform_title(B, sn, prof, adopt_box=True); assert '글상자 → title placeholder' in ch, ch
    B.save('/tmp/tb_box.pptx'); assert T.validate('/tmp/tb_box.pptx', base)

# ---------------------------------------------------------------- CLI
def t_cli():
    for args in (['audit', SRC], ['lint', SRC, '--theme', 'amber'], ['plan', 'quiz', '--cases', 'A', 'B'],
                 ['media', SRC, '-o', out('cli_media')], ['verify', SRC, '--original', SRC],
                 ['restyle', SRC, '-o', out('cli_rs.pptx')],
                 ['polish', SRC, '-o', out('cli_pol.pptx'), '--font', 'office'],
                 ['handout', SRC, '-o', out('cli_ho.md')], ['sync', SRC], ['overflow', SRC]):
        r = cli(*args); assert r.returncode == 0, (args, r.stderr[-400:])



# ---------------------------------------------------------------- 노트-슬라이드 정렬 (v10 정렬 계통)
def t_align_runs():
    d = T.Deck.open(SRC, wd('al'))
    assert isinstance(T.check_notes_alignment(d, stream=io.StringIO()), list)

def t_align_detects_plane_mismatch():
    d = T.Deck.open(SRC, wd('am'))
    d.set_body(7, T.Body().line('Chest CT, axial')); d.set_notes(7, ['sagittal 재구성을 보시겠습니다'])
    probs = T.check_notes_alignment(d, stream=io.StringIO())
    assert any('slide7' in p for p in probs), probs

def t_create_notes_when_missing():
    """notesSlide 가 없는 슬라이드에 set_notes 하면 새로 만든다 (그쪽 계통 _create_notes)."""
    d = T.Deck.open(SRC, wd('cn'))
    sn = d.slide_numbers()[-1]
    rp = d._slide_rels(sn); x = open(rp, encoding='utf8').read()
    m = re.search(r'<Relationship [^>]*notesSlides/(notesSlide\d+\.xml)[^>]*/>', x)
    x = x.replace(m.group(0), ''); open(rp, 'w', encoding='utf8').write(x)
    old_nf = os.path.join(d.dir, 'ppt/notesSlides', m.group(1)); os.remove(old_nf)
    os.remove(os.path.join(d.dir, 'ppt/notesSlides/_rels', m.group(1) + '.rels'))
    assert d.notes_no(sn) is None
    d.set_notes(sn, ['new note']); assert d.notes(sn) == ['new note']
    o = out('cn.pptx'); d.save(o); assert T.validate(o, SRC)

# ---------------------------------------------------------------- 주장 관계도 / 의존 그래프
CLAIMS = [{'id': 'c1', 'statement': 't', 'sites': ['slide:7'], 'keys': ['New title'], 'forbidden': ['철회된표현']}]

def t_mapcheck_ok():
    d = T.Deck.open(SRC, wd('mp')); d.set_title(7, 'New title')
    probs, mx = T.mapcheck(d, CLAIMS, io.StringIO()); assert probs == [], probs; assert 'slide:7' in mx

def t_mapcheck_detects_missing():
    d = T.Deck.open(SRC, wd('mm')); d.set_title(7, 'Something else')
    probs, _ = T.mapcheck(d, CLAIMS, io.StringIO()); assert any('반영되지 않음' in p for p in probs), probs

def t_mapcheck_detects_forbidden():
    d = T.Deck.open(SRC, wd('mf')); d.set_title(7, 'New title'); d.set_body(7, T.Body().line('철회된표현 이 남아 있다'))
    probs, _ = T.mapcheck(d, CLAIMS, io.StringIO()); assert any('철회된 표현' in p for p in probs), probs

def t_mapreport_runs():
    d = T.Deck.open(SRC, wd('mr2')); assert T.mapreport(d, CLAIMS, io.StringIO()) == ['slide:7']

import test_claim_graph as _TCG
from test_claim_graph import GRAPH
for _n in dir(_TCG):
    if _n.startswith('t_'):
        globals()[_n] = getattr(_TCG, _n)

def t_freeze_then_stale_pptx():
    import copy
    d = T.Deck.open(SRC, wd('fz')); d.set_title(7, 'title x'); d.set_notes(7, ['note y']); d.set_notes(8, ['z'])
    g = T.mapfreeze(d, copy.deepcopy(GRAPH), at='2026-01-01')
    r = T.mapstale(d, g, io.StringIO()); assert r['changed'] == [] and r['unverified'] == [], r
    d.set_notes(7, ['note y changed'])
    r = T.mapstale(d, g, io.StringIO()); assert r['changed'] == ['b'] and [x[0] for x in r['suspect']] == ['c'], r

def t_unmapped_sites():
    d = T.Deck.open(SRC, wd('us')); d.set_body(7, T.Body().line('one').line('two').line('three'))
    assert 'slide:7' in T.unmapped_sites(d, [], io.StringIO())
    assert 'slide:7' not in T.unmapped_sites(d, [{'id': 'x', 'sites': ['slide:7']}], io.StringIO())

def t_cli_graph_pptx():
    import copy, json
    p = out('claims_cli.json'); T.save_claims(p, copy.deepcopy(GRAPH))
    for args in (['mapgraph', '--claims', p], ['impact', '--claims', p, 'a'],
                 ['mapfreeze', SRC, '--claims', p, '-o', out('claims_fz.json')],
                 ['mapstale', SRC, '--claims', out('claims_fz.json')],
                 ['extract', SRC, '-o', out('dk_ex.json')], ['scaffold', '--claims', out('dk_ex.json')],
                 ['align', SRC]):
        r = cli(*args); assert r.returncode == 0, (args, r.stderr[-400:])
    assert 'verified' in json.load(open(out('claims_fz.json')))['claims'][0]

# ---------------------------------------------------------------- 실행 (한 번만)
tests = [(n[2:], f) for n, f in list(globals().items()) if n.startswith('t_')]
for name, fn in tests:
    check(name, fn)
w = max(len(n) for _, n, _ in results)
for st, n, msg in results:
    print('%-4s %-*s %s' % (st, w, n, msg))
_p = sum(1 for s, _, _ in results if s == 'PASS')
_s = sum(1 for s, _, _ in results if s == 'SKIP')
_f = sum(1 for s, _, _ in results if s == 'FAIL')
print('\n통과 %d / 건너뜀 %d / 실패 %d  (전체 %d)' % (_p, _s, _f, len(results)))
if _s:
    print('건너뜀은 입력 파일 부재입니다. 결함은 아니지만 검증된 것도 아닙니다.')
if _f:
    print('실패가 있으면 코드 결함입니다.')
    sys.exit(1)
